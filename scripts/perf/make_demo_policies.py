#!/usr/bin/env python3
"""Builds two demo-style SYNTHETIC policies in the style of the user's RA UAT policy (numbered sections, tables, part tags, traceability
matrix, not-selected / withdrawn / inherited lines) plus requirement-level gold labels:
  backend/tests/fixtures/policies/{ac,sc}_policy_raven_demo.docx   docs/perf/corpus/{ac,sc}_raven.gold.json
Gold rule: every requirement of an in-scope control is "met" unless its (control, part) – or a keyword in its text – is a deliberate gap."""
import json
import re
import sys
from pathlib import Path

import docx

sys.path.insert(0, "/app/backend")
FIX = Path("/app/backend/tests/fixtures/policies")
CORPUS = Path("/app/docs/perf/corpus")
HDR = "DEMONSTRATION TEMPLATE – SYNTHETIC. System: The Raven (DoD AWS cloud workload). Approver names are placeholders."

AC = {
    "key": "ac", "title": "Access Control (AC) Policy — DoD AWS Cloud System",
    "doc": [["Item", "Value"], ["Document", "Access Control (AC) Policy"], ["System", "The Raven"], ["Version", "1.0"], ["Owner", "ISSM (placeholder: J. Doe)"],
            ["Review cycle", "At least annually"]],
    "sections": [
        ("1. Purpose", ["This policy establishes access control requirements for The Raven at the system level. It addresses purpose, scope, roles, responsibilities, "
                        "management commitment, coordination among organizational entities, and compliance. (AC-1a.1(a))",
                        "Management is committed to enforcing least privilege and accountable access for every user and service."]),
        ("2. Scope and Applicability", ["This policy applies to all military, civilian and contractor personnel and to all service accounts that access The Raven, "
                                         "including the AWS management console, workloads, and CI/CD pipelines."]),
        ("3. Authority and References", ["This policy is consistent with applicable laws, executive orders, directives, regulations, policies, standards and guidelines, "
                                          "including those listed below. (AC-1a.1(b))"],
         [["Reference", "Relevance"], ["DoDI 8500.01", "Cybersecurity"], ["DoDI 8510.01", "RMF for DoD Systems"], ["NIST SP 800-53 Rev 5", "Control catalog"],
          ["DoD Cloud Computing SRG", "Impact level and connection requirements"]]),
        ("4. Roles and Responsibilities", ["Roles and their access control responsibilities are listed below. Coordination among the ISSM, ISSO, system owner and "
                                            "the AWS Cloud Service Provider is performed through the monthly security working group."],
         [["Role", "Responsibility"], ["ISSM", "Designated official who manages the development, documentation and dissemination of this policy and its procedures"],
          ["ISSO", "Account manager for The Raven; approves and reviews accounts"], ["System Owner", "Defines authorized users and access authorizations"]]),
        ("5. Policy Management", [
            "The ISSM is the designated official who manages the development, documentation, and dissemination of the access control policy and procedures. (AC-1b)",
            "Access control procedures are developed and documented to facilitate the implementation of this policy and the associated access controls. (AC-1a.2)",
            "This policy and the procedures are disseminated to all system owners, ISSOs, administrators and users of The Raven through the eMASS artifact library. (AC-1a)",
            "The policy shall be reviewed and updated at least annually and following significant system changes, security incidents or assessment findings. "
            "The procedures shall be reviewed and updated at least annually and following the same events. (AC-1c.1, AC-1c.2)"]),
        ("6. Account Management (AC-2)", [
            "Account types allowed for use are individual user, privileged administrator, service, and emergency accounts; shared, guest and anonymous accounts are prohibited. (AC-2a)",
            "The ISSO is assigned as account manager for all accounts of The Raven. (AC-2b)",
            "Conditions for group and role membership are defined in the access control procedures; membership requires a documented mission need. (AC-2c)",
            "For each account the authorized users, group and role membership, and access authorizations (privileges) and other attributes are specified in the access request record. (AC-2d)",
            "Requests to create accounts require approval by the System Owner and the ISSO. (AC-2e)",
            "Accounts are created, enabled, modified, disabled and removed in accordance with the access control procedures and the Account Lifecycle table below. (AC-2f)",
            "The use of accounts is monitored through AWS CloudTrail and IAM access analyzer reports. (AC-2g)",
            "Access to the system is authorized based on a valid access authorization, intended system usage, and mission or business functions. (AC-2i)",
            "Accounts are reviewed for compliance with account management requirements every 90 days. (AC-2j)",
            "A process is established for changing shared or group account authenticators when individuals are removed from the group; emergency account credentials are rotated after each use. (AC-2k)",
            "Account management processes are aligned with personnel termination and transfer processes. (AC-2l)"],
         [["Lifecycle event", "Action", "Time limit"], ["Approved request", "Create and enable account", "2 business days"], ["Role change", "Modify privileges", "1 business day"],
          ["Inactivity", "Disable account", "After 35 days of inactivity"], ["Departure", "Remove account", "Same day"]]),
        ("6.1 Account management enhancements", [
            "Temporary and emergency accounts are automatically removed or disabled after 72 hours. (AC-2(2))",
            "Accounts are disabled within 24 hours when they have expired, are no longer associated with a user or individual, are in violation of organizational policy, "
            "or have been inactive for 35 days. (AC-2(3))",
            "AC-2(6) Dynamic Privilege Management is not selected for The Raven.",
            "AC-2(10) is withdrawn in Rev 5 (incorporated into AC-2k)."]),
        ("7. Access Enforcement (AC-3)", ["Approved authorizations for logical access to information and system resources are enforced in accordance with applicable access "
                                           "control policies by AWS IAM policies, security groups and application role checks. (AC-3)"]),
        ("8. Least Privilege (AC-6)", [
            "The principle of least privilege is employed, allowing only authorized accesses for users and processes acting on behalf of users that are necessary to "
            "accomplish assigned organizational tasks. (AC-6)",
            "Access is authorized to security functions and security-relevant information only for the ISSM, ISSO and privileged administrators; the security functions "
            "include IAM, KMS key policy, CloudTrail and GuardDuty configuration. (AC-6(1))",
            "Users of privileged accounts are required to use non-privileged accounts when accessing nonsecurity functions. (AC-6(2))"]),
        ("9. Unsuccessful Logon Attempts (AC-7)", ["A limit of three consecutive invalid logon attempts by a user during a 15-minute period is enforced. (AC-7a)"]),
        ("10. System Use Notification (AC-8)", [
            "The Raven displays the standard DoD Notice and Consent Banner before granting access, providing privacy and security notices consistent with applicable laws "
            "and policies, stating that users are accessing a U.S. Government system, that usage may be monitored, recorded and subject to audit, that unauthorized use is "
            "prohibited and subject to penalties, and that use indicates consent to monitoring and recording. (AC-8a)",
            "The notification message remains on the screen until users acknowledge the usage conditions and take explicit actions to log on. (AC-8b)"]),
        ("11. Device Lock and Session Termination (AC-11, AC-12)", [
            "Further access to the system is prevented by initiating a device lock after 15 minutes of inactivity, and the lock is retained until the user reestablishes access "
            "using established identification and authentication procedures. (AC-11a, AC-11b)",
            "A user session is automatically terminated after 8 hours or after 15 minutes of inactivity for privileged sessions. (AC-12)"]),
        ("12. Remote Access (AC-17)", [
            "Usage restrictions, configuration and connection requirements, and implementation guidance are established and documented for each type of remote access allowed: "
            "AWS Client VPN with FIPS 140-validated TLS and the AWS Systems Manager Session Manager for administrators. (AC-17a)",
            "Each type of remote access to the system is authorized by the ISSO prior to allowing such connections. (AC-17b)",
            "Remote access methods are monitored and controlled by automated mechanisms through VPN connection logs and CloudTrail session records. (AC-17(1))",
            "Cryptographic mechanisms (FIPS 140-validated TLS 1.2 or higher) are implemented to protect the confidentiality and integrity of remote access sessions. (AC-17(2))",
            "Remote accesses are routed through two authorized and managed network access control points (the VPN endpoints in the transit VPC). (AC-17(3))",
            "The execution of privileged commands and access to security-relevant information via remote access is authorized only for operational needs approved by the ISSM. (AC-17(4)a)",
            "AC-17(5), AC-17(7) and AC-17(8) are withdrawn."]),
        ("13. AWS Shared Responsibility and Inheritance", ["Physical access to AWS data centers is inherited from the AWS GovCloud FedRAMP High authorization and is outside this policy."]),
    ],
    "matrix": [["Control", "Title", "Section", "Status"], ["AC-1", "Policy and Procedures", "1–5", "Addressed"], ["AC-2", "Account Management", "6", "Addressed"],
               ["AC-2(2)", "Automated Temporary and Emergency Account Management", "6.1", "Addressed"], ["AC-2(3)", "Disable Accounts", "6.1", "Addressed"],
               ["AC-2(6)", "Dynamic Privilege Management", "6.1", "Not selected"], ["AC-2(10)", "Shared and Group Account Credential Change", "6.1", "Withdrawn"],
               ["AC-3", "Access Enforcement", "7", "Addressed"], ["AC-6", "Least Privilege", "8", "Addressed"], ["AC-6(1)", "Authorize Access to Security Functions", "8", "Addressed"],
               ["AC-6(2)", "Non-privileged Access for Nonsecurity Functions", "8", "Addressed"], ["AC-7", "Unsuccessful Logon Attempts", "9", "Addressed"],
               ["AC-8", "System Use Notification", "10", "Addressed"], ["AC-11", "Device Lock", "11", "Addressed"], ["AC-12", "Session Termination", "11", "Addressed"],
               ["AC-17", "Remote Access", "12", "Addressed"], ["AC-17(1)", "Monitoring and Control", "12", "Addressed"], ["AC-17(2)", "Protection of Confidentiality and Integrity Using Encryption", "12", "Addressed"],
               ["AC-17(3)", "Managed Access Control Points", "12", "Addressed"], ["AC-17(4)", "Privileged Commands and Access", "12", "Addressed"]],
    "in_scope": ["ac-1", "ac-2", "ac-2.2", "ac-2.3", "ac-3", "ac-6", "ac-6.1", "ac-6.2", "ac-7", "ac-8", "ac-11", "ac-12", "ac-17", "ac-17.1", "ac-17.2", "ac-17.3", "ac-17.4"],
    "gaps": {("ac-2", "h"): "Notification of account managers (time periods) is not stated.", ("ac-7", "b"): "No action when the logon limit is exceeded.",
             ("ac-8", "c"): "Publicly accessible system conditions are not stated.", ("ac-17.4", "b"): "Rationale for remote privileged access is not documented in the policy."},
    "kw_gaps": {},
    "not_selected": {"ac-2.6": "AC-2(6) Dynamic Privilege Management is not selected for The Raven."},
    "withdrawn": ["ac-2.10", "ac-17.5", "ac-17.7", "ac-17.8"],
    "negative": ["ra-3", "ra-5", "sc-7", "sc-8", "sc-13", "sc-28", "cp-9", "ir-4", "au-6", "cm-6", "si-4", "pe-3", "mp-6", "ps-3", "ca-7"],
}

SC = {
    "key": "sc", "title": "System and Communications Protection (SC) Policy — DoD AWS Cloud System",
    "doc": [["Item", "Value"], ["Document", "System and Communications Protection (SC) Policy"], ["System", "The Raven"], ["Version", "1.0"], ["Owner", "ISSM (placeholder: J. Doe)"]],
    "sections": [
        ("1. Purpose", ["This policy establishes system and communications protection requirements for The Raven. It addresses purpose, scope, roles, responsibilities, "
                        "management commitment, coordination among organizational entities, and compliance. (SC-1a.1(a))"]),
        ("2. Scope and Applicability", ["This policy applies to all network, boundary, cryptographic and DNS components of The Raven in AWS GovCloud, and to all personnel who administer them."]),
        ("3. Authority and References", ["This policy is consistent with applicable laws, executive orders, directives, regulations, policies, standards and guidelines. (SC-1a.1(b))"],
         [["Reference", "Relevance"], ["DoDI 8510.01", "RMF"], ["CNSSP 15", "Use of public standards for secure information sharing"], ["DoD Cloud Computing SRG", "BCAP and DoDIN connection"]]),
        ("4. Roles and Responsibilities", ["The ISSM, ISSO, network engineer and the AWS Cloud Service Provider coordinate boundary changes through the change control board."],
         [["Role", "Responsibility"], ["ISSM", "Designated official who manages the development, documentation and dissemination of this policy and its procedures"],
          ["Network Engineer", "Maintains boundary protection devices and DNS"], ["Crypto Custodian", "Manages KMS keys and certificates"]]),
        ("5. Policy Management", [
            "The ISSM is the designated official who manages the development, documentation, and dissemination of the system and communications protection policy and procedures. (SC-1b)",
            "System and communications protection procedures are developed and documented to facilitate the implementation of this policy and the associated controls. (SC-1a.2)",
            "This policy and the procedures are disseminated to all system owners, ISSOs and network administrators through the eMASS artifact library. (SC-1a)",
            "The policy shall be reviewed and updated at least annually and following significant system changes, security incidents or assessment findings. "
            "The procedures shall be reviewed and updated at least annually and following the same events. (SC-1c.1, SC-1c.2)"]),
        ("6. Denial-of-Service Protection (SC-5)", ["The Raven protects against or limits the effects of volumetric and application-layer denial-of-service events using AWS Shield Advanced and AWS WAF rate-based rules. (SC-5a)"]),
        ("7. Boundary Protection (SC-7)", [
            "Communications at the external managed interfaces to the system and at key internal managed interfaces are monitored and controlled by AWS Network Firewall and security groups. (SC-7a)",
            "Connections to external networks or systems are made only through managed interfaces consisting of boundary protection devices arranged in accordance with the DoD "
            "Cloud Computing SRG; all DoDIN traffic traverses the BCAP. (SC-7c)",
            "The number of external network connections to the system is limited to the BCAP connection and the AWS Client VPN endpoint. (SC-7(3))",
            "For each external telecommunication service a managed interface is implemented, a traffic flow policy is established for each managed interface, the confidentiality "
            "and integrity of the information being transmitted across each interface is protected, each exception to the traffic flow policy is documented with a supporting "
            "mission or business need and duration of that need, exceptions are reviewed every 90 days and removed when no longer supported, unauthorized exchange of control "
            "plane traffic with external networks is prevented, information is published to enable remote networks to detect unauthorized control plane traffic, and "
            "unauthorized control plane traffic from external networks is filtered. (SC-7(4))",
            "Network communications traffic is denied by default and allowed by exception at managed interfaces. (SC-7(5))",
            "Split tunneling is prevented for remote devices connecting to the system unless the tunnel is securely provisioned using the approved VPN configuration. (SC-7(7))"],
         [["Interface", "Device", "Policy"], ["BCAP / DoDIN", "AWS Network Firewall", "Deny by default"], ["Internet egress", "NAT + Network Firewall", "Allow-list only"]]),
        ("8. Transmission Confidentiality and Integrity (SC-8)", [
            "The confidentiality and integrity of transmitted information is protected. (SC-8)",
            "Cryptographic mechanisms (TLS 1.2 or higher with FIPS 140-validated modules) are implemented to prevent unauthorized disclosure of information and detect changes to information during transmission. (SC-8(1))",
            "SC-9 and SC-14 are withdrawn (incorporated into SC-8 and other controls)."]),
        ("9. Network Disconnect (SC-10)", ["The network connection associated with a communications session is terminated at the end of the session or after 15 minutes of inactivity. (SC-10)"]),
        ("10. Cryptographic Key Establishment and Management (SC-12)", [
            "Cryptographic keys are established and managed when cryptography is employed within the system in accordance with key generation and key storage requirements: "
            "keys are generated in AWS KMS FIPS 140-validated HSMs and stored only in KMS. (SC-12)"]),
        ("11. Cryptographic Protection (SC-13)", ["The cryptographic uses are identified as data in transit, data at rest and digital signatures, and the types of cryptography "
                                                   "required for each use are FIPS 140-validated and NSA-approved cryptography where classified. (SC-13a, SC-13b)"]),
        ("12. Public Key Infrastructure Certificates (SC-17)", ["Public key certificates are issued under the DoD PKI certificate policy or obtained from an approved service provider, "
                                                                 "and only approved trust anchors are included in trust stores or certificate stores managed by the organization. (SC-17a, SC-17b)"]),
        ("13. Secure Name / Address Resolution (SC-20, SC-21, SC-22)", [
            "Additional data origin authentication and integrity verification artifacts are provided along with the authoritative name resolution data the system returns in response "
            "to external name/address resolution queries, using DNSSEC signing in Route 53. (SC-20a)",
            "Data origin authentication and data integrity verification on the name/address resolution responses the system receives from authoritative sources are requested and performed (DNSSEC validation in the Route 53 Resolver). (SC-21)",
            "The systems that collectively provide name/address resolution service are fault-tolerant and implement internal and external role separation. (SC-22)"]),
        ("14. Session Authenticity (SC-23)", ["The authenticity of communications sessions is protected by TLS with mutual authentication for service-to-service traffic. (SC-23)"]),
        ("15. Protection of Information at Rest (SC-28)", [
            "The confidentiality and integrity of all CUI, audit logs and backups at rest is protected. (SC-28)",
            "Cryptographic mechanisms (AWS KMS AES-256 encryption) are implemented to prevent unauthorized disclosure and modification of CUI, audit logs and backups at rest on EBS, S3 and RDS. (SC-28(1))"]),
        ("16. Process Isolation (SC-39)", ["A separate execution domain is maintained for each executing system process; this is inherited from the AWS Nitro hypervisor and supplemented by container isolation. (SC-39)"]),
        ("17. Collaborative Computing Devices", ["SC-15 is not selected for The Raven: the system has no collaborative computing devices."]),
    ],
    "matrix": [["Control", "Title", "Section", "Status"], ["SC-1", "Policy and Procedures", "1–5", "Addressed"], ["SC-5", "Denial-of-Service Protection", "6", "Addressed"],
               ["SC-7", "Boundary Protection", "7", "Addressed"], ["SC-7(3)", "Access Points", "7", "Addressed"], ["SC-7(4)", "External Telecommunications Services", "7", "Addressed"],
               ["SC-7(5)", "Deny by Default — Allow by Exception", "7", "Addressed"], ["SC-7(7)", "Split Tunneling for Remote Devices", "7", "Addressed"],
               ["SC-8", "Transmission Confidentiality and Integrity", "8", "Addressed"], ["SC-8(1)", "Cryptographic Protection", "8", "Addressed"], ["SC-9", "Transmission Confidentiality", "8", "Withdrawn"],
               ["SC-10", "Network Disconnect", "9", "Addressed"], ["SC-12", "Cryptographic Key Establishment and Management", "10", "Addressed"], ["SC-13", "Cryptographic Protection", "11", "Addressed"],
               ["SC-14", "Public Access Protections", "8", "Withdrawn"], ["SC-15", "Collaborative Computing Devices and Applications", "17", "Not selected"],
               ["SC-17", "Public Key Infrastructure Certificates", "12", "Addressed"], ["SC-20", "Secure Name/Address Resolution Service (Authoritative Source)", "13", "Addressed"],
               ["SC-21", "Secure Name/Address Resolution Service (Recursive or Caching Resolver)", "13", "Addressed"], ["SC-22", "Architecture and Provisioning for Name/Address Resolution Service", "13", "Addressed"],
               ["SC-23", "Session Authenticity", "14", "Addressed"], ["SC-28", "Protection of Information at Rest", "15", "Addressed"], ["SC-28(1)", "Cryptographic Protection", "15", "Addressed"],
               ["SC-39", "Process Isolation", "16", "Inherited (AWS)"]],
    "in_scope": ["sc-1", "sc-5", "sc-7", "sc-7.3", "sc-7.4", "sc-7.5", "sc-7.7", "sc-8", "sc-8.1", "sc-10", "sc-12", "sc-13", "sc-17", "sc-20", "sc-21", "sc-22", "sc-23", "sc-28", "sc-28.1", "sc-39"],
    "gaps": {("sc-5", "b"): "The controls employed per DoS type are not stated as a separate objective.", ("sc-7", "b"): "Separation of publicly accessible components is not stated.",
             ("sc-20", "b"): "Chain-of-trust / child-subspace security status is not stated."},
    "kw_gaps": {"sc-12": r"distribut|destruct|access"},
    "not_selected": {"sc-15": "SC-15 is not selected for The Raven: the system has no collaborative computing devices."},
    "withdrawn": ["sc-9", "sc-14"],
    "negative": ["ac-17", "ac-2", "ac-6", "ra-3", "ra-5", "ia-2", "au-6", "cp-9", "ir-4", "cm-6", "pe-3", "si-4", "ps-3", "mp-6", "ca-7"],
}


def build_docx(p: dict, path: Path) -> None:
    d = docx.Document()
    d.add_paragraph(p["title"])
    d.add_paragraph(HDR)
    _table(d, p["doc"])
    for sec in p["sections"]:
        d.add_paragraph(sec[0])
        for para in sec[1]:
            d.add_paragraph(para)
        if len(sec) > 2:
            _table(d, sec[2])
    d.add_paragraph(f"{len(p['sections']) + 1}. Control Traceability Matrix")
    d.add_paragraph("Every in-scope control maps to a section of this policy or is marked not selected, withdrawn or inherited.")
    _table(d, p["matrix"])
    d.add_paragraph(f"{len(p['sections']) + 2}. Approval")
    _table(d, [["Role", "Name", "Signature", "Date"], ["Authorizing Official", "Placeholder AO", "(signed)", "2026-10-01"]])
    d.save(path)


def _table(d, rows: list) -> None:
    t = d.add_table(rows=len(rows), cols=len(rows[0]))
    for i, r in enumerate(rows):
        for j, v in enumerate(r):
            t.cell(i, j).text = v


def _part(req: dict) -> str:
    rid, ref = req["id"], str(req.get("ref") or "")
    if "_obj." in rid:
        m = re.match(r"[a-z]", rid.split("_obj.")[1])
        return m.group(0) if m else "-"
    m = re.match(r"[A-Z]{2}-\d+(?:\s*\(\d+\))?\s*\(?([a-z])\)?", ref)
    return m.group(1) if m else "-"


def gold(p: dict) -> dict:
    from app.db import SessionLocal
    from app.services import policy_coverage as pc
    labels = []
    with SessionLocal() as db:
        for c in p["in_scope"]:
            for r in pc.requirements(db, c):
                lab, note = "met", None
                if (c, _part(r)) in p["gaps"]:
                    lab, note = "not_addressed", p["gaps"][(c, _part(r))]
                elif c in p["kw_gaps"] and re.search(p["kw_gaps"][c], r["text"], re.I):
                    lab, note = "not_addressed", "Only key generation and storage are stated."
                labels.append({"control_id": c, "requirement_id": r["id"], "part": _part(r), "label": lab, "note": note})
    return {"document": f"backend/tests/fixtures/policies/{p['key']}_policy_raven_demo.docx (SYNTHETIC demo, built by scripts/perf/make_demo_policies.py)",
            "labelled_by": "construction: in-scope control parts are written to be addressed except the listed deliberate gaps", "requirements": labels,
            "not_selected": p["not_selected"], "withdrawn": p["withdrawn"], "negative_controls": p["negative"],
            "negative_rule": "never a proposed map, coverage row (gap or covered) or audit finding"}


if __name__ == "__main__":
    for p in (AC, SC):
        build_docx(p, FIX / f"{p['key']}_policy_raven_demo.docx")
        g = gold(p)
        (CORPUS / f"{p['key']}_raven.gold.json").write_text(json.dumps(g, indent=1))
        print(p["key"], len(g["requirements"]), "requirements,", sum(x["label"] != "met" for x in g["requirements"]), "gold gaps")
