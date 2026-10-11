"""Generate SYNTHETIC sample policies (TXT/DOCX/PDF) for the Policy Auditor. Intentional defects for the demo:
35-day inactivity (observed config = 90), 3 lockout attempts, 15-minute session lock, cites NIST SP 800-53 Rev 4,
no remote-access section, no named official."""
from pathlib import Path

OUT = Path(__file__).resolve().parent.parent / "samples" / "policies"
TITLE = "IronVision Demo Agency – Access Control Policy (SYNTHETIC)"
SECTIONS = [
    ("1. Purpose", "This policy establishes access control requirements for agency information systems in accordance with NIST SP 800-53 Rev 4."),
    ("2. Scope", "Applies to all agency systems, including the Logistics Tracking Platform (LTP). All content is synthetic demo data."),
    ("3. Roles", "The policy is maintained by the Office of the CISO. Approving official: TBD"),
    ("4. Account Management", "Accounts that remain inactive are disabled after 35 days of inactivity. Account reviews occur quarterly."),
    ("5. Unsuccessful Logon Attempts", "Systems lock the account after 3 consecutive invalid logon attempts within 15 minutes."),
    ("6. Session Lock", "A session lock is initiated after 15 minutes of inactivity and is retained until the user re-authenticates."),
    ("7. Review", "This policy is reviewed annually."),
]


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    text = TITLE + "\n\n" + "\n\n".join(f"{h}\n{b}" for h, b in SECTIONS) + "\n"
    (OUT / "access_control_policy.txt").write_text(text)
    import docx
    d = docx.Document()
    d.add_heading(TITLE, 1)
    for h, b in SECTIONS:
        d.add_heading(h, 2)
        d.add_paragraph(b)
    d.save(OUT / "access_control_policy.docx")
    from reportlab.lib.pagesizes import letter
    from reportlab.pdfgen import canvas
    c = canvas.Canvas(str(OUT / "access_control_policy.pdf"), pagesize=letter)
    y = 750
    for line in text.splitlines():
        c.drawString(50, y, line[:110])
        y -= 16
    c.save()
    print("wrote", sorted(p.name for p in OUT.iterdir()))


if __name__ == "__main__":
    main()
