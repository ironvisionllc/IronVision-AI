"""eMASS Implementation Plan + Risk Assessment fields (eMASS REST API `PUT /api/systems/{systemId}/controls`, MITRE emass_client
OpenAPI spec). Enum values, character limits and the conditional-required table are copied from that spec; see
docs/RMF_INTERPRETATION_DECISIONS.md §eMASS Implementation Plan.

Status is NOT stored twice: ControlAssignment.implementation_status stays the single source of truth and the eMASS
Implementation Status is derived from it (EmassControlFields.manually_inherited only splits Inherited vs Manually Inherited).
"""
from datetime import date, timedelta

from fastapi import HTTPException
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models import (CommonControlProvider, ControlAssignment, EmassControlFields, Evidence, EvidenceControl, EvidenceSource, Finding,
                        FindingControlMap, RoleAssignment, System, TailoringDecision, User)

STATUSES = ["Planned", "Implemented", "Inherited", "Not Applicable", "Manually Inherited"]
DESIGNATIONS = ["Common", "System-Specific", "Hybrid"]
PROVIDERS = ["DoD", "Component", "Enclave"]
FREQUENCIES = ["Constantly", "Daily", "Weekly", "Monthly", "Quarterly", "Semi-Annually", "Annually", "Every Two Years", "Every Three Years", "Undetermined"]
METHODS = ["Automated", "Semi-Automated", "Manual", "Undetermined"]
LEVELS = ["Very Low", "Low", "Moderate", "High", "Very High"]

SLCM = ["slcm_criticality", "slcm_frequency", "slcm_method", "slcm_reporting", "slcm_tracking", "slcm_comments"]
PLAN_FIELDS = ["control_designation", "estimated_completion_date", "responsible_entities", *SLCM, "na_justification", "common_control_provider"]
RISK_ENUMS = ["severity", "relevance_of_threat", "likelihood", "impact", "residual_risk_level"]
RISK_TEXT = ["impact_description", "vulnerability_summary", "recommendations", "mitigations"]
RISK_FIELDS = RISK_ENUMS + RISK_TEXT
ENUMS = {"control_designation": DESIGNATIONS, "common_control_provider": PROVIDERS, "slcm_frequency": FREQUENCIES, "slcm_method": METHODS,
         **{k: LEVELS for k in RISK_ENUMS}}
LIMITS = {**{k: 2000 for k in ("responsible_entities", "slcm_criticality", "slcm_reporting", "slcm_tracking", "na_justification",
                               "impact_description", "vulnerability_summary", "recommendations", "mitigations")}, "slcm_comments": 4000}
REQUIRED = {"Planned": ["control_designation", "estimated_completion_date", "responsible_entities", *SLCM],
            "Implemented": ["control_designation", "estimated_completion_date", "responsible_entities", *SLCM],
            "Not Applicable": ["na_justification"],
            "Manually Inherited": ["common_control_provider", *SLCM],
            "Inherited": []}
NARRATIVE_LIMIT = 2000

# column label (eMASS UI wording) and eMASS API field name; vulnerability_summary has no API field (UI-only, see docs)
FIELD_META = {
    "implementation_status": ("Implementation Status", "implementationStatus"),
    "common_control_provider": ("Common Control Provider", "commonControlProvider"),
    "na_justification": ("N/A Justification", "naJustification"),
    "control_designation": ("Security Control Designation", "controlDesignation"),
    "estimated_completion_date": ("Estimated Completion Date", "estimatedCompletionDate"),
    "responsible_entities": ("Responsible Entities", "responsibleEntities"),
    "implementation_narrative": ("Implementation Narrative", "implementationNarrative"),
    "slcm_criticality": ("SLCM Criticality", "slcmCriticality"),
    "slcm_frequency": ("SLCM Frequency", "slcmFrequency"),
    "slcm_method": ("SLCM Method", "slcmMethod"),
    "slcm_reporting": ("SLCM Reporting", "slcmReporting"),
    "slcm_tracking": ("SLCM Tracking", "slcmTracking"),
    "slcm_comments": ("SLCM Comments", "slcmComments"),
    "severity": ("Severity", "severity"),
    "relevance_of_threat": ("Relevance of Threat", "relevanceOfThreat"),
    "likelihood": ("Likelihood", "likelihood"),
    "impact": ("Impact", "impact"),
    "impact_description": ("Impact Description", "impactDescription"),
    "residual_risk_level": ("Residual Risk Level", "residualRiskLevel"),
    "vulnerability_summary": ("Vulnerability Summary", None),
    "recommendations": ("Recommendations", "recommendations"),
    "mitigations": ("Mitigations", "mitigations"),
}
AUTOMATED_SOURCES = {"acas", "stig_checklist", "sarif"}
MANUAL_SOURCES = {"manual_upload", "manual_attestation", "policy_audit"}


def emass_status(a: ControlAssignment, f: EmassControlFields | None) -> str:
    s = a.implementation_status
    if s == "implemented":
        return "Implemented"
    if s == "not_applicable":
        return "Not Applicable"
    if s == "inherited":
        return "Manually Inherited" if f is not None and f.manually_inherited else "Inherited"
    return "Planned"


def internal_status(current: str, status: str) -> str:
    if status == "Planned":
        return current if current in ("planned", "partially_implemented", "not_implemented") else "planned"
    return {"Implemented": "implemented", "Not Applicable": "not_applicable", "Inherited": "inherited", "Manually Inherited": "inherited"}[status]


def has_inheritance(a: ControlAssignment) -> bool:
    return a.provider_id is not None or a.responsibility == "inherited"


def fields_for(db: Session, a: ControlAssignment) -> EmassControlFields | None:
    return db.scalar(select(EmassControlFields).where(EmassControlFields.assignment_id == a.id))


def ensure(db: Session, a: ControlAssignment) -> EmassControlFields:
    f = fields_for(db, a)
    if f is None:
        f = EmassControlFields(system_id=a.system_id, assignment_id=a.id, control_id=a.control_id)
        db.add(f)
        db.flush()
    return f


def values(f: EmassControlFields | None) -> dict:
    return {k: (getattr(f, k) if f is not None else None) for k in PLAN_FIELDS + RISK_FIELDS}


def missing_required(status: str, vals: dict) -> list[str]:
    return [k for k in REQUIRED[status] if vals.get(k) in (None, "")]


def missing_all(status: str, vals: dict, narrative: str | None) -> list[str]:
    """Single rule used by the panel, the Workbench filter and the xlsx/validation report. implementationNarrative is in
    eMASS ControlsRequiredFields (required on every controls PUT); Inherited controls take it from the providing system."""
    narr = [] if status == "Inherited" or (narrative or "").strip() else ["implementation_narrative"]
    return missing_required(status, vals) + narr


def narratives(db: Session, system_id) -> dict:
    from app.models import Implementation
    return dict(db.execute(select(Implementation.assignment_id, Implementation.narrative).join(ControlAssignment, ControlAssignment.id == Implementation.assignment_id)
                           .where(ControlAssignment.system_id == system_id, Implementation.is_current.is_(True))).all())


def check_values(vals: dict) -> None:
    """422 on enum / char-limit violations (eMASS business rules)."""
    errs = []
    for k, v in vals.items():
        if v in (None, ""):
            continue
        if k in ENUMS and v not in ENUMS[k]:
            errs.append(f"{FIELD_META[k][0]}: '{v}' is not an eMASS value ({', '.join(ENUMS[k])})")
        if k in LIMITS and isinstance(v, str) and len(v) > LIMITS[k]:
            errs.append(f"{FIELD_META[k][0]} exceeds the eMASS limit of {LIMITS[k]:,} characters ({len(v):,})")
    if errs:
        raise HTTPException(422, "; ".join(errs))


def _narrative(db: Session, a: ControlAssignment) -> str:
    from app.models import Implementation
    i = db.scalar(select(Implementation).where(Implementation.assignment_id == a.id, Implementation.is_current.is_(True)))
    return (i.narrative if i else "") or ""


def provider_info(db: Session, a: ControlAssignment) -> dict | None:
    if not has_inheritance(a):
        return None
    p = db.get(CommonControlProvider, a.provider_id) if a.provider_id else None
    return {"provider": p.name if p else None, "package": p.package if p else None, "basis": p.basis if p else None,
            "responsibility": a.responsibility, "provider_responsibility": a.provider_responsibility_text,
            "system_responsibility": a.system_responsibility_text}


def view(db: Session, a: ControlAssignment, can_edit: bool) -> dict:
    f = fields_for(db, a)
    st = emass_status(a, f)
    vals = values(f)
    narr = _narrative(db, a)
    return {"implementation_status": st, "internal_status": a.implementation_status, "values": vals,
            "missing_required": missing_all(st, vals, narr), "required": REQUIRED[st],
            "inherited_locked": st == "Inherited", "inheritance": provider_info(db, a),
            "narrative_length": len(narr), "narrative_over_limit": len(narr) > NARRATIVE_LIMIT,
            "assisted": (f.assisted if f else None) or {}, "updated_at": f.updated_at.isoformat() if f else None,
            "risk_suggestion": risk_matrix(vals.get("likelihood"), vals.get("impact")),
            "can_edit": can_edit, "options": options()}


def options() -> dict:
    return {"statuses": STATUSES, "designations": DESIGNATIONS, "providers": PROVIDERS, "frequencies": FREQUENCIES, "methods": METHODS,
            "levels": LEVELS, "required": REQUIRED, "limits": LIMITS, "labels": {k: v[0] for k, v in FIELD_META.items()}}


def risk_matrix(likelihood: str | None, impact: str | None) -> dict | None:
    """Display-only suggestion (NIST SP 800-30 Rev 1 Table I-2 style 5x5 lookup); never auto-set."""
    if likelihood not in LEVELS or impact not in LEVELS:
        return None
    li, ii = LEVELS.index(likelihood), LEVELS.index(impact)
    table = [[0, 0, 0, 1, 1], [0, 1, 1, 1, 2], [0, 1, 2, 2, 3], [0, 1, 2, 3, 4], [0, 1, 2, 3, 4]]
    return {"residual_risk_level": LEVELS[table[li][ii]], "basis": f"Likelihood {likelihood} × Impact {impact} (SP 800-30 Table I-2 lookup) – suggestion only"}


def _audit(db, current, a, action, before, after, extra=None, request_meta=None):
    from app.services import audit
    changed = {k for k in set(before) | set(after) if before.get(k) != after.get(k)}
    audit.record(db, action=action, actor_user_id=current.user.id, actor_email=current.user.email, entity_type="control_assignment",
                 entity_id=str(a.id), before={k: _j(before.get(k)) for k in sorted(changed)}, after={k: _j(after.get(k)) for k in sorted(changed)},
                 details={"control_id": a.control_id, **(extra or {})}, **(request_meta or {}))
    return changed


def _j(v):
    return v.isoformat() if isinstance(v, date) else v


def _assisted(f: EmassControlFields, fields: list[str], source: str | None, interaction_id, user) -> None:
    if not source or not fields:
        return
    from datetime import datetime, timezone
    stamp = {"source": source, "interaction_id": str(interaction_id) if interaction_id else None, "by": user.email,
             "at": datetime.now(timezone.utc).isoformat()}
    f.assisted = {**(f.assisted or {}), **{k: stamp for k in fields}}


def save_plan(db: Session, a: ControlAssignment, status: str, vals: dict, current, assist: dict, meta: dict) -> dict:
    vals = {k: (v.strip() if isinstance(v, str) else v) or None for k, v in vals.items() if k in PLAN_FIELDS}
    check_values(vals)
    f = ensure(db, a)
    before_status = emass_status(a, f)
    if status == "Inherited" and not has_inheritance(a):
        raise HTTPException(422, "Inherited requires inheritance data from a providing system or common control provider – none is recorded "
                                 "for this control. Use Manually Inherited and name the Common Control Provider instead.")
    before = {**values(f), "implementation_status": before_status}
    if status == "Inherited":
        if before_status == "Inherited" and any(vals.get(k) != before.get(k) for k in vals):
            raise HTTPException(422, "Inherited controls take their Implementation Plan from the providing system – these fields are read-only")
        vals = {}
    miss = missing_required(status, {**values(f), **vals} if status == "Inherited" else vals)
    if miss:
        raise HTTPException(422, f"eMASS requires these fields when Implementation Status is {status}: "
                                 + ", ".join(FIELD_META[k][0] for k in miss))
    for k, v in vals.items():
        setattr(f, k, v)
    f.manually_inherited = status == "Manually Inherited"
    new_internal = internal_status(a.implementation_status, status)
    prev_internal = a.implementation_status
    a.implementation_status = new_internal
    f.updated_by = current.user.id
    used = [k for k in assist.get("fields", []) if k in vals]
    _assisted(f, used, assist.get("source"), assist.get("interaction_id"), current.user)
    changed = _audit(db, current, a, "emass.plan_updated", before, {**values(f), "implementation_status": status},
                     {"ai_assisted": assist.get("source") == "ai" and bool(used), "suggestion_source": assist.get("source") if used else None,
                      "suggested_fields_accepted": used, "interaction_id": str(assist.get("interaction_id")) if used and assist.get("interaction_id") else None},
                     meta)
    if new_internal != prev_internal:
        from app.services import audit
        audit.record(db, action="implementation.status_changed", actor_user_id=current.user.id, actor_email=current.user.email,
                     entity_type="control_assignment", entity_id=str(a.id), before={"implementation_status": prev_internal},
                     after={"implementation_status": new_internal, "emass_status": status, "via": "emass_implementation_plan"}, **meta)
    return {"changed": sorted(changed), "status_changed": new_internal != prev_internal}


def save_risk(db: Session, a: ControlAssignment, vals: dict, current, assist: dict, meta: dict) -> dict:
    f = ensure(db, a)
    if emass_status(a, f) == "Inherited":
        raise HTTPException(422, "Risk Assessment information cannot be updated when the control is Inherited (eMASS business rule)")
    vals = {k: (v.strip() if isinstance(v, str) else v) or None for k, v in vals.items() if k in RISK_FIELDS}
    check_values(vals)
    before = values(f)
    for k, v in vals.items():
        setattr(f, k, v)
    f.updated_by = current.user.id
    used = [k for k in assist.get("fields", []) if k in vals]
    _assisted(f, used, assist.get("source"), assist.get("interaction_id"), current.user)
    changed = _audit(db, current, a, "emass.risk_updated", before, values(f),
                     {"ai_assisted": assist.get("source") == "ai" and bool(used), "suggestion_source": assist.get("source") if used else None,
                      "suggested_fields_accepted": used}, meta)
    return {"changed": sorted(changed)}


BULK_FIELDS = ["control_designation", "estimated_completion_date", "responsible_entities", *SLCM, "common_control_provider"]


def bulk(db: Session, s: System, control_ids: list[str], vals: dict, current, meta: dict) -> dict:
    vals = {k: (v.strip() if isinstance(v, str) else v) for k, v in vals.items() if k in BULK_FIELDS and v not in (None, "")}
    if not vals:
        raise HTTPException(422, "Choose at least one Implementation Plan field to set")
    check_values(vals)
    from app.services import audit
    done, skipped = [], []
    for a in db.scalars(select(ControlAssignment).where(ControlAssignment.system_id == s.id, ControlAssignment.control_id.in_([c.lower() for c in control_ids]))):
        f = ensure(db, a)
        st = emass_status(a, f)
        if st == "Inherited":
            skipped.append({"control_id": a.control_id, "reason": "Inherited – values come from the providing system"})
            continue
        if st == "Not Applicable":
            skipped.append({"control_id": a.control_id, "reason": "Not Applicable – Implementation Plan fields do not apply"})
            continue
        before = values(f)
        for k, v in vals.items():
            setattr(f, k, v)
        f.updated_by = current.user.id
        _audit(db, current, a, "emass.plan_updated", before, values(f), {"bulk": True}, meta)
        done.append(a.control_id)
    audit.record(db, action="emass.plan_bulk_updated", actor_user_id=current.user.id, actor_email=current.user.email, entity_type="system",
                 entity_id=str(s.id), after={k: _j(v) for k, v in vals.items()}, details={"updated": done, "skipped": skipped}, **meta)
    return {"updated": done, "skipped": skipped, "fields": sorted(vals)}


# ---------- grounding + deterministic suggestions ----------
def _people(db: Session, s: System) -> dict[str, list[str]]:
    out: dict[str, list[str]] = {}
    for role, name in db.execute(select(RoleAssignment.role, User.display_name).join(User, User.id == RoleAssignment.user_id)
                                 .where(RoleAssignment.scope_id == s.id, RoleAssignment.role.in_(("isso", "system_owner", "issm")))):
        out.setdefault(role, []).append(name)
    return out


def grounding(db: Session, s: System, a: ControlAssignment) -> dict:
    from app.api.v1.workbench import _control_poams, _policies_for
    ev = db.execute(select(Evidence, EvidenceSource.source_type).join(EvidenceControl, EvidenceControl.evidence_id == Evidence.id)
                    .outerjoin(EvidenceSource, EvidenceSource.id == Evidence.source_id).where(EvidenceControl.assignment_id == a.id)).all()
    fs = db.scalars(select(Finding).join(FindingControlMap, FindingControlMap.finding_id == Finding.id)
                    .where(Finding.system_id == s.id, FindingControlMap.control_id == a.control_id, Finding.current_state == "open").distinct()).all()
    poams = [p for p in _control_poams(db, s.id, a.control_id) if p["status"] not in ("closed", "risk_accepted", "completed")]
    td = db.scalar(select(TailoringDecision).where(TailoringDecision.system_id == s.id, TailoringDecision.control_id == a.control_id))
    return {
        "system": {"name": s.name, "acronym": s.acronym, "baseline": s.baseline_key, "categorization": s.categorization,
                   "confidentiality": s.confidentiality_impact, "integrity": s.integrity_impact, "availability": s.availability_impact},
        "control": {"id": a.control_id, "internal_status": a.implementation_status, "responsibility": a.responsibility, "origin": a.origin},
        "inheritance": provider_info(db, a),
        "people": _people(db, s),
        "evidence": [{"name": e.name, "source_type": st or e.evidence_type, "state": e.state, "collected": e.collected_at.date().isoformat()} for e, st in ev][:15],
        "open_findings": [{"title": f.title[:160], "cat": f.cat, "severity": f.severity, "kev": bool(f.is_kev)} for f in fs][:15],
        "open_poams": [{"id": p["human_id"], "weakness": (p["weakness"] or "")[:160], "due": str(p["due_date"]) if p["due_date"] else None} for p in poams],
        "policies": _policies_for(db, s.id, a.control_id),
        "tailoring_justification": td.justification if td else None,
        "current_values": {k: _j(v) for k, v in values(fields_for(db, a)).items() if v},
    }


def _max_level(*vals) -> str:
    order = ["low", "moderate", "high"]
    best = max((order.index(v.lower()) for v in vals if v and v.lower() in order), default=1)
    return ["Low", "Moderate", "High"][best]


def rule_suggestions(ctx: dict, status: str) -> list[dict]:
    """Deterministic fallback (no AI): every value is derived from the listed records."""
    out = []
    add = lambda k, v, why: out.append({"field": k, "value": v, "reason": why})  # noqa: E731
    inh, ev, fs, poams, sysd = ctx["inheritance"], ctx["evidence"], ctx["open_findings"], ctx["open_poams"], ctx["system"]
    resp = ctx["control"]["responsibility"]
    if status == "Not Applicable":
        add("na_justification", ctx["tailoring_justification"] or f"Not applicable to {sysd['acronym']}: [state why – e.g. the capability does not exist in the boundary].",
            "From the recorded tailoring decision" if ctx["tailoring_justification"] else "No tailoring decision recorded – placeholder to complete")
    else:
        if status != "Manually Inherited":
            add("control_designation", {"inherited": "Common", "shared": "Hybrid"}.get(resp, "System-Specific"), f"Responsibility is '{resp}' in the inheritance record")
            if status == "Implemented":
                latest = max((e["collected"] for e in ev if e["state"] in ("current", "near_expiration")), default=None)
                add("estimated_completion_date", latest or date.today().isoformat(),
                    f"Implemented: date of the latest current evidence ({latest})" if latest else "Implemented with no current evidence – today's date")
            else:
                due = max((p["due"] for p in poams if p["due"]), default=None)
                add("estimated_completion_date", due or (date.today() + timedelta(days=90)).isoformat(),
                    f"Latest open POA&M milestone due date ({due})" if due else "No open POA&M due date – 90 days from today")
            ppl = ctx["people"]
            who = [f"ISSO ({', '.join(ppl['isso'])})" if ppl.get("isso") else "ISSO", f"System Owner ({', '.join(ppl['system_owner'])})" if ppl.get("system_owner") else "System Owner"]
            if inh and inh.get("provider"):
                who.append(f"Common control provider: {inh['provider']}")
            add("responsible_entities", "; ".join(who), "From the system's role assignments" + (" and inheritance record" if inh and inh.get("provider") else ""))
        else:
            name = (inh or {}).get("provider") or ""
            add("common_control_provider", "DoD" if "dod" in name.lower() else ("Enclave" if "enclave" in name.lower() else "Component"),
                f"Provider recorded as '{name}'" if name else "No provider recorded – Component assumed; confirm")
        crit = "High" if any(f["cat"] == "I" or f["kev"] for f in fs) else _max_level(sysd["confidentiality"], sysd["integrity"], sysd["availability"])
        add("slcm_criticality", crit, "Open CAT I / KEV finding" if crit == "High" and fs else f"Highest CIA impact of the system ({sysd['categorization'] or 'n/a'})")
        auto = [e for e in ev if e["source_type"] in AUTOMATED_SOURCES]
        manual = [e for e in ev if e["source_type"] not in AUTOMATED_SOURCES]
        add("slcm_frequency", "Monthly" if auto else ("Annually" if manual else "Undetermined"),
            f"{len(auto)} scanner evidence item(s) (STIG/ACAS/SARIF) – monthly ConMon scans" if auto else
            (f"{len(manual)} manual/policy evidence item(s) – annual review" if manual else "No evidence linked yet"))
        add("slcm_method", "Automated" if auto and not manual else ("Semi-Automated" if auto else ("Manual" if manual else "Undetermined")),
            f"{len(auto)} automated and {len(manual)} manual evidence item(s)")
        add("slcm_reporting", "Monthly ConMon report to the ISSM" if auto else "Annual security review report to the ISSM", "Matches the SLCM frequency above")
        add("slcm_tracking", "eMASS POA&M" if fs or poams else "ConMon tracker reviewed by the ISSO",
            f"{len(fs)} open finding(s), {len(poams)} open POA&M(s)")
        add("slcm_comments", f"Derived from {len(ev)} evidence item(s), {len(fs)} open finding(s) and {len(poams)} open POA&M(s)"
            + (f"; governing policy: {'; '.join(ctx['policies'])}" if ctx["policies"] else "; no governing policy linked") + ".", "Summary of the records used")
    if status == "Inherited":
        return []
    kev, cat1 = any(f["kev"] for f in fs), any(f["cat"] == "I" for f in fs)
    sev = "Very High" if kev else "High" if cat1 else "Moderate" if any(f["cat"] == "II" for f in fs) else "Low" if fs else "Very Low"
    add("severity", sev, f"Worst open finding: {'KEV' if kev else 'CAT I' if cat1 else (fs[0]['cat'] and 'CAT ' + fs[0]['cat']) if fs else 'none'}")
    lik = "High" if kev or cat1 else "Moderate" if fs else "Low"
    add("likelihood", lik, f"{len(fs)} open finding(s)" + (" incl. KEV" if kev else ""))
    add("relevance_of_threat", "Very High" if kev else "Moderate" if fs else "Low", "KEV = known exploitation" if kev else f"{len(fs)} open finding(s)")
    imp = _max_level(sysd["confidentiality"], sysd["integrity"], sysd["availability"])
    add("impact", imp, f"System categorization C/I/A = {sysd['confidentiality']}/{sysd['integrity']}/{sysd['availability']}")
    rr = risk_matrix(lik, imp)
    add("residual_risk_level", rr["residual_risk_level"], rr["basis"])
    add("impact_description", f"Failure of {ctx['control']['id'].upper()} could affect the {imp.lower()}-impact information of {sysd['name']}.", "From the categorization")
    if fs:
        add("vulnerability_summary", "; ".join(f"{f['title']} ({'CAT ' + f['cat'] if f['cat'] else f['severity']})" for f in fs[:5]), "Open findings mapped to this control")
        add("recommendations", "Remediate the open findings listed above and re-scan to confirm closure.", "Open findings exist")
    if poams:
        add("mitigations", "; ".join(f"{p['id']}: {p['weakness']}" + (f" (due {p['due']})" if p["due"] else "") for p in poams[:5]), "Open POA&Ms for this control")
    return out


def clean_suggestions(items: list[dict], status: str) -> tuple[list[dict], list[dict]]:
    """Keep only known fields with valid eMASS values; never compliance / test results."""
    allowed = set(RISK_FIELDS) | set(PLAN_FIELDS) if status != "Inherited" else set()
    ok, rejected = [], []
    for it in items:
        k, v = it.get("field"), it.get("value")
        why = None
        if k not in allowed:
            why = "not an editable Implementation Plan / Risk Assessment field"
        elif v in (None, ""):
            why = "empty"
        elif k in ENUMS and v not in ENUMS[k]:
            why = f"'{v}' is not an eMASS value"
        elif k in LIMITS and len(str(v)) > LIMITS[k]:
            why = "exceeds eMASS character limit"
        elif k == "estimated_completion_date":
            try:
                date.fromisoformat(str(v))
            except ValueError:
                why = "not an ISO date"
        (rejected if why else ok).append({**it, "rejected_reason": why} if why else {"field": k, "value": v, "reason": str(it.get("reason") or "")[:400]})
    seen, dedup = set(), []
    for it in ok:
        if it["field"] not in seen:
            seen.add(it["field"])
            dedup.append(it)
    return dedup, rejected


# ---------- exports ----------
def rows(db: Session, s: System) -> list[dict]:
    from app.models import Control
    fmap = {f.assignment_id: f for f in db.scalars(select(EmassControlFields).where(EmassControlFields.system_id == s.id))}
    narr = narratives(db, s.id)
    out = []
    for a, title, label in db.execute(select(ControlAssignment, Control.title, Control.label).join(Control, Control.id == ControlAssignment.control_uuid)
                                      .where(ControlAssignment.system_id == s.id).order_by(Control.sort_id)).all():
        f = fmap.get(a.id)
        st = emass_status(a, f)
        vals = values(f)
        n = narr.get(a.id) or ""
        miss = missing_all(st, vals, n)
        out.append({"control_id": a.control_id, "acronym": label, "title": title, "implementation_status": st, **vals,
                    "implementation_narrative": n, "missing": miss, "narrative_over_limit": len(n) > NARRATIVE_LIMIT})
    return out


def validation_report(db: Session, s: System) -> dict:
    rs = rows(db, s)
    bad = [{"control": r["acronym"], "implementation_status": r["implementation_status"], "missing": [FIELD_META[k][0] for k in r["missing"]],
            "narrative_over_limit": r["narrative_over_limit"]} for r in rs if r["missing"] or r["narrative_over_limit"]]
    return {"system": s.acronym, "controls": len(rs), "rows_with_issues": len(bad), "complete": len(rs) - len(bad), "items": bad,
            "rules": "eMASS PUT /api/systems/{systemId}/controls business rules (MITRE emass_client OpenAPI); narrative limit 2,000 characters"}


def api_payload(r: dict) -> dict:
    """One control in eMASS API field names (estimatedCompletionDate as Unix time, per the spec)."""
    from datetime import datetime, timezone
    body = {"acronym": r["acronym"]}
    for k, (_label, api) in FIELD_META.items():
        if api is None or r.get(k) in (None, ""):
            continue
        v = r[k]
        if k == "estimated_completion_date":
            v = int(datetime.combine(v, datetime.min.time(), timezone.utc).timestamp())
        body[api] = v
    return body


XLSX_COLUMNS = ["acronym", "title", "implementation_status", "common_control_provider", "na_justification", "control_designation",
                "estimated_completion_date", "responsible_entities", "implementation_narrative", "slcm_criticality", "slcm_frequency", "slcm_method",
                "slcm_reporting", "slcm_tracking", "slcm_comments", "severity", "relevance_of_threat", "likelihood", "impact", "impact_description",
                "residual_risk_level", "vulnerability_summary", "recommendations", "mitigations"]
BANNER = "eMASS-formatted, compatibility pending human verification – generated by IronVision AI from eMASS API field definitions; not an official eMASS template"


def xlsx(db: Session, s: System) -> bytes:
    import io

    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    rs = rows(db, s)
    wb = Workbook()
    ws = wb.active
    ws.title = "Implementation Plan"
    ws.append([BANNER])
    ws["A1"].font = Font(bold=True, color="9C0006")
    ws.append([("Control Acronym" if k == "acronym" else "Control Title" if k == "title" else FIELD_META[k][0]) for k in XLSX_COLUMNS])
    for c in ws[2]:
        c.font, c.fill = Font(bold=True), PatternFill("solid", fgColor="DDEBF7")
    for r in rs:
        ws.append([_cell(r.get(k)) for k in XLSX_COLUMNS])
    ws.freeze_panes = "B3"
    m = wb.create_sheet("Field mapping")
    m.append(["Column", "eMASS API field (PUT /api/systems/{systemId}/controls)", "Allowed values / limit"])
    for k in XLSX_COLUMNS[2:]:
        label, api = FIELD_META[k]
        m.append([label, api or "(no API field – eMASS UI only; exported for reference)",
                  ", ".join(ENUMS[k]) if k in ENUMS else (", ".join(STATUSES) if k == "implementation_status" else
                                                          f"max {LIMITS.get(k, NARRATIVE_LIMIT if k == 'implementation_narrative' else 0):,} chars" if k in LIMITS or k == "implementation_narrative" else "")])
    v = wb.create_sheet("Validation")
    v.append(["Control", "Implementation Status", "Missing required eMASS fields", "Narrative over 2,000 chars"])
    for it in validation_report(db, s)["items"]:
        v.append([it["control"], it["implementation_status"], ", ".join(it["missing"]), "yes" if it["narrative_over_limit"] else ""])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _cell(v):
    if isinstance(v, str):
        v = v[:32000]
        return "'" + v if v[:1] in ("=", "+", "-", "@") else v
    return v


def oscal_props(f: EmassControlFields | None, a: ControlAssignment) -> tuple[list[dict], str | None]:
    """Props (enum/date values) + markdown remarks (free text) for an OSCAL implemented-requirement."""
    ns = "https://ironvisionai.com/ns/oscal/emass"
    props = [{"name": "emass-implementation-status", "ns": ns, "value": emass_status(a, f)}]
    if f is None:
        return props, None
    for k in ("control_designation", "common_control_provider", "slcm_frequency", "slcm_method", *RISK_ENUMS):
        if getattr(f, k):
            props.append({"name": "emass-" + k.replace("_", "-"), "ns": ns, "value": getattr(f, k)})
    if f.estimated_completion_date:
        props.append({"name": "emass-estimated-completion-date", "ns": ns, "value": f.estimated_completion_date.isoformat()})
    text = [f"**{FIELD_META[k][0]}:** {getattr(f, k).strip()}" for k in ("responsible_entities", "slcm_criticality", "slcm_reporting", "slcm_tracking",
                                                                       "slcm_comments", "na_justification", *RISK_TEXT) if (getattr(f, k) or "").strip()]
    return props, ("eMASS Implementation Plan / Risk Assessment (IronVision AI):\n\n" + "\n\n".join(text)) if text else None
