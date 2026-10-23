"""eMASS HW/SW List template (.xlsm): header-role detection, import (never overwrites existing inventory), bounded export.
Macros are never executed; the package is copied byte-for-byte, so the VBA project and its signatures stay identical."""
import hashlib
import io
import re
import uuid
import zipfile
from datetime import date, datetime

import openpyxl
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.models import Asset, AssetCorrelation, InventoryDiscrepancy, InventoryExport, InventoryImport, SoftwareItem, System
from app.services import jobs
from app.services.inventory import correlate as cr
from app.services.tr import xlsx

HW = {"asset_no": ("#", "asset id"), "component_type": ("component type",), "machine_name": ("machine name",), "nickname": ("nickname",),
      "ip": ("ip address",), "virtual": ("virtual asset",), "manufacturer": ("manufacturer",), "model": ("model number",),
      "serial": ("serial number",), "os": ("os/ios/fw version", "os version", "operating system"), "memory": ("memory",),
      "location": ("location",), "approval": ("approval status",), "critical": ("critical information system asset",)}
SW = {"asset_no": ("#",), "sw_type": ("software type",), "vendor": ("software vendor",), "name": ("software name",), "version": ("version",),
      "subsystem": ("subsystem",), "license": ("license or contract",), "total_licenses": ("total licenses",), "licenses_used": ("licenses used",),
      "license_poc": ("license poc",), "license_expiration": ("license expiration",), "approval": ("approval status",), "approval_date": ("approval date",),
      "critical": ("critical information",), "purpose": ("purpose",)}
LABELS = {"date_exported": "date exported", "exported_by": "exported by", "system_name": "system name"}
INV_FIELDS = ("asset_no", "nickname", "virtual", "manufacturer", "model", "memory", "location", "approval", "critical")
XLSM = "application/vnd.ms-excel.sheet.macroEnabled.12"


def _norm(v) -> str:
    return re.sub(r"\s+", " ", re.sub(r"\(required\)", "", str(v or ""), flags=re.I)).strip().lower().rstrip("?:").strip()


def _val(v):
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip() if v is not None and str(v).strip() else None


def _roles(ws, spec) -> tuple[int | None, dict]:
    for r in range(1, min(ws.max_row, 40) + 1):
        found: dict = {}
        for c in ws[r]:
            h = _norm(c.value)
            role = next((k for k, syn in spec.items() if k not in found and any(h == s or h.startswith(s) for s in syn)), None) if h else None
            if role:
                found[role] = c.column_letter
        if len(found) >= 4:
            return r, found
    return None, {}


def _resolve(wb, f: str) -> list[str]:
    f = f.lstrip("=")
    if f.startswith('"'):
        return [x.strip() for x in f.strip('"').split(",")]
    dn = wb.defined_names.get(f)
    ref = dn.attr_text if dn is not None else f
    m = re.match(r"(\w+)\[#Data\]", ref)
    if m:
        for ws in wb.worksheets:
            if m.group(1) in ws.tables:
                rows = ws[ws.tables[m.group(1)].ref]
                return [str(r[0].value) for r in rows[1:] if r[0].value is not None]
        return []
    m = re.match(r"'?([^'!]+)'?!\$?([A-Z]+)\$?(\d+):\$?([A-Z]+)\$?(\d+)", ref)
    return [str(c.value) for row in wb[m.group(1)][f"{m.group(2)}{m.group(3)}:{m.group(4)}{m.group(5)}"] for c in row if c.value is not None] if m else []


def _vocab(wb, ws, cols: dict) -> dict:
    rev, out = {v: k for k, v in cols.items()}, {}
    for dv in ws.data_validations.dataValidation:
        if dv.type == "list" and dv.formula1:
            vals = _resolve(wb, dv.formula1)
            for rng in str(dv.sqref).split():
                col = re.match(r"[A-Z]+", rng).group(0)
                if col in rev and vals:
                    out[rev[col]] = vals
    return out


def _header_cells(ws, header_row: int) -> dict:
    merged = {(m.min_row, m.min_col): m for m in ws.merged_cells.ranges}
    out = {}
    for r in range(1, header_row):
        for c in ws[r]:
            h = _norm(c.value)
            for role, label in LABELS.items():
                if h.startswith(label) and role not in out:
                    m = merged.get((c.row, c.column))
                    out[role] = f"{openpyxl.utils.get_column_letter((m.max_col if m else c.column) + 1)}{r}"
    return out


def detect(data: bytes) -> dict:
    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    out: dict = {"sheets": wb.sheetnames, "has_vba": "xl/vbaProject.bin" in zipfile.ZipFile(io.BytesIO(data)).namelist()}
    for kind, spec, need in (("hardware", HW, {"machine_name"}), ("software", SW, {"vendor", "name"})):
        for ws in wb.worksheets:
            hr, cols = _roles(ws, spec)
            if hr and need <= set(cols):
                rows = []
                for r in range(hr + 1, ws.max_row + 1):
                    vals = {role: _val(ws[f"{col}{r}"].value) for role, col in cols.items()}
                    if any(v for k, v in vals.items() if k != "asset_no"):
                        rows.append({"row": r, "values": vals})
                out[kind] = {"sheet": ws.title, "header_row": hr, "cols": cols, "vocab": _vocab(wb, ws, cols), "rows": rows,
                             "header_cells": _header_cells(ws, hr)}
                break
    if "hardware" not in out or "software" not in out:
        raise ValueError("Not an eMASS HW/SW List template: could not find the Hardware (Machine Name) and Software (Vendor/Name) header rows")
    return out


def upload(db: Session, system: System, file_name: str, data: bytes, actor) -> InventoryImport:
    from app.services.storage.base import get_storage

    sha = hashlib.sha256(data).hexdigest()
    ver = (db.scalar(select(func.max(InventoryImport.version)).where(InventoryImport.system_id == system.id)) or 0) + 1
    key = f"inventory/{system.org_id}/{system.id}/imports/v{ver}/{sha[:12]}-{re.sub(r'[^A-Za-z0-9._-]', '_', file_name)}"
    get_storage().put(key, data, XLSM if file_name.lower().endswith(".xlsm") else "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    imp = InventoryImport(system_id=system.id, org_id=system.org_id, file_name=file_name, sha256=sha, storage_key=key, version=ver, status="queued",
                          uploaded_by=actor.email)
    db.add(imp)
    db.flush()
    jobs.enqueue(db, "inventory.import", {"import_id": str(imp.id), "actor_email": actor.email}, requested_by=actor.id)
    return imp


def _propose(db: Session, sid, key: str, title: str, recs: list, change: dict, imp: InventoryImport) -> bool:
    if db.scalar(select(InventoryDiscrepancy.id).where(InventoryDiscrepancy.system_id == sid, InventoryDiscrepancy.key == key)):
        return False
    db.add(InventoryDiscrepancy(system_id=sid, kind="import_conflict", key=key, status="open", title=title, records=recs, proposed_change=change,
                                recommendation="Accept to propose the imported value; it is applied only after approval. Ignore to keep the current inventory value.",
                                details={"import_id": str(imp.id), "file": imp.file_name, "version": imp.version}, last_detected_at=cr.now()))
    return True


def run_import(db: Session, imp: InventoryImport, actor_email: str | None) -> dict:
    from app.services.inventory.consistency import rec_asset
    from app.services.storage.base import get_storage

    det = detect(get_storage().get(imp.storage_key))
    sid = imp.system_id
    c = {"hardware_rows": 0, "software_rows": 0, "created": 0, "unchanged": 0, "proposed_changes": 0, "linked_scan_assets": 0, "errors": []}
    imp_rec = {"type": "import", "id": str(imp.id), "label": f"{imp.file_name} v{imp.version}", "source": "hwsw_import", "source_label": "eMASS HW/SW import",
               "timestamp": imp.created_at.isoformat()}
    for r in det["hardware"]["rows"]:
        v = r["values"]
        c["hardware_rows"] += 1
        mn = v.get("machine_name")
        if not mn:
            c["errors"].append({"sheet": det["hardware"]["sheet"], "row": r["row"], "reason": "Machine Name is required"})
            continue
        ips = [x for x in (cr.n_ip(p) for p in re.split(r"[\s,;]+", v.get("ip") or "")) if x]
        facts = {"host": mn, "ips": ips, "serial": v.get("serial")}
        same = [a for a in inventory_assets(db, sid) if cr.n_host(a.hostname or a.name) == cr.n_host(mn)]
        d = {"status": "matched", "asset_id": str(same[0].id), "method": "deterministic:machine_name"} if len(same) == 1 else cr.evaluate(db, sid, facts)
        target = db.get(Asset, uuid.UUID(d["asset_id"])) if d["status"] == "matched" else None
        inv = {k: v[k] for k in INV_FIELDS if v.get(k)}
        if target is not None and not target.is_unmatched:
            diffs = []
            for field, new in (("os", v.get("os")), ("serial_number", v.get("serial")), ("asset_type", v.get("component_type"))):
                if new and (getattr(target, field) or None) != new:
                    diffs.append({"type": "set_field", "asset_id": str(target.id), "field": field, "from": getattr(target, field), "to": new})
            for k, new in inv.items():
                if (target.inventory or {}).get(k) != new:
                    diffs.append({"type": "set_field", "asset_id": str(target.id), "field": f"inventory.{k}", "from": (target.inventory or {}).get(k), "to": new})
            have = {x for x in map(cr.n_ip, [target.ip_address, *(target.ip_addresses or [])]) if x}
            diffs += [{"type": "add_ip", "asset_id": str(target.id), "value": ip} for ip in ips if ip not in have]
            for ch in diffs:
                what = ch.get("field") or "ip"
                c["proposed_changes"] += _propose(db, sid, f"import:{target.id}:{what}:{ch.get('to') or ch.get('value')}",
                                                  f"Import row {r['row']} for “{target.name}”: {what} “{ch.get('from') or '–'}” → “{ch.get('to') or ch.get('value')}”",
                                                  [rec_asset(target), imp_rec], ch, imp)
            c["unchanged"] += not diffs
            continue
        a = Asset(system_id=sid, name=mn, hostname=cr.n_host(mn), ip_address=ips[0] if ips else None, ip_addresses=ips or None,
                  serial_number=v.get("serial"), os=v.get("os"), asset_type=v.get("component_type") or "unknown", inventory={**inv, "template_row": r["row"]},
                  inventory_source="hwsw_import", review_status="confirmed", source="hwsw_import",
                  provenance={"import_id": str(imp.id), "sha256": imp.sha256, "sheet": det["hardware"]["sheet"], "row": r["row"], "imported_by": actor_email})
        db.add(a)
        db.flush()
        cr.record_identifiers(db, a, facts, "hwsw_import", str(imp.id))
        c["created"] += 1
        if target is not None:  # scan-observed host matched on a reliable identifier → deterministic link (findings unchanged)
            cr.link(db, target, a)
            for corr in db.scalars(select(AssetCorrelation).where(AssetCorrelation.observed_asset_id == target.id)):
                corr.chosen_asset_id, corr.method = a.id, d["method"] + " (inventory import)"
                corr.status = "matched" if corr.status == "new" else corr.status
            c["linked_scan_assets"] += 1
    for r in det["software"]["rows"]:
        v = r["values"]
        c["software_rows"] += 1
        if not (v.get("vendor") and v.get("name")):
            c["errors"].append({"sheet": det["software"]["sheet"], "row": r["row"], "reason": "Software Vendor and Software Name are required"})
            continue
        sw = db.scalar(select(SoftwareItem).where(SoftwareItem.system_id == sid, func.lower(SoftwareItem.vendor) == v["vendor"].lower(),
                                                  func.lower(SoftwareItem.name) == v["name"].lower()))
        if sw is None:
            db.add(SoftwareItem(system_id=sid, vendor=v["vendor"], name=v["name"], version=v.get("version"), sw_type=v.get("sw_type"), source="hwsw_import",
                                import_id=imp.id, attributes={**{k: x for k, x in v.items() if x and k not in ("vendor", "name", "version", "sw_type")}, "template_row": r["row"]}))
            c["created"] += 1
        elif v.get("version") and v["version"] != sw.version:
            c["proposed_changes"] += _propose(db, sid, f"import:sw:{sw.id}:version:{v['version']}", f"Import row {r['row']}: {sw.vendor} {sw.name} version “{sw.version}” → “{v['version']}”",
                                              [{"type": "software", "id": str(sw.id), "label": f"{sw.vendor} {sw.name} {sw.version}", "source": sw.source,
                                                "source_label": "eMASS HW/SW import", "timestamp": sw.created_at.isoformat()}, imp_rec],
                                              {"type": "set_software", "software_id": str(sw.id), "field": "version", "from": sw.version, "to": v["version"]}, imp)
        else:
            c["unchanged"] += 1
    imp.detection = {k: ({kk: vv for kk, vv in x.items() if kk != "rows"} | {"row_count": len(x["rows"])} if isinstance(x, dict) else x) for k, x in det.items()}
    imp.counts, imp.status = c, "imported"
    db.flush()
    jobs.enqueue(db, "inventory.check", {"system_id": str(sid)})
    return c


def inventory_assets(db: Session, sid) -> list[Asset]:
    return list(db.scalars(select(Asset).where(Asset.system_id == sid, Asset.is_unmatched.is_(False), Asset.correlated_asset_id.is_(None)).order_by(Asset.name)))


def export(db: Session, system: System, actor) -> InventoryExport:
    from app.services.storage.base import get_storage

    imp = db.scalar(select(InventoryImport).where(InventoryImport.system_id == system.id, InventoryImport.status == "imported").order_by(InventoryImport.version.desc()))
    if imp is None:
        raise ValueError("Upload the eMASS HW/SW List template (.xlsm) first – the export is written into that workbook")
    store = get_storage()
    base = store.get(imp.storage_key)
    det = detect(base)
    hw, sw = det["hardware"], det["software"]
    assets = inventory_assets(db, system.id)
    softs = list(db.scalars(select(SoftwareItem).where(SoftwareItem.system_id == system.id).order_by(SoftwareItem.vendor, SoftwareItem.name)))
    meta = {"date_exported": f"{date.today():%d-%b-%Y}", "exported_by": actor.email, "system_name": system.name}

    def cells(spec, items, valf):
        out = {ref: meta[k] for k, ref in spec["header_cells"].items()}
        start = spec["header_row"] + 1
        last = max([start + len(items) - 1, *[r["row"] for r in spec["rows"]]])
        for i in range(last - start + 1):
            vals = valf(items[i]) if i < len(items) else {}
            for role, col in spec["cols"].items():
                out[f"{col}{start + i}"] = vals.get(role)
        return out

    def hwv(a: Asset):
        inv = a.inventory or {}
        ips = [x for x in dict.fromkeys([a.ip_address, *(a.ip_addresses or [])]) if x]
        return {**{k: inv.get(k) for k in INV_FIELDS}, "component_type": a.asset_type if a.asset_type not in (None, "unknown") else None, "machine_name": a.name,
                "ip": ", ".join(ips) or None, "serial": a.serial_number, "os": a.os}

    def swv(s: SoftwareItem):
        return {**{k: v for k, v in (s.attributes or {}).items() if k in SW}, "vendor": s.vendor, "name": s.name, "version": s.version, "sw_type": s.sw_type}

    hc, sc = cells(hw, assets, hwv), cells(sw, softs, swv)
    mid = xlsx.write(base, hw["sheet"], hc)
    out = xlsx.write(mid, sw["sheet"], sc)
    v1, v2 = xlsx.verify(base, mid, hw["sheet"], set(hc)), xlsx.verify(mid, out, sw["sheet"], set(sc))
    za, zb = zipfile.ZipFile(io.BytesIO(base)), zipfile.ZipFile(io.BytesIO(out))
    vba = {n: {"sha256": hashlib.sha256(za.read(n)).hexdigest(), "unchanged": za.read(n) == zb.read(n)} for n in za.namelist() if "vba" in n.lower()}

    def outside(spec, cs):
        allowed_cols, hdr = set(spec["cols"].values()), set(spec["header_cells"].values())
        return [r for r in cs if r not in hdr and (re.match(r"[A-Z]+", r).group(0) not in allowed_cols or int(re.search(r"\d+", r).group(0)) <= spec["header_row"])]

    warnings = [f"{hw['sheet']} row {hw['header_row'] + 1 + i}: “{val}” is not in the {role} list (eMASS accepts new values)"
                for i, a in enumerate(assets) for role, val in hwv(a).items() if val and role in hw["vocab"] and val not in hw["vocab"][role]]
    rep = {"ok": v1["ok"] and v2["ok"] and all(x["unchanged"] for x in vba.values()) and not outside(hw, hc) and not outside(sw, sc),
           "sheets": {hw["sheet"]: v1, sw["sheet"]: v2}, "vba_parts": vba, "outside_permitted": outside(hw, hc) + outside(sw, sc),
           "changed_cells": [f"{hw['sheet']}!{r}" for r in hc] + [f"{sw['sheet']}!{r}" for r in sc], "warnings": warnings[:50],
           "permitted_range": [f"{s['sheet']}!{c}{s['header_row'] + 1}:{c}…" for s in (hw, sw) for c in sorted(s["cols"].values())] +
                              [f"{s['sheet']}!{r}" for s in (hw, sw) for r in s["header_cells"].values()],
           "human_required": ["Open the file in Microsoft Excel (macros stay disabled) – no repair prompt", "Import it in eMASS (Hardware/Software → Import)",
                              "Spot-check three hardware and three software rows in eMASS"]}
    if not rep["ok"]:
        raise ValueError(f"Export validation failed: {(v1['issues'] + v2['issues'] + rep['outside_permitted'])[:3]}")
    ver = (db.scalar(select(func.max(InventoryExport.version)).where(InventoryExport.system_id == system.id)) or 0) + 1
    name = re.sub(r"\.xls[xm]$", "", imp.file_name, flags=re.I) + f"_export_v{ver}" + (".xlsm" if imp.file_name.lower().endswith(".xlsm") else ".xlsx")
    key = f"inventory/{system.org_id}/{system.id}/exports/v{ver}/{name}"
    store.put(key, out, XLSM)
    e = InventoryExport(system_id=system.id, import_id=imp.id, version=ver, file_name=name, sha256=hashlib.sha256(out).hexdigest(), storage_key=key,
                        exported_by=actor.email, counts={"hardware": len(assets), "software": len(softs)}, verification=rep)
    db.add(e)
    db.flush()
    return e
