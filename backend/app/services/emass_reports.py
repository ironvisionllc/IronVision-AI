"""eMASS report import – shared detection, context prompt and per-type handlers.
Every type: "eMASS <type> import – compatibility pending verification against a real eMASS export".
Types: system_details, control_information (services.emass_import), test_results (existing TR parser), poam, hwsw (existing HW/SW import)."""
import csv
import io
import re
from datetime import date

from fastapi import HTTPException
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models import ControlAssignment, EmassImport, Poam, System
from app.services import emass_import as ei
from app.services import emass_migration as em

TYPES = {"system_details": "System Details", "control_information": "Control Information", "test_results": "Test Results", "poam": "POA&M",
         "hwsw": "Hardware/Software List", "slcm": "SLCM Worksheet", "categorization": "Categorization", "ppsm": "PPSM Boundaries"}
MIGRATION_TYPES = ["test_results", "slcm", "hwsw", "categorization", "ppsm", "system_details"]  # Import Center / migration wizard; Control Information + POA&M deferred
SECTION_PERM = {"System details": "system.write", "Categorization": "system.write", "Roles": "system.write", "Controls": "implementation.write",
                "Test Results": "ingest.upload", "POA&M": "poam.manage", "Inventory": "system.write", "SLCM": "implementation.write",
                "PPSM": "system.write"}


def label(t: str) -> str:
    return f"eMASS {TYPES.get(t, t)} import – compatibility pending verification against a real eMASS export"


_n = ei._n
SYS_FIELDS = {  # field -> (label, aliases, section)
    "name": ("System Name", ["system name", "name", "information system name"], "System details"),
    "acronym": ("System Acronym", ["system acronym", "acronym"], "System details"),
    "emass_id": ("eMASS System ID", ["system id", "emass system id", "emass id", "systemid"], "System details"),
    "description": ("System Description", ["system description", "description", "purpose"], "System details"),
    "organization_name": ("Organization / Component", ["organization", "component", "owning organization", "organization name", "dod component"], "System details"),
    "system_type": ("System Type", ["system type", "registration type", "type"], "System details"),
    "mission_owner": ("Mission Owner", ["mission owner", "program office"], "System details"),
    "authorization_status": ("Authorization Status", ["authorization status", "ato status"], "System details"),
    "authorization_date": ("Authorization Date", ["authorization date", "ato date", "authorization termination date start"], "System details"),
    "authorization_expiration": ("Authorization Termination Date", ["authorization termination date", "ato expiration date", "ato expiration", "authorization expiration"], "System details"),
    "nss_determination": ("National Security System", ["national security system", "nss", "is nss"], "System details"),
    "confidentiality_impact": ("Confidentiality", ["confidentiality", "confidentiality level", "confidentiality impact"], "Categorization"),
    "integrity_impact": ("Integrity", ["integrity", "integrity level", "integrity impact"], "Categorization"),
    "availability_impact": ("Availability", ["availability", "availability level", "availability impact"], "Categorization"),
    "categorization": ("Security Categorization", ["security categorization", "categorization", "overall categorization"], "Categorization"),
    "role_isso": ("ISSO", ["isso", "information system security officer"], "Roles"),
    "role_issm": ("ISSM", ["issm", "information system security manager"], "Roles"),
    "role_system_owner": ("System Owner", ["system owner", "pm/iso", "program manager"], "Roles"),
    "role_ao": ("Authorizing Official", ["authorizing official", "ao"], "Roles"),
}
APPLY_SYS = {"name", "acronym", "emass_id", "description", "organization_name", "system_type", "mission_owner"}
BLOCK_SYS = {"authorization_status": "Set only by the AO authorization decision workflow (Authorize step)",
             "authorization_date": "Set only by the AO authorization decision workflow (Authorize step)",
             "authorization_expiration": "Set only by the AO authorization decision workflow (Authorize step)",
             "nss_determination": "Changes the categorization rules – update it in Prepare / Categorize",
             "Categorization": "Categorization changes go through the Categorize step and AO approval – shown as a proposal only",
             "Roles": "Assign people in Key roles – an import never creates user accounts or role assignments"}
POAM_FIELDS = {
    "emass_remote_id": ("POA&M ID", ["poam id", "poa&m id", "poa&m item id", "poam item id", "displaypoamid", "item id"]),
    "control": ("Control Acronym", ["control acronym", "security control number", "control", "security control"]),
    "cci": ("CCI", ["cci", "ccis"]),
    "weakness": ("Weakness Description", ["weakness description", "vulnerability description", "weakness", "description"]),
    "source_label": ("Source Identifying Weakness", ["source identifying weakness", "source identifying vulnerability", "source of discovery"]),
    "severity": ("Raw Severity", ["raw severity", "severity", "raw severity value"]),
    "residual_risk": ("Residual Risk Level", ["residual risk level", "residual risk"]),
    "status": ("Status", ["status", "poa&m status", "poam status"]),
    "due_date": ("Scheduled Completion Date", ["scheduled completion date", "scheduled completion", "completion date"]),
    "remediation_plan": ("Milestones", ["milestone with completion dates", "milestones with completion dates", "milestones", "milestone"]),
    "poc": ("Point of Contact", ["poc", "point of contact", "poc name"]),
    "resources": ("Resources Required", ["resources required", "resources"]),
    "comments": ("Comments", ["comments", "comment"]),
    "mitigations": ("Mitigations", ["mitigations", "mitigation"]),
}
SEV = {"veryhigh": "critical", "critical": "critical", "high": "high", "cati": "high", "moderate": "medium", "medium": "medium", "catii": "medium",
       "low": "low", "verylow": "low", "catiii": "low"}
RAW = {"verylow": "Very Low", "low": "Low", "moderate": "Moderate", "medium": "Moderate", "high": "High", "veryhigh": "Very High"}
CATS = {"cati": "CAT I", "catii": "CAT II", "catiii": "CAT III"}
SIGS = {
    "system_details": (["system name", "system acronym", "system id", "authorization status", "ato date", "confidentiality", "integrity", "availability",
                        "system type", "registration type", "national security system", "authorization termination date", "system description"], r"system|registration"),
    "control_information": (["control acronym", "implementation status", "security control designation", "slcm criticality", "slcm frequency",
                             "residual risk level", "relevance of threat", "common control provider", "n/a justification", "estimated completion date", "severity"], r"control"),
    "test_results": (["ap acronym", "cci", "cci definition", "compliance status", "date tested", "tested by", "test results", "assessment procedure"], r"test|template"),
    "poam": (["poam id", "poa&m id", "weakness description", "vulnerability description", "scheduled completion date", "milestone with completion dates",
              "source identifying weakness", "raw severity", "resources required", "point of contact", "severity", "status"], r"poa&?m"),
    "hwsw": (["machine name", "asset name", "manufacturer", "model number", "serial number", "software vendor", "software name", "software version",
              "ip address", "operating system"], r"hardware|software"),
    "slcm": (["control acronym", "control title", "criticality", "frequency", "method", "reporting", "tracking", "slcm comments", "order #", "sort #",
              "security control #"], r"final results|security controls|continuous monitoring"),
    "categorization": (["mission type", "applicability", "security category", "package", "availablilty", "confidentiality", "integrity"], r"mission types|categoriz"),
    "ppsm": (["port (required)", "protocol (required)", "data service (required)", "boundary (required)", "fqdn", "purpose", "ditpr id:", "poc name:"], r"ppsm"),
}


def read_sheets(data: bytes, file_name: str) -> list[tuple[str, list[tuple]]]:
    """All sheets as value tuples (macros never loaded; CSV accepted); size/row limits enforced."""
    if file_name.lower().endswith(".csv"):
        if len(data) > ei.MAX_BYTES:
            raise HTTPException(413, "File is larger than 10 MB")
        try:
            rows = [tuple(r) for r in csv.reader(io.StringIO(data.decode("utf-8-sig")))]
        except (UnicodeDecodeError, csv.Error) as exc:
            raise HTTPException(422, "Could not read the CSV file (UTF-8 expected)") from exc
        if len(rows) > ei.MAX_ROWS + ei.HEADER_SCAN_ROWS:
            raise HTTPException(422, f"More than {ei.MAX_ROWS:,} data rows – rejected")
        return [("CSV", [tuple(v if v != "" else None for v in r) for r in rows])]
    if not file_name.lower().endswith((".xlsx", ".xlsm")):
        raise HTTPException(422, "Upload an eMASS export as .xlsx, .xlsm or .csv")
    wb = ei._safe_open(data)
    try:
        if len(wb.sheetnames) > ei.MAX_SHEETS:
            raise HTTPException(422, "Too many worksheets")
        out = []
        for ws in wb.worksheets:
            rows = []
            for r in ws.iter_rows(values_only=True):
                rows.append(r)
                if len(rows) > ei.MAX_ROWS + ei.HEADER_SCAN_ROWS:
                    raise HTTPException(422, f"More than {ei.MAX_ROWS:,} data rows – rejected")
            out.append((ws.title, rows))
        return out
    finally:
        wb.close()


def _header(rows: list[tuple], aliases: dict[str, str], need: str | None = None) -> tuple[int, dict[str, int]] | None:
    best = None
    for i, row in enumerate(rows[:ei.HEADER_SCAN_ROWS]):
        hits: dict[str, int] = {}
        for c, v in enumerate(row):
            k = aliases.get(_n(v)) if v is not None else None
            if k and k not in hits:
                hits[k] = c
        if (need is None or need in hits) and len(hits) >= 2 and (best is None or len(hits) > len(best[1])):
            best = (i, hits)
    return best


def detect(sheets: list[tuple[str, list[tuple]]]) -> dict:
    scores = {}
    for t, (heads, sheet_rx) in SIGS.items():
        hs = {_n(h) for h in heads}
        best = 0
        for title, rows in sheets:
            seen = {_n(v) for row in rows[:ei.HEADER_SCAN_ROWS] for v in row if v is not None} & hs
            best = max(best, len(seen) + (2 if re.search(sheet_rx, title, re.I) else 0))
        scores[t] = best
    if sum(1 for title, _ in sheets if re.search(r"hardware|software", title, re.I)) >= 2:
        scores["hwsw"] += 3
    ranked = sorted(scores.items(), key=lambda x: -x[1])
    (top, s1), (_, s2) = ranked[0], ranked[1]
    conf = "high" if s1 >= 5 and s1 - s2 >= 3 else "medium" if s1 >= 3 and s1 - s2 >= 1 else "low"
    typ = top if s1 >= 2 else None
    return {"type": typ, "confidence": conf if s1 >= 2 else "none", "scores": dict(ranked),
            "candidates": [t for t, sc in ranked if sc >= max(2, s1 - 1)], "ambiguous": s1 - s2 <= 1,
            "sheets": [t for t, _ in sheets], "record_count": record_count(sheets, typ) if typ else None}


def record_count(sheets, typ: str) -> int | None:
    """Data rows below the best-matching header row of the detected type (Control Information: parsed control rows)."""
    if typ == "control_information":
        try:
            return len(ei.parse_sheets(sheets)["rows"])
        except HTTPException:
            return None
    sig = {_n(h) for h in SIGS.get(typ, ([], ""))[0]}
    best = None
    for _, rows in sheets:
        for i, r in enumerate(rows[:ei.HEADER_SCAN_ROWS]):
            hits = sum(1 for v in r if v is not None and _n(v) in sig)
            if hits >= 2 and (best is None or hits > best[0]):
                best = (hits, sum(1 for x in rows[i + 1:] if any(v not in (None, "") for v in x)))
    return best[1] if best else None


def route_type(det: dict, override: str | None = None) -> str:
    """ONE dispatcher for every upload path: explicit choice → detected type → Control Information (its parser explains what is missing)."""
    return override or det.get("type") or "control_information"


def identity(sheets) -> dict:
    """System name / acronym / eMASS ID found in the file (header+row below, or 'Label | Value' pairs, or 'System Name: X' banners)."""
    keys = {_n(a): f for f in ("name", "acronym", "emass_id") for a in SYS_FIELDS[f][1] if a not in ("name", "acronym", "description", "type")}
    keys.update({_n("system acronym"): "acronym", _n("system name"): "name"})
    out: dict = {}
    for _, rows in sheets:
        for i, row in enumerate(rows[:ei.HEADER_SCAN_ROWS]):
            for c, v in enumerate(row):
                if v is None:
                    continue
                m = re.match(r"\s*(system name|system acronym|system id|emass system id)\s*[:=]\s*(.+)", str(v), re.I)
                if m:
                    out.setdefault(keys[_n(m.group(1))], m.group(2).strip())
                    continue
                f = keys.get(_n(v))
                if not f or f in out:
                    continue
                right = row[c + 1] if c + 1 < len(row) else None
                below = rows[i + 1][c] if i + 1 < len(rows) and c < len(rows[i + 1]) else None
                all_labels = {_n(a) for _, al, _ in SYS_FIELDS.values() for a in al}
                header_like = sum(1 for x in row if x not in (None, "")) >= 3
                for val in ((below, right) if header_like else (right, below)):
                    if val not in (None, "") and _n(val) not in all_labels and not str(val).strip().endswith(":"):
                        out[f] = str(val).strip()
                        break
    return out


def mismatch(ident: dict, s: System | None) -> str | None:
    if s is None or not ident:
        return None
    pairs = [("eMASS ID", ident.get("emass_id"), s.emass_id), ("acronym", ident.get("acronym"), s.acronym), ("name", ident.get("name"), s.name)]
    known = [(k, a, b) for k, a, b in pairs if a and b]
    if ident.get("emass_id") and s.emass_id and _n(ident["emass_id"]) != _n(s.emass_id):
        return (f"The file identifies eMASS ID '{ident['emass_id']}' – this does not match the current system {s.acronym or s.name} "
                f"(eMASS ID '{s.emass_id}'), even if the name or acronym is the same. Check the file before continuing")
    if any(_n(a) == _n(b) for _, a, b in known):
        return None
    return "The file identifies " + ", ".join(f"{k} '{a}'" for k, a, _ in known or [(k, a, b) for k, a, b in pairs if a]) + \
        f" – this does not match the current system {s.acronym or s.name}" if (known or any(a for _, a, _ in pairs)) else None


def _field(cur, imp_v, err=None, blocked=None, lab=None) -> dict:
    return {"current": ei._j(cur), "imported": ei._j(imp_v), "change": not err and str(cur if cur is not None else "") != str(imp_v),
            "error": err, "blocked": blocked, "decision": None, "label": lab}


def build_system(db: Session, s: System, sheets) -> tuple[list[dict], dict]:
    aliases = {_n(a): f for f, (_, al, _) in SYS_FIELDS.items() for a in al}
    vals: dict = {}
    for _, rows in sheets:
        h = _header(rows, aliases)
        if h and h[0] + 1 < len(rows):
            i, cols = h
            vals.update({f: rows[i + 1][c] for f, c in cols.items() if c < len(rows[i + 1]) and rows[i + 1][c] not in (None, "") and f not in vals})
        for row in rows:  # 'Label | Value' layout
            if len(row) >= 2 and row[0] is not None and _n(row[0]) in aliases and row[1] not in (None, ""):
                vals.setdefault(aliases[_n(row[0])], row[1])
    secs: dict[str, dict] = {}
    for f, raw in vals.items():
        lab, _, sec = SYS_FIELDS[f]
        v, err = (ei._date(raw) if f in ("authorization_date", "authorization_expiration") else (str(raw).strip()[:2000], None))
        if f in ("confidentiality_impact", "integrity_impact", "availability_impact", "categorization") and not err:
            v = {"l": "Low", "low": "Low", "m": "Moderate", "moderate": "Moderate", "h": "High", "high": "High"}.get(_n(v), v)
        cur = None if f.startswith("role_") else getattr(s, f, None) if s else None
        blocked = None if f in APPLY_SYS else BLOCK_SYS.get(f) or BLOCK_SYS.get(sec)
        secs.setdefault(sec, {})[f] = _field(cur, v, err, blocked if s else None, lab)
    items = [{"section": sec, "key": sec.lower().replace(" ", "_"), "label": sec, "create": False, "fields": fs, "required_issue": None} for sec, fs in secs.items()]
    return items, {"matched": len(items), "fields_read": sorted(vals)}


def build_poam(db: Session, s: System, sheets) -> tuple[list[dict], dict]:
    aliases = {_n(a): f for f, (_, al) in POAM_FIELDS.items() for a in al}
    found = None
    for title, rows in sheets:
        h = _header(rows, aliases, need="weakness")
        if h and (found is None or len(h[1]) > len(found[2])):
            found = (title, rows, h[1], h[0])
    if not found:
        raise HTTPException(422, "No POA&M header row (Weakness / Vulnerability Description) found")
    _, rows, cols, hi = found
    assigned = set(db.scalars(select(ControlAssignment.control_id).where(ControlAssignment.system_id == s.id)))
    existing = list(db.scalars(select(Poam).where(Poam.system_id == s.id)))
    items, rep = [], {"invalid": [], "not_in_system": [], "matched_existing": 0, "new": 0}
    for n, row in enumerate(rows[hi + 1:], start=hi + 2):
        r = {f: (row[c] if c < len(row) else None) for f, c in cols.items()}
        if not any(v not in (None, "") for v in r.values()):
            continue
        cids = [c for c in (ei.norm_acronym(x) for x in re.split(r"[,;\n]+", str(r.get("control") or ""))) if c]
        weak = str(r.get("weakness") or "").strip()
        p = next((x for x in existing if r.get("emass_remote_id") and x.emass_remote_id and str(x.emass_remote_id) == str(r["emass_remote_id"]).strip()), None) or \
            next((x for x in existing if weak and _n(x.weakness) == _n(weak) and set(x.control_ids or []) & set(cids)), None)
        fs = {}
        for f, raw in r.items():
            if raw in (None, "") or f == "cci":
                continue
            lab = POAM_FIELDS[f][0]
            v, err, blocked = str(raw).strip(), None, None
            if f == "control":
                v = [ei.fmt_acronym(c) for c in sorted(cids)]
                err = None if cids else f"'{raw}' is not a Control Acronym"
                if cids and (bad := [c for c in cids if c not in assigned]):
                    err = f"Not assigned to this system: {', '.join(ei.fmt_acronym(c) for c in bad)}"
                    rep["not_in_system"].append({"row": n, "control": ", ".join(ei.fmt_acronym(c) for c in bad)})
            elif f == "due_date":
                v, err = ei._date(raw)
            elif f == "severity":  # eMASS Raw Severity stays on the eMASS scale; CAT I–III stays a CAT – never converted to technical severity
                v = RAW.get(_n(raw)) or CATS.get(_n(raw))
                err = None if v else f"'{raw}' is not a Raw Severity (Very Low … Very High) or CAT I–III"
            elif f == "residual_risk":
                v = RAW.get(_n(raw))
                err = None if v else f"'{raw}' is not a Residual Risk Level (Very Low … Very High)"
            elif f == "status":
                blocked = ("Closing or risk-accepting a POA&M goes through the closure / risk-acceptance workflow"
                           if _n(raw) in ("completed", "closed", "riskaccepted") else None)
                if not blocked:
                    continue
            elif f == "poc":
                blocked = "The POA&M owner is assigned in the app – the import never creates or maps user accounts"
            elif f in ("source_label",):
                v = v[:50]
            cur = None
            if p is not None:
                cur = {"control": [ei.fmt_acronym(c) for c in sorted(p.control_ids or [])], "residual_risk": p.residual_risk_level,
                       "severity": p.raw_severity if not str(v).startswith("CAT") else p.cat}.get(f, getattr(p, f, None)) if f != "poc" else None
            fs[f] = _field(cur, v, err, blocked, lab)
            if err:
                rep["invalid"].append({"row": n, "control": ", ".join(ei.fmt_acronym(c) for c in cids), "field": lab, "value": str(raw)[:100], "error": err})
        if p is None and ("weakness" not in fs or "control" not in fs or fs["control"]["error"]):
            issue = "A new POA&M needs a weakness and controls assigned to this system"
        else:
            issue = None
        rep["matched_existing" if p else "new"] += 1
        note = ("eMASS marks this item Completed/Closed/Risk Accepted – creating it here starts a new open POA&M; closure still needs the ISSM workflow"
                if p is None and _n(r.get("status") or "") in ("completed", "closed", "riskaccepted") else None)
        items.append({"section": "POA&M", "key": p.human_id if p else f"new-row-{n}", "label": f"{p.human_id} (eMASS {p.emass_remote_id or '—'})" if p else
                      f"New POA&M from row {n}", "create": p is None, "poam_id": str(p.id) if p else None, "row": n, "fields": fs, "required_issue": issue, "note": note})
    rep["matched"] = len(items)
    return items, rep


def build_file_item(section: str, typ: str, data: bytes, file_name: str) -> tuple[list[dict], dict]:
    """TR and HW/SW reuse the existing importers: the session proposes creating an import that then follows the existing rules."""
    if typ == "test_results":
        from app.services.tr import xlsx
        try:
            det = xlsx.detect(data)
        except Exception as exc:  # noqa: BLE001
            raise HTTPException(422, f"Not a Test Results workbook: {exc}") from exc
        what = f"Test Results workbook – sheet '{det.get('sheet')}', header row {det.get('header_row')}"
        note = "Imported Compliance Status / Tested By / Date Tested stay PROPOSALS – the SCA approves each row (separation of duties); nothing is approved automatically."
    else:
        from app.services.inventory import hwsw
        det = hwsw.detect(data)
        if not det.get("hardware") and not det.get("software"):
            raise HTTPException(422, "Not a Hardware/Software List (no Hardware or Software header row found)")
        what = f"HW/SW import – {len((det.get('hardware') or {}).get('rows', []))} hardware, {len((det.get('software') or {}).get('rows', []))} software rows"
        note = "Processed by the existing HW/SW import: macros are never run" + (" (this file contains a VBA project – ignored)" if det.get("has_vba") else "") + \
            "; changes stay proposals until a second person approves them (separation of duties)."
    if file_name.lower().endswith(".csv"):
        raise HTTPException(422, f"{TYPES[typ]} imports need the eMASS .xlsx/.xlsm template")
    return [{"section": section, "key": typ, "label": what, "create": True, "note": note,
             "fields": {"import": _field(None, what, lab="Will be created")}, "required_issue": None}], {"matched": 1}


def build(db: Session, s: System, typ: str, data: bytes, file_name: str) -> tuple[list[dict], dict]:
    sheets = read_sheets(data, file_name)
    if typ == "control_information":
        items, rep = ei.build(db, s, ei.parse_sheets(sheets))
        for i in items:
            i.update(section="Controls", key=i["control_id"], label=i["control"], create=False)
        return items, rep
    if typ == "system_details":
        return build_system(db, s, sheets)
    if typ == "poam":
        return build_poam(db, s, sheets)
    if typ in em.TYPES:
        return em.build(db, s, typ, sheets)
    return build_file_item("Test Results" if typ == "test_results" else "Inventory", typ, data, file_name)


def apply_other(db: Session, imp: EmassImport, s: System, current, meta: dict, src: dict) -> list[dict]:
    from app.services import audit
    from app.services import workflow as wf
    from app.services.storage.base import get_storage
    results = []
    if imp.report_type == "slcm":
        return em.apply_slcm_all(db, imp, s, current, meta, src)
    for i in imp.items or []:
        acc = {k: f for k, f in i["fields"].items() if ei._ok(f) and f["decision"] == "accept"}
        if not acc or i["section"] == "Controls":
            continue
        if imp.report_type in em.TYPES:
            results.append(em.apply_item(db, imp, s, i, acc, current, meta, src))
        elif i["section"] == "System details":
            before = {k: ei._j(getattr(s, k)) for k in acc}
            for k, f in acc.items():
                setattr(s, k, f["imported"])
            audit.record(db, action="system.updated", actor_user_id=current.user.id, actor_email=current.user.email, entity_type="system", entity_id=str(s.id),
                         before=before, after={k: f["imported"] for k, f in acc.items()}, details=src, **meta)
            results.append({"control": "System details", "applied": sorted(acc), "failed": None})
        elif i["section"] == "POA&M":
            vals = {k: f["imported"] for k, f in acc.items()}
            if "severity" in vals:
                sv = vals.pop("severity")
                vals["cat" if sv.startswith("CAT") else "raw_severity"] = sv
            if "residual_risk" in vals:
                vals["residual_risk_level"] = vals.pop("residual_risk")
            if "control" in vals:
                vals["control_ids"] = [ei.norm_acronym(c) for c in vals.pop("control")]
            if "due_date" in vals:
                vals["due_date"] = date.fromisoformat(vals["due_date"])
            if risk := [k for k in ("raw_severity", "residual_risk_level", "mitigations", "comments") if k in vals]:
                vals["risk_source"] = {k: f"emass_import:{imp.id}" for k in risk}
            if i["create"]:
                if i["required_issue"] or not {"weakness", "control_ids"} <= set(vals):
                    results.append({"control": i["label"], "applied": [], "failed": i["required_issue"] or "Accept the weakness and controls to create a POA&M"})
                    continue
                p = wf.create_poam(db, system=s, actor=current.user, data={**vals, "status": "new"}, findings=[])
                audit.record(db, action="poam.created", actor_user_id=current.user.id, actor_email=current.user.email, entity_type="poam", entity_id=str(p.id),
                             after={k: ei._j(v) for k, v in vals.items()}, details=src, **meta)
            else:
                p = db.get(Poam, i["poam_id"])
                before = {k: ei._j(getattr(p, k)) for k in vals}
                for k, v in vals.items():
                    setattr(p, k, v)
                wf.history(db, p, current.user, "updated", status=p.status)
                audit.record(db, action="poam.updated", actor_user_id=current.user.id, actor_email=current.user.email, entity_type="poam", entity_id=str(p.id),
                             before=before, after={k: ei._j(v) for k, v in vals.items()}, details=src, **meta)
            results.append({"control": i["label"], "applied": sorted(acc), "failed": None})
        elif i["section"] in ("Test Results", "Inventory"):
            data = get_storage().get(imp.storage_key)
            if i["section"] == "Test Results":
                from app.services.tr import service as trs
                obj = trs.import_workbook(db, s, imp.file_name, data, current.user)
                act = "tr.workbook_uploaded"
            else:
                from app.services.inventory import hwsw
                obj = hwsw.upload(db, s, imp.file_name, data, current.user)
                act = "inventory.import_uploaded"
            audit.record(db, action=act, actor_user_id=current.user.id, actor_email=current.user.email, entity_type=act.split(".")[0],
                         entity_id=str(obj.id), after={"file_name": imp.file_name, "sha256": imp.sha256}, details=src, **meta)
            results.append({"control": i["section"], "applied": ["import"], "failed": None, "created_id": str(obj.id)})
    return results


def prefill(imp: EmassImport, sheets) -> dict:
    """Values for the existing registration form – nothing is created."""
    if imp.report_type == "categorization" or (imp.detected or {}).get("type") == "categorization":
        p = em.parse_categorization(sheets)
        hint = {k: v for k, v in zip("cia", p["impacts"]) if v}
        return {"system": {"name": p["meta"]["package"]} if p["meta"].get("package") else {}, "categorization_hint": hint or None,
                "mission_types": [x["mission_type"] for x in p["selected"]],
                "note": "Pre-filled from the eMASS Categorization template (high-water mark of the applicable mission types). Registration validation and AO approval still apply."}
    items, _ = build_system(None, None, sheets)  # type: ignore[arg-type]
    vals = {k: f["imported"] for i in items for k, f in i["fields"].items() if not f["error"]}
    out = {k: vals[k] for k in ("name", "acronym", "emass_id", "description", "organization_name", "mission_owner") if vals.get(k)}
    ident = (imp.detected or {}).get("identity") or {}
    for k in ("name", "acronym", "emass_id"):
        if ident.get(k) and k not in out:
            out[k] = ident[k]
    cat = {k[0]: vals[f"{k}_impact"] for k in ("confidentiality", "integrity", "availability") if vals.get(f"{k}_impact")}
    return {"system": out, "categorization_hint": cat or None,
            "note": "Pre-filled from the eMASS file. Registration validation, role rules and AO approval of the categorization still apply."}


SAMPLES = {
    "system_details": [["SYNTHETIC SAMPLE – eMASS System Details export layout (compatibility pending verification). Not real eMASS data."], [],
                       ["System Name", "System Acronym", "System ID", "System Description", "Organization", "System Type", "Authorization Status",
                        "ATO Date", "Authorization Termination Date", "National Security System", "Confidentiality", "Integrity", "Availability", "ISSO", "Authorizing Official"],
                       ["Logistics Tracking Platform", "LTP", "SYN-48213", "SYNTHETIC – tracks shipments and depot inventory (updated description from eMASS)",
                        "SYNTHETIC Logistics Command", "Major Application", "Authorization to Operate (ATO)", "2025-03-01", "2028-03-01", "No", "Moderate", "Moderate", "Low",
                        "Alex Rivera", "Dana Brooks"]],
    "poam": [["SYNTHETIC SAMPLE – eMASS POA&M export layout (compatibility pending verification). Not real eMASS data."], [],
             ["POA&M ID", "Control Acronym", "CCI", "Weakness Description", "Source Identifying Weakness", "Raw Severity", "Residual Risk Level", "Status",
              "Scheduled Completion Date", "Milestone with Completion Dates", "POC", "Resources Required", "Comments"],
             ["SYN-9001", "AC-2", "000015", "SYNTHETIC: Inactive accounts are not disabled after 35 days", "eMASS assessment", "Moderate", "Low", "Ongoing",
              "06/30/2027", "1) Configure automatic disable – 03/31/2027; 2) Verify – 06/30/2027", "Alex Rivera", "2 admin hours/week", "SYNTHETIC sample"],
             ["SYN-9002", "SI-2", "", "SYNTHETIC: Patch cadence exceeds 30 days for edge servers", "ACAS", "High", "Moderate", "Completed", "01/15/2027", "Patch and rescan", "", "", ""],
             ["SYN-9003", "ZZ-9", "", "SYNTHETIC: control not in the system", "Manual", "Extreme", "", "Ongoing", "13/45/2027", "", "", "", ""]],
}


def sample_xlsx(typ: str, system_rows: list[dict] | None = None) -> bytes:
    if typ == "control_information":
        return ei.sample_xlsx(system_rows)
    if typ in em.TYPES:
        return em.sample(typ)
    if typ in ("test_results", "hwsw"):
        from pathlib import Path

        from app.services.tr import xlsx
        base = Path(__file__).resolve().parents[2] / "tests" / "fixtures" / "emass"
        if typ == "test_results":  # real blank template populated with SYNTHETIC rows (blank-template export of a sandbox system)
            return (base / "SYNTHETIC_TRExport.xlsx").read_bytes()
        tpl = (base / "HWSWList_Template-eMASS3.xlsm").read_bytes()
        hw = {"B8": "Server", "C8": "syn-web01", "E8": "10.10.0.11", "G8": "SYNTHETIC Dell", "H8": "R750", "I8": "SYN-SN-001", "J8": "RHEL 9",
              "B9": "Server", "C9": "syn-db01", "E9": "10.10.0.12", "I9": "SYN-SN-002", "J9": "RHEL 9"}
        sw = {"B8": "COTS Application", "C8": "SYNTHETIC Apache", "D8": "HTTP Server", "E8": "2.4.58"}
        return xlsx.write(xlsx.write(tpl, "Hardware", hw), "Software", sw)
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "System Details" if typ == "system_details" else "POA&M"
    for r in SAMPLES[typ]:
        ws.append(r)
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()
