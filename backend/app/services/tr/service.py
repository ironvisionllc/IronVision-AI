"""eMASS Test Result workflow: import → detect/match → generate proposals from real data → SCA review → export into a copy of
the original workbook. Evidence only ever SUGGESTS a Compliance Status; the final value comes from an assessor determination
(Assessment record or an SCA decision in review) and Tested By / Date Tested come from that same determination."""
import hashlib
import re
import uuid
from datetime import date, datetime, timezone

from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.models import (Assessment, Cci, CciControlMap, Control, ControlAssignment, Evidence, EvidenceControl, Finding, FindingControlMap, Poam,
                        PolicyRequirementCoverage, System, TrExport, TrRow, TrTemplate, TrWorkbook, User)
from app.services.tr import xlsx

AP_RX = re.compile(r"^\s*([A-Za-z]{2})-(\d+)\s*(?:\((\d+)\))?(?:\.(\d+))?\s*$")
RESULT_MAP = {"satisfied": "Compliant", "other_than_satisfied": "Non-Compliant", "not_satisfied": "Non-Compliant", "not_applicable": "Not Applicable"}
STIG_STATE = {"open": "Open", "pass": "NotAFinding", "not_applicable": "Not_Applicable", "closed": "NotAFinding"}


def now():
    return datetime.now(timezone.utc)


def parse_ap(ap: str | None) -> tuple[str | None, str | None]:
    """'AC-1.1' → ('ac-1', '1'); 'AC-2(1).3' → ('ac-2.1', '3'); 'AC-2(1)' → ('ac-2.1', None)."""
    m = AP_RX.match(ap or "")
    if not m:
        return None, None
    fam, num, enh, proc = m.groups()
    return f"{fam.lower()}-{int(num)}" + (f".{int(enh)}" if enh else ""), proc


def norm_cci(v: str | None) -> str | None:
    d = re.sub(r"\D", "", str(v or ""))
    return f"CCI-{int(d):06d}" if d else None


def _ref(db: Session) -> tuple[dict, set]:
    from app.services.content.loader import active_catalog_id

    cmap: dict = {}
    for cci, ctl in db.execute(select(CciControlMap.cci_id, CciControlMap.control_id)):
        cmap.setdefault(cci, set()).add(ctl)
    cat = active_catalog_id(db)
    return cmap, set(db.scalars(select(Control.control_id).where(Control.catalog_version_id == cat)))


def match(cmap: dict, catalog: set, ap: str | None, cci_raw: str | None, ctl_raw: str | None) -> dict:
    ctl, _ = parse_ap(ap)
    ctl = ctl or parse_ap(ctl_raw)[0]
    cci = norm_cci(cci_raw)
    out = {"cci_id": cci, "control_id": None}
    if cci and cci not in cmap:
        return {**out, "match_status": "unmatched", "match_reason": f"{cci} is not in the loaded DISA CCI list"}
    if ap and not parse_ap(ap)[0]:
        if cci and len(cmap[cci]) == 1:
            return {**out, "control_id": next(iter(cmap[cci])), "match_status": "ambiguous",
                    "match_reason": f"AP Acronym “{ap}” is not recognised; {cci} maps to {next(iter(cmap[cci])).upper()} – confirm"}
        return {**out, "match_status": "unmatched", "match_reason": f"AP Acronym “{ap}” is not recognised"}
    if ctl and ctl not in catalog:
        return {**out, "match_status": "unmatched", "match_reason": f"{ctl.upper()} is not in the control catalog"}
    if ctl and cci:
        if ctl in cmap[cci]:
            return {**out, "control_id": ctl, "match_status": "matched", "match_reason": f"{(ap or ctl).upper()} and {cci} agree"}
        if ctl.split(".")[0] in {c.split(".")[0] for c in cmap[cci]}:
            return {**out, "control_id": ctl, "match_status": "ambiguous",
                    "match_reason": f"{cci} maps to {', '.join(sorted(c.upper() for c in cmap[cci]))}, the row says {ctl.upper()} – confirm the enhancement"}
        return {**out, "control_id": ctl, "match_status": "ambiguous",
                "match_reason": f"{cci} maps to {', '.join(sorted(c.upper() for c in cmap[cci]))}, not {ctl.upper()}"}
    if cci:
        cands = sorted(cmap[cci])
        if len(cands) == 1:
            return {**out, "control_id": cands[0], "match_status": "ambiguous", "match_reason": f"No AP Acronym; {cci} maps to {cands[0].upper()} – confirm"}
        return {**out, "match_status": "ambiguous", "match_reason": f"No AP Acronym; {cci} maps to several controls: {', '.join(c.upper() for c in cands)}"}
    if ctl:
        return {**out, "control_id": ctl, "match_status": "ambiguous", "match_reason": f"No CCI – {ctl.upper()} row needs a CCI or confirmation"}
    return {**out, "match_status": "unmatched", "match_reason": "Row has neither a recognisable AP Acronym nor a CCI"}


def import_workbook(db: Session, system: System, file_name: str, data: bytes, actor) -> TrWorkbook:
    from app.services import jobs
    from app.services.storage.base import get_storage

    sha = hashlib.sha256(data).hexdigest()
    ver = (db.scalar(select(func.max(TrWorkbook.version)).where(TrWorkbook.system_id == system.id, TrWorkbook.file_name == file_name)) or 0) + 1
    key = f"tr/{system.org_id}/{system.id}/originals/{sha}/v{ver}/{file_name}"
    get_storage().put(key, data, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    wb = TrWorkbook(org_id=system.org_id, system_id=system.id, file_name=file_name, sha256=sha, storage_key=key, version=ver, status="queued",
                    uploaded_by=actor.email)
    db.add(wb)
    db.flush()
    jobs.enqueue(db, "tr.detect", {"workbook_id": str(wb.id)}, requested_by=getattr(actor, "id", None))
    return wb


def _apply_template(db: Session, wb: TrWorkbook, det: dict) -> dict:
    t = db.scalar(select(TrTemplate).where(TrTemplate.org_id == wb.org_id, TrTemplate.fingerprint == det["fingerprint"]).order_by(TrTemplate.version.desc()))
    d = {k: det[k] for k in ("sheet", "header_row", "roles", "vocabulary", "date_format", "date_write", "limits", "readonly_cols")}
    if t:
        wb.template_id = t.id
        d = {**d, **{k: t.definition[k] for k in ("roles", "vocabulary") if k in t.definition}}
        d["template"] = {"id": str(t.id), "name": t.name, "version": t.version}
    return d


def detect_and_match(db: Session, wb: TrWorkbook, roles_override: dict | None = None) -> None:
    from app.services.storage.base import get_storage

    data = get_storage().get(wb.storage_key)
    det = xlsx.detect(data)
    wb.detection = {k: v for k, v in det.items() if k != "rows"}
    wb.definition = _apply_template(db, wb, det) if not roles_override else {**(wb.definition or {}), "roles": roles_override}
    if roles_override and roles_override != det["roles"]:  # re-read the identifiers with the user's column roles
        import io

        import openpyxl
        ws = openpyxl.load_workbook(io.BytesIO(data))[det["sheet"]]
        det["rows"] = [{"row": r, **{k: (str(ws[f"{c}{r}"].value).strip() if ws[f"{c}{r}"].value is not None else None) for k, c in roles_override.items()
                                     if k in xlsx.ID_ROLES}} for r in range(det["header_row"] + 1, ws.max_row + 1)
                       if any(ws[f"{roles_override[k]}{r}"].value not in (None, "") for k in ("ap_acronym", "cci", "control_acronym") if k in roles_override)]
    db.query(TrRow).filter(TrRow.workbook_id == wb.id).delete()
    cmap, catalog = _ref(db)
    seen: dict = {}
    for r in det["rows"]:
        m = match(cmap, catalog, r.get("ap_acronym"), r.get("cci"), r.get("control_acronym"))
        flags = []
        k = (r.get("ap_acronym"), m["cci_id"])
        if m["cci_id"] and k in seen:
            m = {**m, "match_status": "ambiguous", "match_reason": f"Duplicate of row {seen[k]} (same AP Acronym and CCI) – eMASS accepts one result per AP"}
            flags.append("duplicate")
        seen.setdefault(k, r["row"])
        db.add(TrRow(workbook_id=wb.id, row_number=r["row"], control_acronym=r.get("control_acronym"), ap_acronym=r.get("ap_acronym"),
                     cci_id=m["cci_id"], control_id=m["control_id"], match_status=m["match_status"], match_reason=m["match_reason"], flags=flags))
    db.flush()
    wb.status = "mapped"
    refresh_counts(db, wb)


def refresh_counts(db: Session, wb: TrWorkbook) -> dict:
    by = lambda col: dict(db.execute(select(col, func.count()).where(TrRow.workbook_id == wb.id).group_by(col)).all())  # noqa: E731
    wb.counts = {"rows": sum(by(TrRow.match_status).values()), "match": by(TrRow.match_status), "state": by(TrRow.state)}
    return wb.counts


def save_template(db: Session, wb: TrWorkbook, name: str, actor) -> TrTemplate:
    fp = wb.detection["fingerprint"]
    ver = (db.scalar(select(func.max(TrTemplate.version)).where(TrTemplate.org_id == wb.org_id, TrTemplate.fingerprint == fp)) or 0) + 1
    t = TrTemplate(org_id=wb.org_id, name=name, version=ver, fingerprint=fp, definition=wb.definition, created_by=actor.email)
    db.add(t)
    db.flush()
    wb.template_id = t.id
    wb.definition = {**wb.definition, "template": {"id": str(t.id), "name": name, "version": ver}}
    return t


def populate_from_system(db: Session, wb: TrWorkbook) -> int:
    """Blank template (no AP rows): add one row per CCI of each assigned control. AP numbers are DERIVED (control.n) – flagged."""
    if db.scalar(select(func.count()).select_from(TrRow).where(TrRow.workbook_id == wb.id)):
        return 0
    first = wb.definition["header_row"] + 1
    defs = dict(db.execute(select(Cci.cci_id, Cci.definition)).all())
    n = 0
    for cid in db.scalars(select(ControlAssignment.control_id).where(ControlAssignment.system_id == wb.system_id).order_by(ControlAssignment.control_id)):
        ccis = sorted(set(db.scalars(select(CciControlMap.cci_id).where(CciControlMap.control_id == cid))))
        for i, cci in enumerate(ccis, start=1):
            ctl = cid.upper().replace(".", "(") + (")" if "." in cid else "")
            db.add(TrRow(workbook_id=wb.id, row_number=first + n, added=True, control_acronym=ctl, ap_acronym=f"{ctl}.{i}", cci_id=cci, control_id=cid,
                         match_status="matched", match_reason="Row added from the system's control set (template had no AP rows)",
                         flags=["derived_ap"], proposed={"cci_definition": defs.get(cci)}))
            n += 1
    db.flush()
    return n


def _name(u: User | None) -> str | None:
    if u is None:
        return None
    parts = re.sub(r"\s*\([^)]*\)", "", u.display_name or "").split()  # "Dana Brooks (SCA persona)" → "Brooks, Dana"
    return f"{parts[-1]}, {' '.join(parts[:-1])}" if len(parts) > 1 else (u.display_name or u.email)


def generate(db: Session, wb: TrWorkbook, populate: bool = False) -> dict:
    from app.services.policy_coverage import applicable_policies

    if populate:
        populate_from_system(db, wb)
    sid = wb.system_id
    system = db.get(System, sid)
    assigned = {a.control_id: a for a in db.scalars(select(ControlAssignment).where(ControlAssignment.system_id == sid))}
    assess: dict = {}
    for a in db.scalars(select(Assessment).where(Assessment.system_id == sid).order_by(Assessment.created_at)):
        assess[a.control_id] = a
    stig: dict = {}
    for f in db.scalars(select(Finding).where(Finding.system_id == sid, Finding.source_type == "stig_checklist")):
        for c in f.ccis or []:
            stig.setdefault(c, []).append(f)
    acas: dict = {}
    for f, ctl in db.execute(select(Finding, FindingControlMap.control_id).join(FindingControlMap, FindingControlMap.finding_id == Finding.id)
                             .where(Finding.system_id == sid, Finding.source_type == "acas", Finding.current_state == "open",
                                    FindingControlMap.valid_to.is_(None), FindingControlMap.review_status != "rejected")):
        acas.setdefault(ctl, []).append(f)
    pols = {p.id: p for p in applicable_policies(db, system)}
    cov: dict = {}
    for c in db.scalars(select(PolicyRequirementCoverage).where(PolicyRequirementCoverage.policy_id.in_(list(pols)), PolicyRequirementCoverage.requirement_type == "cci")):
        cov.setdefault(c.requirement_id, []).append(c)
    evs: dict = {}
    for aid, e in db.execute(select(EvidenceControl.assignment_id, Evidence).join(Evidence, Evidence.id == EvidenceControl.evidence_id)
                             .where(EvidenceControl.assignment_id.in_([a.id for a in assigned.values()]))):
        evs.setdefault(aid, []).append(e)
    poams: dict = {}
    for p in db.scalars(select(Poam).where(Poam.system_id == sid)):
        for c in p.control_ids or []:
            poams.setdefault(c, []).append(p)
    limit = (wb.definition.get("limits") or {}).get("test_results") or 4000
    for row in db.scalars(select(TrRow).where(TrRow.workbook_id == wb.id)):
        if row.state == "approved":
            continue
        flags = [f for f in (row.flags or []) if f in ("duplicate", "derived_ap")]
        if row.match_status in ("ambiguous", "unmatched") or not row.control_id:
            row.proposed = {**{k: v for k, v in (row.proposed or {}).items() if k == "cci_definition"}, "compliance_status": None, "suggestion": None,
                            "date_tested": None, "tested_by": None, "test_results": None, "determination_source": None}
            row.sources, row.state = [], "missing"
            row.flags = flags + ["needs_mapping_review"]
            continue
        cid, cci = row.control_id, row.cci_id
        a = assigned.get(cid)
        src, parts = [], []
        det = tested_by = tested_on = None
        asm = assess.get(cid)
        if asm is not None:
            det = RESULT_MAP.get(asm.result)
            tested_by, tested_on = _name(db.get(User, asm.assessor_id)), asm.created_at.date().isoformat()
            src.append({"type": "assessment", "id": str(asm.id), "label": f"Assessment {asm.result.replace('_', ' ')} by {tested_by} ({tested_on})"})
            if asm.findings_text:
                parts.append(f"Assessor notes: {asm.findings_text.strip()[:600]}")
        st = stig.get(cci, []) if cci else []
        opn = [f for f in st if f.current_state == "open"]
        seen_l = set()
        for f in st:
            label = f"STIG {f.vuln_id or f.rule_id} {STIG_STATE.get(f.current_state, f.current_state)} ({f.benchmark or 'checklist'}, {f.last_seen_at:%Y-%m-%d})"
            if label in seen_l or len(seen_l) >= 8:
                continue
            seen_l.add(label)
            src.append({"type": "finding", "id": str(f.id), "label": label})
        if st:
            parts.append("; ".join(s["label"] for s in src if s["type"] == "finding") + (f" (+{len(st) - len(seen_l)} more on other assets)" if len(st) > len(seen_l) else ""))
        for f in acas.get(cid, [])[:5]:
            src.append({"type": "finding", "id": str(f.id), "label": f"ACAS plugin {f.plugin_id} open ({f.severity}): {f.title[:70]}"})
            parts.append(src[-1]["label"])
        for c in (cov.get(cci, []) if cci else []):
            p = pols[c.policy_id]
            loc = (c.snippets or [{}])[0].get("location", "")
            src.append({"type": "policy", "id": str(p.id), "label": f"Policy “{p.title}” v{p.version} {loc} – {c.status.upper()}"})
            parts.append(src[-1]["label"])
        stale = False
        for e in (evs.get(a.id, []) if a else [])[:5]:
            stale = stale or e.state not in ("current", "near_expiration")
            src.append({"type": "evidence", "id": str(e.id), "label": f"Evidence {e.file_name} ({e.collected_at:%Y-%m-%d}, SHA-256 {e.sha256[:12]}…){' – STALE' if e.state not in ('current', 'near_expiration') else ''}"})
        for p in poams.get(cid, [])[:3]:
            src.append({"type": "poam", "id": str(p.id), "label": f"POA&M {p.human_id} “{p.weakness[:60]}” ({p.status})"})
            parts.append(src[-1]["label"])
        sugg = "Non-Compliant" if opn or acas.get(cid) else ("Compliant" if any(f.current_state in ("pass", "not_applicable") for f in st)
                                                             or any(c.status == "met" for c in cov.get(cci, [])) else None)
        if stale:
            flags.append("stale_evidence")
        if det == "Compliant" and (opn or acas.get(cid)):
            flags.append("conflict")
        if det is None:
            flags.append("assessor_input_required")
        narrative = "; ".join(parts)[:limit] or None
        row.proposed = {**{k: v for k, v in (row.proposed or {}).items() if k == "cci_definition"}, "compliance_status": det, "suggestion": sugg,
                        "date_tested": tested_on, "tested_by": tested_by, "test_results": narrative,
                        "determination_source": f"assessment:{asm.id}" if det else None}
        row.sources, row.flags = src, flags
        row.state = "conflict" if "conflict" in flags else ("missing" if not src else ("edited" if row.state == "edited" else "auto"))
    wb.status, wb.generated_at = "generated", now()
    return refresh_counts(db, wb)


def approve(db: Session, wb: TrWorkbook, rows: list[TrRow], actor, justification: str, use_suggestion: bool = False, status: str | None = None,
            testers: dict | None = None, test_results: str | None = None, sod_details: dict | None = None) -> dict:
    """SCA only (enforced by the API). Compliance Status = the recorded assessment determination or the SCA's own determination.
    Tested By / Date Tested are ALWAYS entered (or explicitly confirmed) by the SCA per row – never defaulted to the approver or today.
    `testers` = {row_id: {"tested_by": str, "date_tested": date}} – validated by the API."""
    vocab = wb.definition["vocabulary"]["compliance_status"]
    done, skipped = [], []
    for r in rows:
        p = {**(r.proposed or {}), **(r.final or {})}
        if r.match_status in ("ambiguous", "unmatched"):
            skipped.append({"row": r.row_number, "reason": "Resolve the control/CCI mapping first"})
            continue
        det, how = p.get("compliance_status"), "assessment record"
        if status or (not det and use_suggestion and p.get("suggestion")):
            det, how = status or p["suggestion"], "SCA determination in TR review"
        if not det:
            skipped.append({"row": r.row_number, "reason": "No assessor determination – choose a Compliance Status or accept the suggestion"})
            continue
        if det not in vocab:
            skipped.append({"row": r.row_number, "reason": f"“{det}” is not an allowed value ({', '.join(vocab)})"})
            continue
        t = testers[str(r.id)]
        by, on = t["tested_by"].strip(), t["date_tested"].isoformat()
        prop = r.proposed or {}
        src = (f"confirmed from {prop['determination_source']}" if prop.get("determination_source") and prop.get("tested_by") == by and prop.get("date_tested") == on
               else "entered by the SCA at approval")
        r.final = {"compliance_status": det, "tested_by": by, "date_tested": on, "tester_confirmed": True, "tester_source": src,
                   "test_results": test_results if test_results is not None else p.get("test_results"), "determined_by": how,
                   "approved_by": actor.email, "approved_at": now().isoformat(),
                   **({"sod_label": sod_details[str(r.id)]["label"]} if sod_details and (sod_details.get(str(r.id)) or {}).get("label") else {}),
                   **({"ai_assisted": p["ai_assisted"]} if p.get("ai_assisted") else {})}
        r.state, r.decided_by, r.decided_at, r.justification = "approved", actor.email, now(), justification
        done.append(r.row_number)
    return {"approved": done, "skipped": skipped}


def actor_user(db: Session, actor):
    return actor if isinstance(actor, User) else db.get(User, actor.id)


def edit(db: Session, wb: TrWorkbook, r: TrRow, values: dict, actor, justification: str) -> TrRow:
    allowed = {"test_results", "compliance_status"}
    vals = {k: v for k, v in values.items() if k in allowed}
    if vals.get("compliance_status") and vals["compliance_status"] not in wb.definition["vocabulary"]["compliance_status"]:
        raise ValueError(f"“{vals['compliance_status']}” is not an allowed Compliance Status")
    if "test_results" in vals and len(vals["test_results"] or "") > ((wb.definition.get("limits") or {}).get("test_results") or 4000):
        raise ValueError("Test Results exceed the template's character limit")
    r.final = {**(r.final or {}), **vals, "tester_confirmed": False}   # tester/date are entered again at (re-)approval
    if vals.get("compliance_status"):
        r.final["determined_by"] = "SCA determination in TR review"
    r.state, r.decided_by, r.decided_at, r.justification = "edited", actor.email, now(), justification
    refresh_counts(db, wb)
    return r


def sandbox_rows(rows: list[TrRow]) -> int:
    return sum(1 for r in rows if r.state == "approved" and (r.final or {}).get("sod_label"))


def export(db: Session, wb: TrWorkbook, actor) -> TrExport:
    from app.services.storage.base import get_storage

    store = get_storage()
    orig = store.get(wb.storage_key)
    d = wb.definition
    roles, sheet = d["roles"], d["sheet"]
    permitted = set(roles[r] for r in xlsx.RESULT_ROLES if r in roles)
    cells: dict = {}
    rows = list(db.scalars(select(TrRow).where(TrRow.workbook_id == wb.id).order_by(TrRow.row_number)))
    mode = export_mode(rows)
    ev_ids = set()
    for r in rows:
        if mode == "blank" and r.added:  # blank-template mode: identifiers of the rows added from the system
            for role, val in (("control_acronym", r.control_acronym), ("ap_acronym", r.ap_acronym), ("cci", (r.cci_id or "").replace("CCI-", "")),
                              ("cci_definition", (r.proposed or {}).get("cci_definition"))):
                if role in roles:
                    cells[f"{roles[role]}{r.row_number}"] = val
        f = r.final or {}
        if r.state != "approved" or not f.get("tester_confirmed"):
            continue
        on = date.fromisoformat(f["date_tested"]) if f.get("date_tested") else None
        val = {"compliance_status": f.get("compliance_status"), "tested_by": (f.get("tested_by") or "")[: (d.get("limits") or {}).get("tested_by") or 150] or None,
               "test_results": f.get("test_results"),
               "date_tested": on if d.get("date_write") == "serial" else (on.strftime("%d-%b-%Y") if on else None)}
        for role, v in val.items():
            if role in roles:
                cells[f"{roles[role]}{r.row_number}"] = v
        ev_ids.update(s["id"] for s in (r.sources or []) if s["type"] == "evidence")
    out = xlsx.write(orig, sheet, cells)
    rep = xlsx.verify(orig, out, sheet, set(cells))
    rep.update(permission_report(d, wb, rows, mode, cells))
    rep["sandbox_same_person_rows"] = sandbox_rows(rows)
    rep["ok"] = rep["ok"] and not rep["outside_permitted"]
    if not rep["ok"]:
        raise ValueError(f"Export validation failed: {rep['issues'][:3] or rep['outside_permitted'][:3]}")
    ver = (db.scalar(select(func.max(TrExport.version)).where(TrExport.workbook_id == wb.id)) or 0) + 1
    sha = hashlib.sha256(out).hexdigest()
    name = re.sub(r"\.xlsx$", "", wb.file_name, flags=re.I) + f"_TR_export_v{ver}.xlsx"
    key = f"tr/{wb.org_id}/{wb.system_id}/exports/{wb.id}/v{ver}/{name}"
    store.put(key, out, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    counts = refresh_counts(db, wb)["state"]
    t = (d.get("template") or {}).get("version")
    exp = TrExport(workbook_id=wb.id, system_id=wb.system_id, version=ver, file_name=name, sha256=sha, storage_key=key, exported_by=actor.email,
                   counts={**counts, "written_rows": sum(1 for r in rows if r.state == "approved")}, evidence_ids=sorted(ev_ids), verification=rep,
                   template_version=t)
    db.add(exp)
    for r in rows:
        if r.state == "approved":
            r.last_exported_at = now()
    db.flush()
    return exp


def export_mode(rows: list[TrRow]) -> str:
    """blank = the template had no data rows and rows were added from the system; prepopulated = eMASS-exported rows (only M–P written)."""
    return "blank" if rows and all(r.added for r in rows) else "prepopulated"


def permission_report(d: dict, wb: TrWorkbook, rows: list[TrRow], mode: str, cells: dict) -> dict:
    roles, sheet, first = d["roles"], d["sheet"], d["header_row"] + 1
    last = max(wb.detection.get("last_row") or 0, max((r.row_number for r in rows), default=0))
    res_cols = {roles[k]: k for k in xlsx.RESULT_ROLES if k in roles}
    id_cols = {roles[k]: k for k in xlsx.ID_ROLES if k in roles} if mode == "blank" else {}
    added = {r.row_number for r in rows if r.added}
    ranges = [{"range": f"{sheet}!{c}{first}:{c}{last}", "column": k, "why": "Result column (“Enter Test Results Here”) – written only for SCA-approved rows "
               "with a confirmed tester and test date"} for c, k in sorted(res_cols.items())]
    if mode == "blank":
        ar = f"{min(added)}:{max(added)}" if added else "–"
        ranges += [{"range": f"{sheet}!{c}{min(added)}:{c}{max(added)}" if added else f"{sheet}!{c}", "column": k,
                    "why": f"Blank-template mode only: the template had no data rows, so rows {ar} were added from the system's control set and eMASS needs "
                           "Control Acronym / AP Acronym / CCI / CCI Definition to identify them. Written only on those added rows."} for c, k in sorted(id_cols.items())]

    def ok(ref):
        col, row = re.match(r"[A-Z]+", ref).group(0), int(re.search(r"\d+", ref).group(0))
        return row >= first and (col in res_cols or (col in id_cols and row in added))

    return {"mode": mode, "mode_label": {"blank": "Blank-template mode (rows added from the system)", "prepopulated": "Pre-populated mode (eMASS rows; only M–P written)"}[mode],
            "permitted_ranges": ranges, "permitted_range": [x["range"] for x in ranges],
            "changed_cells": sorted(cells, key=lambda ref: (int(re.search(r"\d+", ref).group(0)), ref)),
            "outside_permitted": [ref for ref in cells if not ok(ref)],
            "proof": "Every cell outside the permitted ranges was compared between the original and the export (value, formula, style) and every other "
                     "package part was compared byte-for-byte; see checked_cells / issues." + (" Columns other than " + ", ".join(sorted(set(res_cols) | set(id_cols))) +
                                                                                             " are never written in this mode." if mode == "blank" else
                                                                                             " Identifier and control-information columns (all columns except " +
                                                                                             ", ".join(sorted(res_cols)) + ") are never written in this mode."),
}


_ = (uuid,)
