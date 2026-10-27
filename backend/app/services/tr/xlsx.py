"""Workbook engine for eMASS Excel templates. Detection maps columns by header text (never by position); writing edits only
the target cells' XML inside a byte-for-byte copy of the original package, so styles, merges, validations, defined names,
tables, VBA and every other part stay identical; verify() proves it."""
import hashlib
import io
import re
import zipfile
from datetime import date, datetime

import openpyxl
from lxml import etree
from openpyxl.utils import column_index_from_string, get_column_letter

NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
RNS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
ROLES = {
    "control_acronym": {"control acronym", "control", "control id", "security control"},
    "ap_acronym": {"ap acronym", "ap", "ap id", "ap number", "assessment procedure acronym", "ap #"},
    "cci": {"cci", "cci id", "cci number", "cci #", "cci no"},
    "cci_definition": {"cci definition", "cci text", "cci description"},
    "implementation_guidance": {"implementation guidance", "guidance"},
    "assessment_procedures": {"assessment procedures", "assessment procedure", "test procedure", "test procedures"},
    "compliance_status": {"compliance status", "compliance", "test status", "status", "result"},
    "date_tested": {"date tested", "test date", "tested on", "date of test"},
    "tested_by": {"tested by", "tester", "assessor", "tested by name"},
    "test_results": {"test results", "test result", "results", "results narrative", "test result narrative"},
}
RESULT_ROLES = ("compliance_status", "date_tested", "tested_by", "test_results")
ID_ROLES = ("control_acronym", "ap_acronym", "cci", "cci_definition")
DEFAULT_VOCAB = ["Compliant", "Non-Compliant", "Not Applicable"]
SKIP_SHEETS = re.compile(r"example|instruction|assignment values|lists?\b", re.I)


def norm(s) -> str:
    return re.sub(r"[^a-z0-9#]+", " ", str(s or "").lower()).strip()


def role_of(header) -> str | None:
    n = norm(header)
    return next((r for r, syn in ROLES.items() if n in syn), None)


def _banners(ws, row: int) -> dict[int, str]:
    """Group banner text per column from the row above the header (merged ranges or single cells)."""
    out = {}
    if row < 1:
        return out
    for m in ws.merged_cells.ranges:
        if m.min_row <= row <= m.max_row and ws.cell(m.min_row, m.min_col).value:
            for c in range(m.min_col, m.max_col + 1):
                out[c] = str(ws.cell(m.min_row, m.min_col).value)
    for c in range(1, ws.max_column + 1):
        if c not in out and ws.cell(row, c).value:
            out[c] = str(ws.cell(row, c).value)
    return out


def _group(text: str | None) -> str | None:
    t = (text or "").lower()
    if "enter" in t:
        return "editable"
    if "read-only" in t or "read only" in t or "provided by" in t or "latest" in t:
        return "readonly"
    return None


def _header_row(ws) -> tuple[int, int]:
    best = (0, 0)
    for r in range(1, min(ws.max_row, 40) + 1):
        n = len({role_of(ws.cell(r, c).value) for c in range(1, min(ws.max_column, 80) + 1)} - {None})
        if n > best[1]:
            best = (r, n)
    return best


def _list_validations(ws) -> list[tuple[list, list[str]]]:
    out = []
    for dv in ws.data_validations.dataValidation:
        if dv.type == "list" and dv.formula1 and dv.formula1.startswith('"'):
            cols = sorted({column_index_from_string(re.match(r"\$?([A-Z]+)", part).group(1)) for part in str(dv.sqref).split()})
            out.append((cols, [v.strip() for v in dv.formula1.strip('"').split(",")]))
    return out


def _limits(ws, col: int) -> int | None:
    for dv in ws.data_validations.dataValidation:
        if dv.type == "textLength" and any(column_index_from_string(re.match(r"\$?([A-Z]+)", p).group(1)) == col for p in str(dv.sqref).split()):
            try:
                return int(dv.formula1)
            except (TypeError, ValueError):
                return None
    return None


def detect(data: bytes, min_score: int = 3) -> dict:
    wb = openpyxl.load_workbook(io.BytesIO(data), keep_vba=False)
    sheets = []
    for ws in wb.worksheets:
        r, n = _header_row(ws)
        sheets.append({"title": ws.title, "state": ws.sheet_state, "header_row": r, "score": n})
    cands = [s for s in sheets if s["score"] >= min_score and s["state"] == "visible" and not SKIP_SHEETS.search(s["title"])] or [s for s in sheets if s["score"] >= min_score]
    if not cands:
        raise ValueError("No Test Result header row found (looked for AP Acronym, CCI, Compliance Status, Date Tested, Tested By, Test Results in the first 40 rows of every sheet)")
    main = max(cands, key=lambda s: (s["score"], s["title"].lower() == "template"))
    ws = wb[main["title"]]
    hr = main["header_row"]
    banners = _banners(ws, hr - 1)
    cols, roles = [], {}
    for c in range(1, ws.max_column + 1):
        h = ws.cell(hr, c).value
        if h is None:
            continue
        cols.append({"col": get_column_letter(c), "idx": c, "header": str(h), "role": role_of(h), "banner": banners.get(c), "group": _group(banners.get(c))})
    for role in ROLES:
        hits = [x for x in cols if x["role"] == role]
        if not hits:
            continue
        pick = next((x for x in hits if x["group"] == "editable"), None) if role in RESULT_ROLES else None
        pick = pick or next((x for x in hits if x["group"] != "readonly"), hits[0])
        roles[role] = pick["col"]
        for x in hits:
            if x is not pick:
                x["role"] = f"latest_{role}" if role in RESULT_ROLES else None
    # Value vocabulary: list validation on the column, else the same header in another sheet (e.g. Example), else Instructions text.
    vocab, vsrc = None, None
    cs = roles.get("compliance_status")
    if cs:
        idx = column_index_from_string(cs)
        vocab = next((v for cols_, v in _list_validations(ws) if idx in cols_), None)
        vsrc = f"data validation on {ws.title}!{cs}" if vocab else None
        for other in wb.worksheets:
            if vocab or other.title == ws.title:
                continue
            ohr = _header_row(other)[0]
            for cols_, v in _list_validations(other):
                if any(role_of(other.cell(ohr, c).value) == "compliance_status" for c in cols_):
                    vocab, vsrc = v, f"data validation on {other.title} (same header)"
                    break
        if not vocab:
            for other in wb.worksheets:
                for row in other.iter_rows(max_row=60, max_col=3, values_only=True):
                    m = next((re.search(r"\(([A-Z][\w-]+(?:\s*,\s*[A-Z][\w -]+)+)\)", str(v)) for v in row if v and "compliance status" in str(v).lower()), None)
                    if m:
                        vocab, vsrc = [x.strip() for x in m.group(1).split(",")], f"{other.title} sheet text"
                        break
                if vocab:
                    break
    vocab, vsrc = vocab or DEFAULT_VOCAB, vsrc or "eMASS default (no list found in the workbook)"
    date_fmt, date_write = "dd-mmm-yyyy", "text"
    if roles.get("date_tested"):
        fmt = ws[f"{roles['date_tested']}{hr + 1}"].number_format
        if re.search(r"[dy]", fmt or "", re.I) and fmt != "General":
            date_fmt, date_write = fmt, "serial"
    limits = {r: _limits(ws, column_index_from_string(c)) for r, c in roles.items() if r in RESULT_ROLES}
    rows = []
    for r in range(hr + 1, ws.max_row + 1):
        vals = {role: ws[f"{c}{r}"].value for role, c in roles.items() if role in ID_ROLES}
        if any(v not in (None, "") for k, v in vals.items() if k in ("ap_acronym", "cci", "control_acronym")):
            rows.append({"row": r, **{k: (str(v).strip() if v is not None else None) for k, v in vals.items()}})
    fp = hashlib.sha256("|".join([ws.title, str(hr)] + [f"{x['col']}={norm(x['header'])}" for x in cols]).encode()).hexdigest()
    return {"sheets": sheets, "sheet": ws.title, "header_row": hr, "columns": cols, "roles": roles, "fingerprint": fp,
            "vocabulary": {"compliance_status": vocab, "source": vsrc}, "date_format": date_fmt, "date_write": date_write, "limits": limits,
            "rows": rows, "last_row": ws.max_row, "readonly_cols": [x["col"] for x in cols if x["group"] == "readonly"]}


def sheet_part(z: zipfile.ZipFile, title: str) -> str:
    wbx = etree.fromstring(z.read("xl/workbook.xml"))
    rid = next(s.get(f"{{{RNS}}}id") for s in wbx.iter(f"{{{NS}}}sheet") if s.get("name") == title)
    rels = etree.fromstring(z.read("xl/_rels/workbook.xml.rels"))
    target = next(r.get("Target") for r in rels if r.get("Id") == rid)
    return target.lstrip("/") if target.startswith("/") else f"xl/{target}"


def excel_serial(d: date) -> int:
    return (d - date(1899, 12, 30)).days


def _cell(row_el, ref: str):
    for c in row_el.iter(f"{{{NS}}}c"):
        if c.get("r") == ref:
            return c
    col = column_index_from_string(re.match(r"[A-Z]+", ref).group(0))
    new = etree.Element(f"{{{NS}}}c", r=ref)
    for i, c in enumerate(row_el.findall(f"{{{NS}}}c")):
        if column_index_from_string(re.match(r"[A-Z]+", c.get("r")).group(0)) > col:
            row_el.insert(i, new)
            return new
    row_el.append(new)
    return new


def _row(sheet_data, r: int):
    for i, el in enumerate(sheet_data.findall(f"{{{NS}}}row")):
        n = int(el.get("r"))
        if n == r:
            return el
        if n > r:
            new = etree.Element(f"{{{NS}}}row", r=str(r))
            sheet_data.insert(i, new)
            return new
    new = etree.SubElement(sheet_data, f"{{{NS}}}row", r=str(r))
    return new


def write(data: bytes, sheet_title: str, cells: dict[str, object]) -> bytes:
    """Copy the package byte-for-byte and change only the given cells (A1 refs) in one sheet. Strings are inline (sharedStrings untouched)."""
    zin = zipfile.ZipFile(io.BytesIO(data))
    part = sheet_part(zin, sheet_title)
    root = etree.fromstring(zin.read(part))
    sd = root.find(f"{{{NS}}}sheetData")
    for ref, val in cells.items():
        r = int(re.search(r"\d+", ref).group(0))
        c = _cell(_row(sd, r), ref)
        for child in list(c):
            c.remove(child)
        c.attrib.pop("t", None)
        if val is None or val == "":
            continue
        if isinstance(val, (date, datetime)):
            etree.SubElement(c, f"{{{NS}}}v").text = str(excel_serial(val if isinstance(val, date) and not isinstance(val, datetime) else val.date()))
        elif isinstance(val, (int, float)):
            etree.SubElement(c, f"{{{NS}}}v").text = str(val)
        else:
            c.set("t", "inlineStr")
            t = etree.SubElement(etree.SubElement(c, f"{{{NS}}}is"), f"{{{NS}}}t")
            t.text = str(val)
            t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    new_xml = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as zout:
        for info in zin.infolist():
            zout.writestr(info, new_xml if info.filename == part else zin.read(info.filename))
    return out.getvalue()


def verify(orig: bytes, out: bytes, sheet_title: str, targets: set[str]) -> dict:
    """Prove that only the target cells changed: every other zip part byte-identical; in the target sheet every non-target cell,
    row and non-cell element identical; re-opens with openpyxl with the same sheets, merges, validations and defined names."""
    issues = []
    za, zb = zipfile.ZipFile(io.BytesIO(orig)), zipfile.ZipFile(io.BytesIO(out))
    if zb.testzip() is not None:
        issues.append("zip CRC check failed")
    if [i.filename for i in za.infolist()] != [i.filename for i in zb.infolist()]:
        issues.append("zip part list or order changed")
    part = sheet_part(za, sheet_title)
    for n in za.namelist():
        if n != part and za.read(n) != zb.read(n):
            issues.append(f"part changed: {n}")
    a, b = etree.fromstring(za.read(part)), etree.fromstring(zb.read(part))
    for ca, cb in zip(a, b):
        if ca.tag != cb.tag:
            issues.append(f"element order changed: {ca.tag}")
        elif not ca.tag.endswith("sheetData") and etree.tostring(ca) != etree.tostring(cb):
            issues.append(f"sheet element changed: {etree.QName(ca).localname}")
    cells_b = {c.get("r"): c for c in b.iter(f"{{{NS}}}c")}
    checked = 0
    for c in a.iter(f"{{{NS}}}c"):
        ref = c.get("r")
        if ref in targets:
            if cells_b.get(ref) is not None and cells_b[ref].get("s") != c.get("s"):
                issues.append(f"style changed on target {ref}")
            continue
        checked += 1
        if ref not in cells_b or etree.tostring(cells_b[ref]) != etree.tostring(c):
            issues.append(f"non-target cell changed: {ref}")
    rows_a = {r.get("r"): dict(r.attrib) for r in a.iter(f"{{{NS}}}row")}
    for r in b.iter(f"{{{NS}}}row"):
        if r.get("r") in rows_a and dict(r.attrib) != rows_a[r.get("r")]:
            issues.append(f"row attributes changed: {r.get('r')}")
    try:
        wa = openpyxl.load_workbook(io.BytesIO(orig))
        wb2 = openpyxl.load_workbook(io.BytesIO(out))
        if wa.sheetnames != wb2.sheetnames:
            issues.append("sheet names changed")
        if dict(wa.defined_names.items()).keys() != dict(wb2.defined_names.items()).keys():
            issues.append("defined names changed")
        for s in wa.sheetnames:
            if {str(m) for m in wa[s].merged_cells.ranges} != {str(m) for m in wb2[s].merged_cells.ranges}:
                issues.append(f"merged cells changed on {s}")
            if [(d.type, d.formula1, str(d.sqref)) for d in wa[s].data_validations.dataValidation] != \
               [(d.type, d.formula1, str(d.sqref)) for d in wb2[s].data_validations.dataValidation]:
                issues.append(f"data validations changed on {s}")
    except Exception as exc:  # the output must re-open cleanly
        issues.append(f"openpyxl could not re-open the file: {exc}")
    return {"ok": not issues, "issues": issues[:50], "checked_cells": checked, "target_cells": len(targets), "parts": len(za.namelist())}
