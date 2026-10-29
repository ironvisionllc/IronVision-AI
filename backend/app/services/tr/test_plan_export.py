"""Export Test Plan: every assessment procedure (control × CCI) of the system's SELECTED, non-withdrawn controls in the same
eMASS-style Test Result layout that the Test Plan importer reads (sheet "Template", header row 2), so the file re-imports cleanly."""
import hashlib
import io
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models import Cci, CciControlMap, Control, ControlAssignment, ControlObjective, System, TrRow, TrWorkbook

SHEET = "Template"
HEADER_ROW = 2
COLUMNS = ["Control Acronym", "AP Acronym", "CCI", "CCI Definition", "Assessment Procedures", "Test Method", "Inherited",
           "Compliance Status", "Date Tested", "Tested By", "Test Results"]
LABEL = "eMASS-style layout – pending verification against a real eMASS export"
INHERITED = {"inherited": "Yes", "shared": "Hybrid (shared)", "system": "No"}


def control_acronym(cid: str) -> str:
    return cid.upper().replace(".", "(") + (")" if "." in cid else "")


def _procedures(db: Session, control_uuid) -> str:
    lines = [f"{o.label} {o.prose.strip()}" for o in db.scalars(select(ControlObjective).where(ControlObjective.control_uuid == control_uuid)
                                                               .order_by(ControlObjective.sort_order)) if o.prose and o.prose.strip()]
    return "\n".join(lines)[:4000]


def _methods(c: Control) -> str:
    return "; ".join(m["method"].title() for m in (c.assessment_methods or []) if isinstance(m, dict) and m.get("method"))


def _tr_rows(db: Session, s: System) -> dict:
    """Latest Test Results row per (control, CCI); an approved row wins over a later unapproved one."""
    out: dict = {}
    for r, _w in db.execute(select(TrRow, TrWorkbook).join(TrWorkbook, TrRow.workbook_id == TrWorkbook.id)
                            .where(TrWorkbook.system_id == s.id, TrWorkbook.status != "failed", TrRow.control_id.is_not(None))
                            .order_by(TrWorkbook.version, TrWorkbook.created_at, TrRow.row_number)):
        k = (r.control_id, r.cci_id)
        if r.state == "approved" or (k not in out or out[k].state != "approved"):
            out[k] = r
    return out


def rows(db: Session, s: System) -> list[dict]:
    defs = dict(db.execute(select(Cci.cci_id, Cci.definition)).all())
    tr = _tr_rows(db, s)
    out = []
    q = (select(ControlAssignment, Control).join(Control, Control.id == ControlAssignment.control_uuid)
         .where(ControlAssignment.system_id == s.id, Control.is_withdrawn.is_(False)).order_by(Control.sort_id, Control.control_id))
    for a, c in db.execute(q):
        ccis = sorted(set(db.scalars(select(CciControlMap.cci_id).where(CciControlMap.control_id == a.control_id))))
        proc, meth = _procedures(db, c.id), _methods(c)
        for i, cci in enumerate(ccis, start=1):
            t = tr.get((a.control_id, cci))
            fin = (t.final or {}) if t is not None and t.state == "approved" else {}
            d = fin.get("date_tested")
            out.append({"Control Acronym": control_acronym(a.control_id), "AP Acronym": (t.ap_acronym if t is not None and t.ap_acronym else f"{control_acronym(a.control_id)}.{i}"),
                        "CCI": cci, "CCI Definition": defs.get(cci), "Assessment Procedures": proc or None, "Test Method": meth or None,
                        "Inherited": INHERITED.get(a.responsibility, a.responsibility), "Compliance Status": fin.get("compliance_status"),
                        "Date Tested": date.fromisoformat(d).strftime("%d-%b-%Y") if d else None, "Tested By": fin.get("tested_by"),
                        "Test Results": fin.get("test_results")})
    return out


def _text(cell, v) -> None:
    cell.value = v
    if isinstance(v, str):
        cell.data_type = "s"  # never a formula (=, +, -, @ at the start stay text)


def build(db: Session, s: System) -> tuple[bytes, list[dict], str]:
    data = rows(db, s)
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET
    ws.cell(1, 1, f"Test Plan – {s.acronym or s.name} – exported {date.today():%d-%b-%Y} by IronVision AI. {LABEL}.").font = Font(italic=True)
    for i, h in enumerate(COLUMNS, start=1):
        c = ws.cell(HEADER_ROW, i, h)
        c.font, c.fill = Font(bold=True), PatternFill("solid", fgColor="DDE7F0")
        ws.column_dimensions[c.column_letter].width = {"CCI Definition": 50, "Assessment Procedures": 70, "Test Results": 50}.get(h, 18)
    for r, row in enumerate(data, start=HEADER_ROW + 1):
        for i, h in enumerate(COLUMNS, start=1):
            _text(ws.cell(r, i), row[h])
            ws.cell(r, i).alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = ws.cell(HEADER_ROW + 1, 1)
    b = io.BytesIO()
    wb.save(b)
    out = b.getvalue()
    return out, data, hashlib.sha256(out).hexdigest()


def summary(data: list[dict]) -> dict:
    return {"rows": len(data), "controls": len({r["Control Acronym"] for r in data}), "inherited": sum(1 for r in data if r["Inherited"] == "Yes"),
            "with_results": sum(1 for r in data if r["Compliance Status"]), "columns": COLUMNS, "label": LABEL}
