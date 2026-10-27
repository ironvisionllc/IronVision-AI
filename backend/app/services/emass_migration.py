"""eMASS System Import & Migration – SLCM worksheet, Categorization template and PPSM boundaries on the shared eMASS import
(emass_reports / emass_import sessions): parse → per-field review (imported values preselected) → apply with provenance, audit and undo.
Also: exports for round trips, AI column assist (redacted, suggestion-only), copy-setup from another system."""
import asyncio
import io
import re
from datetime import date, datetime, timezone
from pathlib import Path

from fastapi import HTTPException
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models import (CategorizationRecord, Control, ControlAssignment, EmassControlFields, IscmProfile, IscmStrategy, OdpValue, Policy,
                        PpsmEntry, System, TailoringDecision)
from app.services import emass_fields as ef
from app.services import emass_import as ei

_n = ei._n
TYPES = {"slcm": "SLCM Worksheet", "categorization": "Categorization", "ppsm": "PPSM Boundaries"}
DEFERRED = {"control_information": "Control Information", "poam": "POA&M"}
FIXTURES = Path(__file__).resolve().parents[2] / "tests" / "fixtures" / "emass"
LEVELS = ["Low", "Moderate", "High"]

SLCM_COLS = {"control": ["control acronym", "security control #", "security control number", "control"],
             "slcm_criticality": ["criticality", "slcm criticality"], "slcm_frequency": ["frequency", "slcm frequency"],
             "slcm_method": ["method", "slcm method"], "slcm_reporting": ["reporting", "slcm reporting"],
             "slcm_tracking": ["tracking", "slcm tracking"], "slcm_comments": ["slcm comments"]}
SLCM_LABEL = {"slcm_criticality": "SLCM Criticality", "slcm_frequency": "SLCM Frequency", "slcm_method": "SLCM Method",
              "slcm_reporting": "SLCM Reporting", "slcm_tracking": "SLCM Tracking", "slcm_comments": "SLCM Comments"}
SLCM_SHEETS = ("final results", "all security controls")
PPSM_COLS = {"asset_id": ["#", "asset id #", "asset id"], "port": ["port (required)", "port"], "protocol": ["protocol (required)", "protocol"],
             "data_service": ["data service (required)", "data service"], "boundary": ["boundary (required)", "boundary", "boundaries"],
             "fqdn": ["fqdn"], "purpose": ["purpose"]}
PPSM_LABEL = {"port": "Port", "protocol": "Protocol", "data_service": "Data Service", "boundary": "Boundary", "fqdn": "FQDN", "purpose": "Purpose"}
PPSM_LIMIT = {"port": 20, "protocol": 20, "data_service": 100, "fqdn": 100, "purpose": 1000}
PPSM_REQUIRED = ("port", "protocol", "data_service", "boundary")
BOUNDARIES = ["Ext to DoD GW (In)", "DoD GW to Ext (Out)", "DoD GW to DoD DMZ or RDTE (In)", "DoD DMZ or RDTE to DoD GW (Out)", "DoD GW to DISN (In)",
              "DISN to DoD GW (Out)", "DISN to Enclave GW (In)", "Enclave GW to DISN (Out)", "Enclave GW to Enclave DMZ (In)",
              "Enclave DMZ to Enclave GW (Out)", "Enclave GW to Enclave (In)", "Enclave to Enclave GW (Out)", "Ext ISP to Enclave GW (In)",
              "Enclave GW to Ext ISP (Out)", "Multiple AO E2E Tunnel", "Single AO E2E Tunnel"]
_BND = {**{_n(b): b for b in BOUNDARIES}, **{str(i + 1): b for i, b in enumerate(BOUNDARIES)}}
REQUIRED_FIELDS = {"slcm": ("control",), "ppsm": PPSM_REQUIRED}
COLS = {"slcm": SLCM_COLS, "ppsm": PPSM_COLS}


# ---------- cell helpers ----------
def clean(v) -> str | None:
    """Cell text: '---' placeholders → empty; an apostrophe we added to neutralize a formula on export is removed again."""
    if v is None:
        return None
    s = str(v).strip()
    if s in ("", "---", "--", "-"):
        return None
    if len(s) > 1 and s[0] == "'" and s[1] in "=+-@":
        s = s[1:]
    return s


def safe_cell(v):
    """Formula/CSV injection neutralization for exported text."""
    if isinstance(v, str) and v[:1] in ("=", "+", "-", "@", "\t", "\r"):
        return "'" + v
    return v


def _f(cur, imp_v, err=None, blocked=None, lab=None, sheet=None, row=None, col=None) -> dict:
    from app.services.emass_reports import _field
    f = _field(cur, imp_v, err, blocked, lab)
    f["src"] = {"sheet": sheet, "row": row, "col": ei._col(col) if col is not None else None}
    return f


def preselect(items: list[dict]) -> None:
    """Imported values win by default: every valid change is preselected 'Use imported' unless it came from a low-trust AI mapping."""
    for i in items:
        if i.get("required_issue") or i.get("ai_mapped"):
            continue
        for f in i["fields"].values():
            if ei._ok(f) and f["decision"] is None:
                f["decision"] = "accept"


def _pick_sheet(sheets, cols: dict, need: str, prefer=()) -> tuple | None:
    from app.services.emass_reports import _header
    aliases = {_n(a): f for f, al in cols.items() for a in al}
    best = None
    for title, rows in sheets:
        h = _header(rows, aliases, need=need)
        if not h:
            continue
        rank = (next((k for k, p in enumerate(prefer) if p in title.lower()), 99), -len(h[1]))
        if best is None or rank < best[0]:
            best = (rank, title, rows, h[0], h[1])
    return best[1:] if best else None


# ---------- redaction + AI column assist ----------
_EMAIL = re.compile(r"[\w.+-]+@[\w-]+(\.[\w-]+)+")
_PHONE = re.compile(r"(\+?\d[\d\s().-]{7,}\d)")
_IP4 = re.compile(r"\b\d{1,3}(\.\d{1,3}){3}(/\d{1,2})?\b")
_IP6 = re.compile(r"\b([0-9a-f]{0,4}:){2,7}[0-9a-f]{0,4}\b", re.I)
_FQDN = re.compile(r"\b[a-z0-9-]+(\.[a-z0-9-]+)+\.(mil|local|internal|lan|corp|intranet|int|gov|smil|sgov|home|ad|net|com|org)\b", re.I)
_PERSON_HDR = re.compile(r"name|poc|owner|contact|tested by|e-?mail|phone|user|person|assigned|reviewer|approver", re.I)


def redact(text) -> str:
    s = str(text or "")
    for rx, tag in ((_EMAIL, "[EMAIL]"), (_IP6, "[IP]"), (_IP4, "[IP]"), (_FQDN, "[HOST]"), (_PHONE, "[PHONE]")):
        s = rx.sub(tag, s)
    return s[:120]


def ai_payload(header: tuple, samples: list[tuple]) -> str:
    out = []
    for c, h in enumerate(header):
        if h in (None, ""):
            continue
        vals = [] if _PERSON_HDR.search(str(h)) else [redact(r[c]) for r in samples if c < len(r) and r[c] not in (None, "")][:3]
        out.append(f"{c}: {redact(h)} | " + ("[REDACTED – person data]" if _PERSON_HDR.search(str(h)) else "; ".join(vals)))
    return "\n".join(out)


def ai_enabled(db: Session, org_id) -> bool:
    from app.services import rmf
    return bool(rmf.get_setting(db, org_id, "emass_ai", {"enabled": True}).get("enabled"))


def ai_suggest(report_type: str, header: tuple, samples: list[tuple]) -> list[dict] | None:
    """Column-mapping SUGGESTIONS from the configured AI provider (redacted input). None when AI is unavailable."""
    from app.services.ai.service import get_ai_service
    svc = get_ai_service()
    if not svc.configured:
        return None
    try:
        res = asyncio.run(svc.run("emass.column_map", "1.0.0", {"report_type": TYPES[report_type], "fields": ", ".join(COLS[report_type]),
                                                                 "columns": ai_payload(header, samples)}))
    except Exception:  # noqa: BLE001 – deterministic import must keep working
        return None
    return [m.model_dump() for m in res.output.mappings if m.field in COLS[report_type]]


def _ai_header(db, s, typ, sheets, rep) -> tuple | None:
    """No deterministic header: ask AI for a mapping of the most table-like sheet; only high/medium suggestions are used, all routed to review."""
    rep["ai"] = {"enabled": bool(s and ai_enabled(db, s.org_id)), "used": False}
    if not rep["ai"]["enabled"]:
        return None
    for title, rows in sheets:
        for hi, row in enumerate(rows[:ei.HEADER_SCAN_ROWS]):
            if sum(1 for v in row if v not in (None, "")) < 3:
                continue
            sug = ai_suggest(typ, row, rows[hi + 1:hi + 4]) or []
            use = {m["field"]: m for m in sug if m["confidence"] in ("high", "medium")}
            rep["ai"]["suggestions"] = sug
            rep["ai"]["low_confidence"] = [m for m in sug if m["confidence"] == "low"]
            if all(k in use for k in REQUIRED_FIELDS[typ]):
                rep["ai"].update(used=True, sheet=title, header_row=hi + 1)
                return title, rows, hi, {k: m["column"] for k, m in use.items()}, use
            return None
    return None


# ---------- SLCM ----------
def _enum(v: str, opts: list[str]) -> str | None:
    return next((o for o in opts if _n(o) == _n(v)), None)


def build_slcm(db: Session, s: System, sheets) -> tuple[list[dict], dict]:
    from app.services import iscm
    rep: dict = {"not_in_system": [], "invalid": [], "duplicates": [], "exceptions": []}
    found = _pick_sheet(sheets, SLCM_COLS, "control", SLCM_SHEETS)
    ai = None
    if found:
        title, rows, hi, cols = found
    elif (got := _ai_header(db, s, "slcm", sheets, rep)):
        title, rows, hi, cols, ai = got
    else:
        raise HTTPException(422, "No SLCM header row (Control Acronym + Criticality / Frequency / Method …) found")
    rep.update(sheet=title, header_row=hi + 1)
    assigned = {a.control_id: a for a in db.scalars(select(ControlAssignment).where(ControlAssignment.system_id == s.id))}
    fields = {f.assignment_id: f for f in db.scalars(select(EmassControlFields).where(EmassControlFields.system_id == s.id))}
    res = iscm.resolver(db, s)
    items, seen = [], {}
    for n, row in enumerate(rows[hi + 1:], start=hi + 2):
        raw = {k: (row[c] if c < len(row) else None) for k, c in cols.items()}
        if not any(clean(v) for v in raw.values()):
            continue
        cid = ei.norm_acronym(clean(raw.get("control")))
        if not cid:
            rep["exceptions"].append({"row": n, "control": str(raw.get("control"))[:40], "issue": "Not a control acronym"})
            continue
        if cid in seen:
            rep["duplicates"].append({"row": n, "control": ei.fmt_acronym(cid), "first_row": seen[cid]})
            continue
        seen[cid] = n
        a = assigned.get(cid)
        if a is None:
            rep["not_in_system"].append({"row": n, "control": ei.fmt_acronym(cid)})
            continue
        f = fields.get(a.id)
        cur = res.apply(cid, ef.emass_status(a, f), f, ef.values(f))[0]
        fs = {}
        for k, c in cols.items():
            if k == "control":
                continue
            v = clean(raw.get(k))
            if v is None:
                continue
            err = None
            if k == "slcm_frequency":
                v, err = (_enum(v, ef.FREQUENCIES), None) if _enum(v, ef.FREQUENCIES) else (v, f"'{v}' is not an eMASS SLCM Frequency")
            elif k == "slcm_method":
                v, err = (_enum(v, ef.METHODS), None) if _enum(v, ef.METHODS) else (v, f"'{v}' is not an eMASS SLCM Method")
            elif len(v) > ef.LIMITS.get(k, 2000):
                err = f"Longer than {ef.LIMITS.get(k, 2000)} characters"
            fs[k] = _f(cur.get(k), v, err, None, SLCM_LABEL[k], title, n, c)
            if ai:
                fs[k]["ai"] = ai.get(k)
            if err:
                rep["invalid"].append({"row": n, "control": ei.fmt_acronym(cid), "field": SLCM_LABEL[k], "value": v[:100], "error": err})
        items.append({"section": "SLCM", "key": cid, "control_id": cid, "label": ei.fmt_acronym(cid), "create": False, "row": n,
                      "fields": fs, "required_issue": None, "ai_mapped": bool(ai)})
    rep["matched"] = len(items)
    rep["missing_from_file"] = [ei.fmt_acronym(c) for c in sorted(set(assigned) - set(seen))][:2000]
    return items, rep


def apply_slcm(db: Session, s: System, i: dict, acc: dict, current, meta: dict, src: dict) -> dict:
    from app.services import audit
    a = db.scalar(select(ControlAssignment).where(ControlAssignment.system_id == s.id, ControlAssignment.control_id == i["control_id"]))
    if a is None:
        return {"control": i["label"], "applied": [], "failed": "control no longer in the system"}
    f = ef.ensure(db, a)
    before = {k: getattr(f, k) for k in acc}
    inherit_before = dict(f.slcm_inherit or {})
    follow = dict(inherit_before)
    for k, fd in acc.items():
        setattr(f, k, fd["imported"])
        follow.pop(k, None)
    f.slcm_inherit = follow or None
    f.updated_by = current.user.id
    audit.record(db, action="emass.slcm_imported", actor_user_id=current.user.id, actor_email=current.user.email, entity_type="control_assignment",
                 entity_id=str(a.id), before=before, after={k: fd["imported"] for k, fd in acc.items()},
                 details={**src, "row": i.get("row"), "provenance": {k: fd.get("src") for k, fd in acc.items()}}, **meta)
    return {"control": i["label"], "key": i["key"], "applied": sorted(acc), "failed": None,
            "undo": {"kind": "slcm", "assignment_id": str(a.id), "before": before, "inherit_before": inherit_before}}


def slcm_rows(db: Session, s: System) -> list[list]:
    from app.services import iscm
    res = iscm.resolver(db, s)
    fields = {f.assignment_id: f for f in db.scalars(select(EmassControlFields).where(EmassControlFields.system_id == s.id))}
    out = []
    for a, title in db.execute(select(ControlAssignment, Control.title).join(Control, Control.id == ControlAssignment.control_uuid)
                               .where(ControlAssignment.system_id == s.id)).all():
        f = fields.get(a.id)
        v = res.apply(a.control_id, ef.emass_status(a, f), f, ef.values(f))[0]
        out.append((a.control_id, title, [v.get(k) for k in SLCM_LABEL]))
    out.sort(key=lambda r: ei.fmt_acronym(r[0]))
    return [[n + 1, ei.fmt_acronym(cid), t, *vals, n + 1] for n, (cid, t, vals) in enumerate(out)]


def slcm_export(db: Session, s: System) -> bytes:
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Final Results"
    ws.append(["Implementation Status / Security Control Status"])
    ws.append(["Sort #", "Control Acronym", "Control Title", "Criticality", "Frequency", "Method", "Reporting", "Tracking", "SLCM Comments", "Order #"])
    for r in slcm_rows(db, s):
        ws.append([safe_cell(v) for v in r])
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


# ---------- Categorization ----------
def parse_categorization(sheets) -> dict:
    for title, rows in sheets:
        hi = next((k for k, r in enumerate(rows[:ei.HEADER_SCAN_ROWS]) if {"missiontype", "applicability"} <= {_n(v) for v in r if v is not None}), None)
        if hi is None:
            continue
        meta = {}
        for r in rows[:hi]:
            if r and r[0] is not None and _n(r[0]) in ("package", "date") and len(r) > 1:
                v = clean(r[1])
                if v and not v.lower().startswith("insert"):
                    meta[_n(r[0])] = v
        cat, sel, bad = None, [], []
        for n, r in enumerate(rows[hi + 1:], start=hi + 2):
            r = tuple(r) + (None,) * 6
            if r[0] not in (None, ""):
                cat = str(r[0]).strip()
            name = clean(r[1]) or (str(r[0]).strip() if r[0] else None)
            app = _n(r[2])
            if not name or app not in ("yes", "no", ""):
                if app not in ("yes", "no", ""):
                    bad.append({"row": n, "control": name or "", "issue": f"Applicability '{r[2]}' is not Yes/No"})
                continue
            if app == "yes":
                lv = [next((x for x in LEVELS if _n(x) == _n(r[c])), None) for c in (3, 4, 5)]
                sel.append({"row": n, "mission_type": f"{cat} – {name}" if cat and cat != name else name, "impacts": lv})
        hw = [max((x["impacts"][k] for x in sel if x["impacts"][k]), key=LEVELS.index, default=None) for k in range(3)]
        return {"sheet": title, "header_row": hi + 1, "meta": meta, "selected": sel, "impacts": hw, "exceptions": bad}
    raise HTTPException(422, "No 'Mission Type | Applicability' table found – not an eMASS Categorization template")


def build_categorization(db: Session, s: System, sheets) -> tuple[list[dict], dict]:
    p = parse_categorization(sheets)
    rep = {"sheet": p["sheet"], "header_row": p["header_row"], "exceptions": p["exceptions"], "mission_types": p["selected"], "package": p["meta"]}
    if not p["selected"]:
        rep["exceptions"].append({"row": None, "control": "", "issue": "No mission type is marked Applicability = Yes – nothing to propose"})
        return [], {**rep, "matched": 0}
    approved = db.scalar(select(CategorizationRecord).where(CategorizationRecord.system_id == s.id, CategorizationRecord.status == "approved")
                         .order_by(CategorizationRecord.revision.desc()).limit(1)) if s else None
    blocked = ("An approved categorization is never overwritten by an import – this stays a proposal; change it in Categorize (AO re-approval)"
               if approved else None)
    fs = {}
    for k, (fld, lab) in enumerate((("confidentiality_impact", "Confidentiality"), ("integrity_impact", "Integrity"), ("availability_impact", "Availability"))):
        rows = [x["row"] for x in p["selected"] if x["impacts"][k] == p["impacts"][k]]
        fs[fld] = _f(getattr(s, fld, None) if s else None, p["impacts"][k], None if p["impacts"][k] else "No Low/Moderate/High value for the applicable mission types",
                     blocked, lab, p["sheet"], rows[0] if rows else None, 3 + k)
    mt = [x["mission_type"] for x in p["selected"]]
    cur_mt = ((s.environment_detail or {}).get("emass_mission_types") if s else None)
    fs["emass_mission_types"] = _f(cur_mt, mt, None, blocked, "eMASS Mission Types (recorded; SP 800-60 information types are chosen in Categorize)",
                                  p["sheet"], p["selected"][0]["row"], 0)
    item = {"section": "Categorization", "key": "categorization", "label": "Categorization proposal", "create": False, "fields": fs,
            "required_issue": None, "note": "Overall = high-water mark of the applicable mission types. The baseline is not changed here – re-select it in Select."}
    return [item], {**rep, "matched": 1}


def apply_categorization(db: Session, s: System, i: dict, acc: dict, current, meta: dict, src: dict) -> dict:
    from app.services import audit
    from app.services.systems import high_water_mark
    before = {k: ei._j(getattr(s, k)) for k in ("confidentiality_impact", "integrity_impact", "availability_impact", "categorization")}
    before["emass_mission_types"] = (s.environment_detail or {}).get("emass_mission_types")
    for k, fd in acc.items():
        if k == "emass_mission_types":
            s.environment_detail = {**(s.environment_detail or {}), "emass_mission_types": fd["imported"]}
        else:
            setattr(s, k, fd["imported"])
    s.categorization = high_water_mark(s.confidentiality_impact, s.integrity_impact, s.availability_impact)
    audit.record(db, action="system.categorization_imported", actor_user_id=current.user.id, actor_email=current.user.email, entity_type="system",
                 entity_id=str(s.id), before=before, after={k: fd["imported"] for k, fd in acc.items()},
                 details={**src, "provenance": {k: fd.get("src") for k, fd in acc.items()}}, **meta)
    return {"control": "Categorization", "key": i["key"], "applied": sorted(acc), "failed": None, "undo": {"kind": "system", "before": before}}


# ---------- PPSM ----------
def _port_ok(v: str) -> bool:
    if _n(v) in ("any", "all"):
        return True
    m = re.fullmatch(r"(\d{1,5})(\s*-\s*(\d{1,5}))?", v)
    return bool(m) and all(0 <= int(x) <= 65535 for x in (m.group(1), m.group(3) or m.group(1))) and int(m.group(1)) <= int(m.group(3) or m.group(1))


def ppsm_key(asset_id, port, protocol, service) -> str:
    return f"asset:{asset_id}" if asset_id else f"{_n(port)}/{_n(protocol)}/{_n(service)}"[:200]


def build_ppsm(db: Session, s: System, sheets) -> tuple[list[dict], dict]:
    rep: dict = {"invalid": [], "duplicates": [], "exceptions": [],
                 "not_imported": ["Information System Owner / POC Name / POC Phone / POC E-Mail – person data is never imported"]}
    found = _pick_sheet(sheets, PPSM_COLS, "port", ("ppsm",))
    ai = None
    if found:
        title, rows, hi, cols = found
    elif (got := _ai_header(db, s, "ppsm", sheets, rep)):
        title, rows, hi, cols, ai = got
    else:
        raise HTTPException(422, "No PPSM header row (Port / Protocol / Data Service / Boundary) found")
    rep.update(sheet=title, header_row=hi + 1)
    existing = {e.record_key: e for e in db.scalars(select(PpsmEntry).where(PpsmEntry.system_id == s.id))} if s else {}
    by_asset = {e.asset_id: e for e in existing.values() if e.asset_id}
    items, seen = [], {}
    headers = {_n(h) for k, al in PPSM_COLS.items() for h in al}
    for n, row in enumerate(rows[hi + 1:], start=hi + 2):
        raw = {k: clean(row[c]) if c < len(row) else None for k, c in cols.items()}
        if not any(v for k, v in raw.items() if k != "asset_id") or any(_n(v) in headers for v in raw.values() if v):
            continue
        aid = raw.get("asset_id") if raw.get("asset_id") and re.fullmatch(r"\d{1,20}", raw["asset_id"]) else None
        key = ppsm_key(aid, raw.get("port"), raw.get("protocol"), raw.get("data_service"))
        if key in seen:
            rep["duplicates"].append({"row": n, "control": key, "first_row": seen[key]})
            continue
        seen[key] = n
        e = by_asset.get(aid) if aid else None
        e = e or existing.get(ppsm_key(None, raw.get("port"), raw.get("protocol"), raw.get("data_service")))
        fs, issues = {}, []
        for k in ("port", "protocol", "data_service", "boundary", "fqdn", "purpose"):
            v, err = raw.get(k), None
            if v is None:
                if k in PPSM_REQUIRED:
                    issues.append(f"{PPSM_LABEL[k]} is required")
                continue
            if k == "boundary":
                parts = [p.strip() for p in v.split(";") if p.strip()]
                norm = [_BND.get(_n(re.sub(r"^\s*\d+\.\s*", "", p))) or _BND.get(re.match(r"\s*(\d+)", p).group(1) if re.match(r"\s*(\d+)", p) else "") for p in parts]
                bad = [p for p, x in zip(parts, norm) if not x]
                v, err = (sorted(set(x for x in norm if x), key=BOUNDARIES.index), None) if not bad else (v, f"Not a PPSM boundary (see Glossary): {', '.join(bad)}")
            elif len(v) > PPSM_LIMIT[k]:
                err = f"Longer than {PPSM_LIMIT[k]} characters (eMASS limit)"
            elif k == "port" and not _port_ok(v):
                err = f"'{v}' is not a port or port range (0–65535)"
            elif k == "protocol":
                v = v.upper()
            cur = None
            if e is not None:
                cur = e.boundaries if k == "boundary" else getattr(e, k)
            fs[k] = _f(cur, v, err, None, PPSM_LABEL[k], title, n, cols[k])
            if ai:
                fs[k]["ai"] = ai.get(k)
            if err:
                rep["invalid"].append({"row": n, "control": key, "field": PPSM_LABEL[k], "value": str(v)[:100], "error": err})
                issues.append(err)
        if e is None:  # a new row: every present field is a change
            for fd in fs.values():
                fd["change"] = not fd["error"]
        label = f"{raw.get('port') or '?'}/{(raw.get('protocol') or '?').upper()} {raw.get('data_service') or ''}".strip()
        items.append({"section": "PPSM", "key": "ppsm:" + key, "record_key": key, "asset_id": aid, "label": label, "create": e is None,
                      "ppsm_id": str(e.id) if e else None, "row": n, "fields": fs, "ai_mapped": bool(ai),
                      "required_issue": "; ".join(issues) if issues and e is None else None})
    rep["matched"] = len(items)
    rep["new"] = sum(1 for i in items if i["create"])
    return items, rep


def _ppsm_vals(e: PpsmEntry) -> dict:
    return {"port": e.port, "protocol": e.protocol, "data_service": e.data_service, "boundaries": list(e.boundaries or []), "fqdn": e.fqdn,
            "purpose": e.purpose, "asset_id": e.asset_id, "record_key": e.record_key}


def apply_ppsm(db: Session, s: System, i: dict, acc: dict, current, meta: dict, src: dict, import_id) -> dict:
    from app.services import audit
    vals = {("boundaries" if k == "boundary" else k): fd["imported"] for k, fd in acc.items()}
    e = db.get(PpsmEntry, i["ppsm_id"]) if i.get("ppsm_id") else db.scalar(select(PpsmEntry).where(PpsmEntry.system_id == s.id, PpsmEntry.record_key == i["record_key"]))
    if e is None:
        if i.get("required_issue") or not {"port", "protocol", "data_service", "boundaries"} <= set(vals):
            return {"control": i["label"], "key": i["key"], "applied": [], "failed": i.get("required_issue") or "Accept Port, Protocol, Data Service and Boundary to create a PPSM row"}
        e = PpsmEntry(system_id=s.id, record_key=i["record_key"], asset_id=i.get("asset_id"), source="emass_import", **vals)
        db.add(e)
        before = None
    else:
        before = _ppsm_vals(e)
        for k, v in vals.items():
            setattr(e, k, v)
    e.import_id, e.source_row, e.updated_by = import_id, i.get("row"), current.user.id
    db.flush()
    audit.record(db, action="ppsm.created" if before is None else "ppsm.updated", actor_user_id=current.user.id, actor_email=current.user.email,
                 entity_type="ppsm_entry", entity_id=str(e.id), before=before, after={k: ei._j(v) for k, v in vals.items()},
                 details={**src, "row": i.get("row"), "provenance": {k: fd.get("src") for k, fd in acc.items()}}, **meta)
    return {"control": i["label"], "key": i["key"], "applied": sorted(acc), "failed": None, "created_id": str(e.id) if before is None else None,
            "undo": {"kind": "ppsm", "id": str(e.id), "before": before}}


def ppsm_list(db: Session, s: System) -> list[dict]:
    return [{"id": str(e.id), **_ppsm_vals(e), "source": e.source, "import_id": str(e.import_id) if e.import_id else None, "source_row": e.source_row}
            for e in db.scalars(select(PpsmEntry).where(PpsmEntry.system_id == s.id).order_by(PpsmEntry.port, PpsmEntry.protocol))]


def ppsm_export(db: Session, s: System) -> bytes:
    """Fill the eMASS PPSM Boundaries template (data validation, glossary and labels kept; macros none)."""
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO((FIXTURES / "ppsm_boundaries_template.xlsx").read_bytes()))
    ws = wb["PPSM"]
    ws["C2"] = date.today().isoformat()  # labels sit in merged A:B, values in merged C:D
    ws["C5"] = safe_cell(s.name)
    for n, e in enumerate(ppsm_list(db, s)):
        r = 9 + n
        for c, v in enumerate([e["asset_id"], e["port"], e["protocol"], e["data_service"], "; ".join(e["boundaries"]), e["fqdn"], e["purpose"]], start=1):
            ws.cell(row=r, column=c, value=safe_cell(v))
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


# ---------- dispatch / undo ----------
BUILD = {"slcm": build_slcm, "categorization": build_categorization, "ppsm": build_ppsm}


def build(db: Session, s: System, typ: str, sheets) -> tuple[list[dict], dict]:
    items, rep = BUILD[typ](db, s, sheets)
    preselect(items)
    return items, rep


def apply_item(db, imp, s, i, acc, current, meta, src) -> dict:
    if i["section"] == "SLCM":
        return apply_slcm(db, s, i, acc, current, meta, src)
    if i["section"] == "PPSM":
        return apply_ppsm(db, s, i, acc, current, meta, src, imp.id)
    return apply_categorization(db, s, i, acc, current, meta, src)


def revert(db: Session, imp, current, meta: dict, keys: list[str] | None) -> list[dict]:
    """Undo applied changes of this session (all, or the given item keys) – before values restored, each step audited."""
    from app.services import audit
    s = db.get(System, imp.system_id)
    done_keys = set((imp.report or {}).get("reverted_keys") or [])
    out = []
    for r in (imp.report or {}).get("apply_results") or []:
        u = r.get("undo")
        if not u or r.get("key") in done_keys or (keys and r.get("key") not in keys):
            continue
        if u["kind"] == "slcm":
            f = db.scalar(select(EmassControlFields).where(EmassControlFields.assignment_id == u["assignment_id"]))
            if f is None:
                continue
            after = {k: getattr(f, k) for k in u["before"]}
            for k, v in u["before"].items():
                setattr(f, k, v)
            f.slcm_inherit = u["inherit_before"] or None
            ent = ("control_assignment", u["assignment_id"])
        elif u["kind"] == "ppsm":
            e = db.get(PpsmEntry, u["id"])
            if e is None:
                continue
            after = _ppsm_vals(e)
            if u["before"] is None:
                db.delete(e)
            else:
                for k, v in u["before"].items():
                    setattr(e, k, v)
            ent = ("ppsm_entry", u["id"])
        else:
            after = {k: ei._j(getattr(s, k)) for k in ("confidentiality_impact", "integrity_impact", "availability_impact", "categorization")}
            for k, v in u["before"].items():
                if k == "emass_mission_types":
                    det = dict(s.environment_detail or {})
                    det.pop("emass_mission_types", None) if v is None else det.update(emass_mission_types=v)
                    s.environment_detail = det or None
                else:
                    setattr(s, k, v)
            ent = ("system", str(s.id))
        audit.record(db, action="emass.import_reverted", actor_user_id=current.user.id, actor_email=current.user.email, entity_type=ent[0],
                     entity_id=ent[1], before={k: ei._j(v) for k, v in after.items()}, after={k: ei._j(v) for k, v in (u["before"] or {}).items()},
                     details={"import_id": str(imp.id), "import_sha256": imp.sha256, "item": r.get("key")}, **meta)
        done_keys.add(r.get("key"))
        out.append({"key": r.get("key"), "control": r.get("control")})
    all_keys = {r.get("key") for r in (imp.report or {}).get("apply_results") or [] if r.get("undo")}
    imp.report = {**(imp.report or {}), "reverted_keys": sorted(k for k in done_keys if k)}
    if all_keys and all_keys <= done_keys:
        imp.status = "reverted"
    return out


# ---------- copy setup from an existing system ----------
def setup_template(s: System) -> dict:
    """Registration values copied from another system (categorization inputs → same baseline); nothing about assessments or evidence."""
    return {"categorization_mode": s.categorization_mode, "confidentiality_impact": s.confidentiality_impact, "integrity_impact": s.integrity_impact,
            "availability_impact": s.availability_impact, "baseline_key": s.baseline_key, "information_types": s.information_types,
            "nss_determination": s.nss_determination, "operating_environment": s.operating_environment, "deployment_type": s.deployment_type,
            "overlay_keys": s.overlay_keys or [], "source": {"id": str(s.id), "name": s.name, "acronym": s.acronym}}


def copy_setup(db: Session, src: System, dst: System, actor) -> dict:
    from app.services import audit
    from app.services import rmf_steps as rs
    tag = f"copied from {src.acronym or src.name}"
    out: dict = {"overlays": [], "tailoring": [], "odps": 0, "iscm": None, "policies": 0, "skipped": []}
    for key in src.overlay_keys or []:
        if key in (dst.overlay_keys or []):
            continue
        try:
            rs.set_overlay(db, dst, key, True, actor)
            out["overlays"].append(key)
        except ValueError as exc:
            out["skipped"].append(f"Overlay {key}: {exc}")
    for t in db.scalars(select(TailoringDecision).where(TailoringDecision.system_id == src.id).order_by(TailoringDecision.created_at)):
        if t.overlay_key:
            continue
        try:
            r = rs.tailor(db, dst, t.action, [t.control_id], f"{t.justification} ({tag})", actor, t.applicability, t.compensating_control_id)
            out["tailoring"] += [f"{t.action} {ei.fmt_acronym(c)}" for c in r["controls"]]
        except ValueError as exc:
            out["skipped"].append(f"Tailoring {t.action} {ei.fmt_acronym(t.control_id)}: {exc}")
    have = {(o.control_id, o.parameter) for o in db.scalars(select(OdpValue).where(OdpValue.system_id == dst.id))}
    dst_ctl = set(db.scalars(select(ControlAssignment.control_id).where(ControlAssignment.system_id == dst.id)))
    for o in db.scalars(select(OdpValue).where(OdpValue.system_id == src.id)):
        if (o.control_id, o.parameter) in have or o.control_id not in dst_ctl:
            continue
        db.add(OdpValue(system_id=dst.id, control_id=o.control_id, parameter=o.parameter, value=o.value, numeric_value=o.numeric_value,
                        unit=o.unit, source=f"{o.source} ({tag})"[:255]))  # approvals are never copied
        out["odps"] += 1
    st = db.scalar(select(IscmStrategy).where(IscmStrategy.system_id == src.id))
    if st and db.scalar(select(IscmStrategy).where(IscmStrategy.system_id == dst.id)) is None:
        nst = IscmStrategy(system_id=dst.id, status="draft", version=1, narrative=st.narrative, scope=st.scope, defaults=dict(st.defaults or {}),
                           reporting=list(st.reporting or []), review_activities=list(st.review_activities or []),
                           evidence_requirements=list(st.evidence_requirements or []))
        db.add(nst)
        db.flush()
        for p in db.scalars(select(IscmProfile).where(IscmProfile.strategy_id == st.id)):
            db.add(IscmProfile(strategy_id=nst.id, name=p.name, match_type=p.match_type, match_value=p.match_value, values=dict(p.values or {}),
                               sort_order=p.sort_order))
        out["iscm"] = "draft strategy + profiles copied (needs its own approval)"
    for p in db.scalars(select(Policy).where(Policy.org_id == dst.org_id)):
        ids = [str(x) for x in (p.system_ids or [])]
        if str(src.id) in ids and str(dst.id) not in ids:
            p.system_ids = ids + [str(dst.id)]
            out["policies"] += 1
    audit.record(db, action="system.setup_copied", actor_user_id=actor.id, actor_email=actor.email, entity_type="system", entity_id=str(dst.id),
                 after=out, details={"source_system_id": str(src.id), "copied": ["overlays", "tailoring", "odps", "iscm", "policy links"],
                                     "never_copied": ["assessments", "evidence", "POA&Ms", "test results", "approvals", "inventory"]})
    return out


def sample(typ: str) -> bytes:
    """SYNTHETIC samples on the real eMASS layouts."""
    import openpyxl
    if typ == "slcm":
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.title = "Final Results"
        ws.append(["Implementation Status / Security Control Status"])
        ws.append(["Sort #", "Control Acronym", "Control Title", "Criticality", "Frequency", "Method", "Reporting", "Tracking", "SLCM Comments", "Order #"])
        for n, (c, crit, fr, me) in enumerate([("AC-2", "CRWG Yellow Criticality Control", "Quarterly", "Automated"), ("AC-2(1)", "CRWG White Criticality Control", "Annually", "Manual"),
                                               ("AU-6", "CRWG Red Criticality Control", "Weekly", "Automated"), ("ZZ-9", "---", "Hourly", "Manual")], start=1):
            ws.append([n, c, "SYNTHETIC", crit, fr, me, "SYNTHETIC: ISSO reports to the ISSM", "POA&M", "SYNTHETIC sample", n])
        b = io.BytesIO()
        wb.save(b)
        return b.getvalue()
    wb = openpyxl.load_workbook(io.BytesIO((FIXTURES / ("categorization_template.xlsx" if typ == "categorization" else "ppsm_boundaries_template.xlsx")).read_bytes()))
    if typ == "categorization":
        ws = wb["Mission Types"]
        ws["B1"] = "SYNTHETIC Logistics Tracking Platform"
        for r in (7, 14):  # Border Control and Transportation (M/M/M), Disaster Preparedness (L/L/L)
            ws.cell(row=r, column=3, value="Yes")
    else:
        ws = wb["PPSM"]
        ws["C5"] = "SYNTHETIC Logistics Tracking Platform"
        for n, row in enumerate([(None, "443", "TCP", "HTTPS", "1. Ext to DoD GW (In); 11. Enclave GW to Enclave (In)", "syn-web.example.test", "SYNTHETIC web front end"),
                                 (None, "1521", "TCP", "Oracle DB", "11. Enclave GW to Enclave (In)", None, "SYNTHETIC database"),
                                 (None, "70000", "UDP", "Bad port", "99. Nowhere", None, "SYNTHETIC invalid row")]):
            for c, v in enumerate(row, start=1):
                ws.cell(row=9 + n, column=c, value=v)
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


def now() -> datetime:
    return datetime.now(timezone.utc)
