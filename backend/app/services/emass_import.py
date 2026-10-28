"""eMASS Control Information import – compatibility pending verification against a real eMASS export.

Parses an eMASS Control Information export (or our own eMASS-formatted Implementation Plan xlsx), detects the header row and
columns BY HEADER NAME (case/whitespace/punctuation-insensitive, with aliases incl. the MITRE emass_client API names), matches
rows to the system's controls by Control Acronym and builds per-field proposals. Nothing is saved until apply(); every applied
field goes through services.emass_fields (single status sync, eMASS required-field rules, Inherited rule, audit).
"""
import hashlib
import io
import re
import zipfile
from datetime import date, datetime, timedelta, timezone

from fastapi import HTTPException
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.models import ControlAssignment, EmassImport, Implementation, System
from app.services import emass_fields as ef

LABEL = "eMASS Control Information import – compatibility pending verification against a real eMASS export"
MAX_BYTES = 10 * 1024 * 1024
MAX_UNCOMPRESSED = 120 * 1024 * 1024
MAX_ROWS = 10_000
MAX_SHEETS = 40
HEADER_SCAN_ROWS = 30
NARRATIVE_MAX = 20_000

IMPORT_FIELDS = ["implementation_status", *ef.PLAN_FIELDS, "implementation_narrative", *ef.RISK_FIELDS]
EXTRA_ALIASES = {
    "acronym": ["control acronym", "control", "acronym", "control id", "security control", "security control acronym", "control identifier", "control number"],
    "implementation_status": ["implementation status", "impl status", "implementationstatus"],
    "common_control_provider": ["common control provider", "ccp"],
    "na_justification": ["n/a justification", "na justification", "not applicable justification"],
    "control_designation": ["security control designation", "control designation", "designation"],
    "estimated_completion_date": ["estimated completion date", "ecd", "estimated completion", "scheduled completion date"],
    "responsible_entities": ["responsible entities", "responsible entity"],
    "implementation_narrative": ["implementation narrative", "narrative"],
    "slcm_criticality": ["slcm criticality", "criticality"],
    "slcm_frequency": ["slcm frequency", "frequency"],
    "slcm_method": ["slcm method", "method"],
    "slcm_reporting": ["slcm reporting", "reporting"],
    "slcm_tracking": ["slcm tracking", "tracking"],
    "slcm_comments": ["slcm comments", "slcm comment"],
    "relevance_of_threat": ["relevance of threat", "threat relevance"],
    "residual_risk_level": ["residual risk level", "residual risk"],
    "impact_description": ["impact description"],
    "vulnerability_summary": ["vulnerability summary"],
}
# known eMASS export columns that are not imported (reported as ignored, not unknown)
IGNORED = ["system name", "system acronym", "control title", "control information", "control family", "compliance status", "ap acronym", "cci",
           "test method", "date tested", "tested by", "test results", "last updated", "last updated by", "security control baseline", "control set",
           "overlays", "inherited from", "inherited", "assessment procedures", "aps"]


def _n(h) -> str:
    return re.sub(r"[^a-z0-9]", "", str(h or "").lower())


def _alias_map() -> dict[str, str]:
    m: dict[str, str] = {}
    for f in IMPORT_FIELDS:
        label, api = ef.FIELD_META[f]
        for a in [label, api or "", f, *EXTRA_ALIASES.get(f, [])]:
            if a:
                m.setdefault(_n(a), f)
    for a in EXTRA_ALIASES["acronym"]:
        m[_n(a)] = "acronym"
    return m


ALIASES = _alias_map()


def norm_acronym(raw) -> str | None:
    """'AC-2 (1)', 'ac-02(1)', 'AC-2(1)' and our internal 'ac-2.1' all become 'ac-2.1'."""
    s = re.sub(r"\s+", "", str(raw or "")).upper()
    m = re.fullmatch(r"([A-Z]{2})-0*(\d+)(?:[(.]0*(\d+)\)?)?", s)
    if not m:
        return None
    return f"{m.group(1).lower()}-{int(m.group(2))}" + (f".{int(m.group(3))}" if m.group(3) else "")


def fmt_acronym(cid: str) -> str:
    base, _, enh = cid.partition(".")
    return base.upper() + (f"({enh})" if enh else "")


def _safe_open(data: bytes):
    import openpyxl
    if len(data) > MAX_BYTES:
        raise HTTPException(413, f"File is larger than {MAX_BYTES // (1024 * 1024)} MB")
    if data[:4] == b"\xd0\xcf\x11\xe0":
        raise HTTPException(422, "Encrypted or legacy .xls workbooks are not accepted – export as .xlsx")
    if data[:2] != b"PK":
        raise HTTPException(422, "Not an .xlsx/.xlsm workbook")
    try:
        z = zipfile.ZipFile(io.BytesIO(data))
        names = z.namelist()
        if "xl/workbook.xml" not in names or "[Content_Types].xml" not in names:
            raise HTTPException(422, "Not an Excel workbook (xl/workbook.xml missing)")
        if sum(i.file_size for i in z.infolist()) > MAX_UNCOMPRESSED or len(names) > 2000:
            raise HTTPException(422, "Workbook expands to an unsafe size – rejected")
        # keep_vba=False: macros (xl/vbaProject.bin) are never loaded or executed
        return openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_vba=False, keep_links=False)
    except HTTPException:
        raise
    except Exception as exc:  # noqa: BLE001 – any parser failure is a clean 422
        raise HTTPException(422, f"Could not read the workbook ({type(exc).__name__})") from exc


def parse(data: bytes, file_name: str = "upload.xlsx") -> dict:
    from app.services.emass_reports import read_sheets
    return parse_sheets(read_sheets(data, file_name))


def parse_sheets(sheets: list) -> dict:
    best = None
    data_sheets = [x for x in sheets if _n(x[0]) not in AUX_SHEETS] or sheets  # our own export: Field mapping / Validation are not data
    for title, rows in data_sheets:
        for i, row in enumerate(rows[:HEADER_SCAN_ROWS], start=1):
            hits = {c: ALIASES[_n(v)] for c, v in enumerate(row) if v is not None and _n(v) in ALIASES}
            if "acronym" in hits.values() and len(hits) >= 2 and (best is None or len(hits) > len(best[3])):
                best = (title, rows, i, hits, row)
    if best is None:
        raise HTTPException(422, "No header row with a Control Acronym column and at least one eMASS field was found in the first "
                                 f"{HEADER_SCAN_ROWS} rows of any sheet")
    title, all_rows, hrow, hits, header = best
    cols: dict[str, int] = {}
    dup_cols = []
    for c, f in hits.items():
        if f in cols:
            dup_cols.append(str(header[c]))
        else:
            cols[f] = c
    ign = {_n(x) for x in IGNORED}
    unknown = [str(v) for c, v in enumerate(header) if v is not None and c not in hits and _n(v) not in ign]
    ignored = [str(v) for c, v in enumerate(header) if v is not None and c not in hits and _n(v) in ign]
    if len(all_rows) - hrow > MAX_ROWS:
        raise HTTPException(422, f"More than {MAX_ROWS:,} data rows – rejected")
    rows = []
    for n, row in enumerate(all_rows[hrow:], start=hrow + 1):
        if not any(v not in (None, "") for v in row):
            continue
        vals = {f: (row[c] if c < len(row) else None) for f, c in cols.items()}
        rows.append({"row": n, "acronym": vals.pop("acronym"), "values": vals})
    return {"sheet": title, "header_row": hrow, "columns": {f: {"header": str(header[c]), "col": _col(c)} for f, c in cols.items()},
            "unknown_headers": unknown, "ignored_headers": ignored, "duplicate_headers": dup_cols, "rows": rows}


AUX_SHEETS = {"fieldmapping", "validation", "readme", "instructions", "readmefirst"}
PLACEHOLDERS = {"", "na", "none", "nil", "tbd", "notapplicable"}  # after _n(): "N/A ", "n/a", "-", "--" …


def _col(i: int) -> str:
    s = ""
    i += 1
    while i:
        i, r = divmod(i - 1, 26)
        s = chr(65 + r) + s
    return s


def _enum(raw: str, opts: list[str], extra: dict | None = None) -> str | None:
    k = _n(raw)
    for o in opts:
        if _n(o) == k:
            return o
    return (extra or {}).get(k)


def normalize(field: str, raw) -> tuple[object, str | None]:
    """→ (value, error). Empty cells are 'no value in file' (never a clear)."""
    if raw is None or (isinstance(raw, str) and not raw.strip()):
        return None, None
    if field == "implementation_status":
        v = _enum(str(raw), ef.STATUSES, {"na": "Not Applicable", "notapplicable": "Not Applicable"})
        return (v, None) if v else (None, f"'{raw}' is not an eMASS Implementation Status ({', '.join(ef.STATUSES)})")
    if field in ef.ENUMS:
        v = _enum(str(raw), ef.ENUMS[field])
        return (v, None) if v else (None, f"'{raw}' is not an eMASS value ({', '.join(ef.ENUMS[field])})")
    if field == "estimated_completion_date":
        return _date(raw)
    v = str(raw).strip()
    lim = NARRATIVE_MAX if field == "implementation_narrative" else ef.LIMITS.get(field, 2000)
    if len(v) > lim:
        return None, f"{len(v):,} characters – over the eMASS limit of {lim:,}"
    return v, None


def _date(raw) -> tuple[date | None, str | None]:
    if isinstance(raw, datetime):
        return raw.date(), None
    if isinstance(raw, date):
        return raw, None
    if isinstance(raw, (int, float)) and 20000 < raw < 80000:
        return (datetime(1899, 12, 30) + timedelta(days=float(raw))).date(), None
    s = str(raw).strip()
    for fmt in ("%Y-%m-%d", "%m/%d/%Y", "%d-%b-%Y", "%d %b %Y", "%Y/%m/%d", "%m/%d/%y"):
        try:
            return datetime.strptime(s, fmt).date(), None
        except ValueError:
            pass
    return None, f"'{raw}' is not a date (use YYYY-MM-DD or MM/DD/YYYY)"


def _j(v):
    return v.isoformat() if isinstance(v, date) else v


def _current(db: Session, a: ControlAssignment) -> dict:
    f = ef.fields_for(db, a)
    narr = db.scalar(select(Implementation.narrative).where(Implementation.assignment_id == a.id, Implementation.is_current.is_(True)))
    return {"implementation_status": ef.emass_status(a, f), **{k: _j(v) for k, v in ef.effective(db, a, f)[0].items()}, "implementation_narrative": narr}


def build(db: Session, s: System, parsed: dict) -> tuple[list[dict], dict]:
    assigned = {a.control_id: a for a in db.scalars(select(ControlAssignment).where(ControlAssignment.system_id == s.id))}
    seen: dict[str, int] = {}
    items, rep = [], {"not_in_system": [], "unparseable": [], "duplicates": [], "invalid": [], "blocked": [], "required_issues": [],
                      "warnings": [], "info": [], "status_not_applied": []}
    for r in parsed["rows"]:
        cid = norm_acronym(r["acronym"])
        if not cid:
            rep["unparseable"].append({"row": r["row"], "value": str(r["acronym"])[:80]})
            continue
        if cid in seen:
            rep["duplicates"].append({"row": r["row"], "control": fmt_acronym(cid), "first_row": seen[cid]})
            continue
        seen[cid] = r["row"]
        a = assigned.get(cid)
        if a is None:
            rep["not_in_system"].append({"row": r["row"], "control": fmt_acronym(cid)})
            continue
        cur = _current(db, a)
        inherited_now = cur["implementation_status"] == "Inherited"
        fields = {}
        for f in IMPORT_FIELDS:
            if f not in r["values"]:
                continue
            v, err = normalize(f, r["values"][f])
            if v is None and err is None:
                continue
            v = _j(v)
            blocked = None
            if not err:
                if f == "implementation_status" and v == "Inherited" and not ef.has_inheritance(a):
                    blocked = "Inherited requires an inheritance record (providing system / common control provider) – none recorded; use Manually Inherited"
                elif f in ef.RISK_FIELDS and (inherited_now or (r["values"].get("implementation_status") and normalize("implementation_status", r["values"]["implementation_status"])[0] == "Inherited")):
                    blocked = "Risk Assessment cannot be changed for an Inherited control (eMASS rule)"
                elif f in ef.PLAN_FIELDS and inherited_now and str(cur.get(f) or "") != str(v):
                    blocked = "Inherited control – Implementation Plan values come from the providing system"
            fields[f] = {"current": cur.get(f), "imported": v if not err else str(r["values"][f])[:200], "change": not err and str(cur.get(f) or "") != str(v),
                         "error": err, "blocked": blocked, "decision": None}
            if err:
                rep["invalid"].append({"row": r["row"], "control": fmt_acronym(cid), "field": ef.FIELD_META[f][0], "value": str(r["values"][f])[:120], "error": err})
            if blocked:
                rep["blocked"].append({"row": r["row"], "control": fmt_acronym(cid), "field": ef.FIELD_META[f][0], "reason": blocked})
        _row_rules(fields, cur, r["row"], fmt_acronym(cid), rep)
        item = {"control_id": cid, "control": fmt_acronym(cid), "row": r["row"], "fields": fields, "required_issue": None}
        item["required_issue"] = required_issue(item, cur)
        if item["required_issue"]:
            rep["required_issues"].append({"row": r["row"], "control": item["control"], "issue": item["required_issue"]})
        items.append(item)
    rep["matched"] = len(items)
    rep["missing_from_file"] = sorted(fmt_acronym(c) for c in assigned if c not in seen)
    rep["rows_read"] = len(parsed["rows"])
    rep["with_changes"] = sum(1 for i in items if any(f["change"] and not f["error"] and not f["blocked"] for f in i["fields"].values()))
    rows_of = lambda *ks: {x["row"] for k in ks for x in rep[k] if x.get("row")}  # noqa: E731
    err = rows_of("invalid", "unparseable", "duplicates")
    rep["summary"] = {"rows_read": rep["rows_read"], "rows_to_apply": rep["with_changes"], "rows_not_in_system": len(rep["not_in_system"]),
                      "rows_with_warnings": len(rows_of("warnings", "status_not_applied", "required_issues", "blocked") - err),
                      "rows_with_errors": len(err), "placeholders_ignored": len(rep["info"])}
    for i in items:  # imported values win by default: valid, unblocked changes are preselected "Use imported"
        for f in i["fields"].values():
            if _ok(f) and f["decision"] is None:
                f["decision"] = "accept"
    return items, rep


def _row_rules(fields: dict, cur: dict, row: int, ctl: str, rep: dict) -> None:
    """Per-row eMASS rules – never a whole-import failure: fields that don't apply to the row's status are dropped (placeholders → info,
    real values → warning); a status whose required fields are missing is not applied (row warning); everything else stays importable."""
    sf = fields.get("implementation_status")
    st = sf["imported"] if sf and _ok(sf) else cur["implementation_status"]
    for k, applies, where in (("na_justification", st == "Not Applicable", "Not Applicable"),
                              ("common_control_provider", st in ("Inherited", "Manually Inherited"), "Inherited / Manually Inherited")):
        f = fields.get(k)
        if f is None or applies:
            continue
        lab, v = ef.FIELD_META[k][0], str(f["imported"] or "").strip()
        if f["error"]:  # enum error on a field that doesn't apply to this status → not an error, just not applied
            rep["invalid"] = [x for x in rep["invalid"] if not (x["row"] == row and x["field"] == lab)]
        if _n(v) in PLACEHOLDERS:
            rep["info"].append({"row": row, "control": ctl, "field": lab, "value": v, "note": f"Placeholder ignored – {lab} does not apply to {st}"})
        else:
            rep["warnings"].append({"row": row, "control": ctl, "field": lab, "value": v[:120],
                                    "reason": f"{lab} only applies to {where} controls – not applied (row is {st})"})
        del fields[k]
    if sf and _ok(sf) and sf["imported"] != "Inherited":
        vals = {k: (fields[k]["imported"] if k in fields and _ok(fields[k]) else cur.get(k)) for k in ef.PLAN_FIELDS}
        miss = ef.missing_required(sf["imported"], vals)
        if miss:
            labs = [ef.FIELD_META[k][0] for k in miss]
            how = "add them to the file" + (", or set the SLCM fields via the system's Continuous Monitoring Strategy" if set(miss) & set(ef.SLCM) else "")
            sf["blocked"] = f"{sf['imported']} not applied – missing: {', '.join(labs)} ({how}). The other values of this row are still imported."
            rep["status_not_applied"].append({"row": row, "control": ctl, "status": sf["imported"], "missing": labs})
            rep["warnings"].append({"row": row, "control": ctl, "field": "Implementation Status", "value": sf["imported"], "reason": sf["blocked"]})


def _ok(f: dict) -> bool:
    return f["change"] and not f["error"] and not f["blocked"]


def required_issue(item: dict, cur: dict, only_accepted: bool = False) -> str | None:
    """Would the resulting Implementation Plan violate the eMASS required-field rules for its status?"""
    take = (lambda f: _ok(f) and f["decision"] == "accept") if only_accepted else _ok
    fs = item["fields"]
    touched = [k for k in ["implementation_status", *ef.PLAN_FIELDS] if k in fs and take(fs[k])]
    if not touched:
        return None
    st = fs["implementation_status"]["imported"] if "implementation_status" in touched else cur["implementation_status"]
    if st == "Inherited":
        return None
    vals = {k: (fs[k]["imported"] if k in touched else cur.get(k)) for k in ef.PLAN_FIELDS}
    miss = ef.missing_required(st, vals)
    return (f"Still incomplete for eMASS (Implementation Status {st}) – missing: {', '.join(ef.FIELD_META[k][0] for k in miss)}. "
            "The imported values are saved; the row stays flagged in the Workbench and the eMASS validation report.") if miss else None


def store(db: Session, org_id, system: System | None, file_name: str, data: bytes, user, report_type: str, detected: dict) -> EmassImport:
    from app.services.storage.base import get_storage
    sha = hashlib.sha256(data).hexdigest()
    safe = re.sub(r"[^A-Za-z0-9._ -]", "_", file_name)[:200]
    key = f"emass-import/{org_id}/{system.id if system else 'unassigned'}/{sha}/{safe}"
    get_storage().put(key, data, "text/csv" if safe.lower().endswith(".csv") else "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    imp = EmassImport(org_id=org_id, system_id=system.id if system else None, file_name=safe, sha256=sha, size_bytes=len(data), storage_key=key,
                      status="detected", uploaded_by=user.id, uploaded_by_email=user.email, report_type=report_type or "control_information",
                      detected=detected, items=[], counts=tally([]))
    db.add(imp)
    db.flush()
    return imp


def target(db: Session, imp: EmassImport, s: System, report_type: str, data: bytes) -> None:
    """Build the review items for one system (nothing is saved to the system)."""
    from app.services import emass_reports as er
    items, rep = er.build(db, s, report_type, data, imp.file_name)
    det = {"label": er.label(report_type)}
    if report_type == "control_information":
        parsed = parse(data, imp.file_name)
        rows = parsed.pop("rows")
        det.update(parsed, data_rows=len(rows))
    prev = db.scalar(select(EmassImport.id).where(EmassImport.org_id == s.org_id, EmassImport.system_id == s.id, EmassImport.sha256 == imp.sha256,
                                                  EmassImport.status == "applied", EmassImport.id != imp.id).limit(1))
    if prev:  # idempotent retry: same file already applied – current values already match, so nothing new is proposed
        rep = {**rep, "duplicate_of": str(prev)}
    imp.system_id, imp.report_type, imp.status = s.id, report_type, "review"
    imp.detection, imp.report, imp.items, imp.counts = det, rep, items, tally(items)


def create(db: Session, s: System, file_name: str, data: bytes, user, report_type: str | None = None) -> EmassImport:
    """Per-system upload – routed by the detected type through the same dispatcher as the global and Import Center uploads."""
    from app.services import emass_reports as er
    sheets = er.read_sheets(data, file_name)
    det = {**er.detect(sheets), "identity": er.identity(sheets), "context_system_id": str(s.id), "context_system": s.acronym or s.name}
    typ = er.route_type(det, report_type)
    imp = store(db, s.org_id, s, file_name, data, user, typ, det)
    target(db, imp, s, typ, data)
    return imp


def tally(items: list[dict]) -> dict:
    c = {"controls": len(items), "changed_fields": 0, "accepted": 0, "rejected": 0, "undecided": 0, "invalid": 0, "blocked": 0}
    for i in items:
        for f in i["fields"].values():
            c["invalid"] += bool(f["error"])
            c["blocked"] += bool(f["blocked"])
            if _ok(f):
                c["changed_fields"] += 1
                c[{"accept": "accepted", "reject": "rejected"}.get(f["decision"], "undecided")] += 1
    return c


def decide(db: Session, imp: EmassImport, decisions: list[dict], bulk: str | None, allowed: set | None = None, section: str | None = None) -> None:
    if imp.status != "review":
        raise HTTPException(409, f"This import is {imp.status} – decisions can no longer change")
    items = [{**i, "fields": {k: dict(v) for k, v in i["fields"].items()}} for i in imp.items or []]
    by = {str(i.get("key") or i.get("control_id")).lower(): i for i in items}
    if bulk == "accept_safe":  # accept every changed, valid, unblocked field (required-field gaps are row warnings, not blockers)
        for i in items:
            if allowed is None or i.get("section", "Controls") in allowed:
                for f in i["fields"].values():
                    if _ok(f) and f["decision"] is None:
                        f["decision"] = "accept"
    elif bulk in ("keep_existing", "use_imported"):  # global or per-section "keep existing" / "use imported"
        for i in items:
            if (section is None or i.get("section", "Controls") == section) and (allowed is None or i.get("section", "Controls") in allowed):
                for f in i["fields"].values():
                    if _ok(f):
                        f["decision"] = "reject" if bulk == "keep_existing" else "accept"
    elif bulk in ("reject_all", "reset"):
        for i in items:
            for f in i["fields"].values():
                if _ok(f):
                    f["decision"] = "reject" if bulk == "reject_all" else None
    for d in decisions:
        i = by.get(str(d.get("key") or d.get("control_id") or "").lower())
        if i is None:
            raise HTTPException(422, f"{d.get('control_id')} is not in this import")
        keys = [k for k, f in i["fields"].items() if _ok(f)] if d.get("field") in (None, "*") else [d["field"]]
        for k in keys:
            f = i["fields"].get(k)
            if f is None:
                raise HTTPException(422, f"{k} is not in this import for {i.get('label')}")
            if d.get("decision") == "accept" and not _ok(f):
                raise HTTPException(422, f"{i.get('label')} {f.get('label') or ef.FIELD_META.get(k, (k,))[0]}: " + (f["error"] or f["blocked"] or "no change to accept"))
            if d.get("decision") == "accept" and allowed is not None and i.get("section", "Controls") not in allowed:
                raise HTTPException(403, f"You cannot apply {i.get('section')} changes ({i.get('section')} needs a different role)")
            f["decision"] = d.get("decision")
    imp.items = items
    imp.counts = tally(items)
    if imp.report_type == "slcm" and imp.system_id:
        from app.services import emass_migration as em
        imp.report = {**(imp.report or {}), "strategy_plan": em.plan_public(em.slcm_plan(db, db.get(System, imp.system_id), items))}


def apply(db: Session, imp: EmassImport, current, meta: dict) -> dict:
    if imp.status != "review":
        raise HTTPException(409, f"This import is already {imp.status}")
    from app.services import emass_reports as er
    src = {"source": f"eMASS import {er.TYPES.get(imp.report_type, imp.report_type)} {imp.file_name} (sha256 {imp.sha256[:12]}…)", "import_id": str(imp.id),
           "import_sha256": imp.sha256, "import_file": imp.file_name, "report_type": imp.report_type}
    results = er.apply_other(db, imp, db.get(System, imp.system_id), current, meta, src)
    for i in imp.items or []:
        acc = {k: f for k, f in i["fields"].items() if _ok(f) and f["decision"] == "accept"}
        if not acc or i.get("section", "Controls") != "Controls":
            continue
        a = db.scalar(select(ControlAssignment).where(ControlAssignment.system_id == imp.system_id, ControlAssignment.control_id == i["control_id"]))
        if a is None:
            results.append({"control": i["control"], "applied": [], "failed": "control no longer in the system"})
            continue
        cur = _current(db, a)
        done, failed = [], []
        plan_keys = [k for k in acc if k == "implementation_status" or k in ef.PLAN_FIELDS]
        if plan_keys:
            st = acc["implementation_status"]["imported"] if "implementation_status" in acc else cur["implementation_status"]
            vals = {k: (acc[k]["imported"] if k in acc else cur.get(k)) for k in ef.PLAN_FIELDS}
            vals["estimated_completion_date"] = date.fromisoformat(vals["estimated_completion_date"]) if vals.get("estimated_completion_date") else None
            try:
                with db.begin_nested():
                    ef.save_plan(db, a, st, vals, current, {"source": "emass_import", "fields": [k for k in plan_keys if k != "implementation_status"]}, meta, src,
                                 enforce_required=False)  # status changes were already gated per row in build (_row_rules)
                done += plan_keys
            except HTTPException as exc:
                failed.append(f"Implementation Plan not saved: {exc.detail}")
        if "implementation_narrative" in acc:
            latest = db.scalar(select(func.max(Implementation.version)).where(Implementation.assignment_id == a.id)) or 0
            for old in db.scalars(select(Implementation).where(Implementation.assignment_id == a.id, Implementation.is_current.is_(True))):
                old.is_current = False
            db.add(Implementation(assignment_id=a.id, version=latest + 1, narrative=acc["implementation_narrative"]["imported"],
                                  implementation_status=a.implementation_status, author_id=current.user.id, review_status="draft"))
            from app.services import audit
            audit.record(db, action="implementation.saved", actor_user_id=current.user.id, actor_email=current.user.email, entity_type="control_assignment",
                         entity_id=str(a.id), before={"version": latest, "narrative": (cur.get("implementation_narrative") or "")[:500]},
                         after={"version": latest + 1, "narrative": acc["implementation_narrative"]["imported"][:500], "review_status": "draft"}, details=src, **meta)
            done.append("implementation_narrative")
        risk_keys = [k for k in acc if k in ef.RISK_FIELDS]
        if risk_keys:
            try:
                with db.begin_nested():
                    ef.save_risk(db, a, {k: acc[k]["imported"] for k in risk_keys}, current, {"source": "emass_import", "fields": risk_keys}, meta, src)
                done += risk_keys
            except HTTPException as exc:
                failed.append(f"Risk Assessment not saved: {exc.detail}")
        results.append({"control": i["control"], "key": i.get("key") or i["control_id"], "applied": done, "failed": "; ".join(failed) or None,
                        "warning": i.get("required_issue"),
                        "undo": {"kind": "control_info", "assignment_id": str(a.id), "before": {k: cur.get(k) for k in done},
                                 "imported": {k: acc[k]["imported"] for k in done}} if done else None})
    imp.status = "applied"
    imp.applied_by_email = current.user.email
    imp.finished_at = datetime.now(timezone.utc)
    imp.counts = {**tally(imp.items or []), "applied_fields": sum(len(r["applied"]) for r in results),
                  "failed_controls": sum(1 for r in results if r["failed"])}
    imp.report = {**(imp.report or {}), "apply_results": results}
    return {"results": results, "counts": imp.counts}


SAMPLE_HEADERS = ["Control Acronym", "Control Title", "Implementation Status", "Security Control Designation", "Estimated Completion Date",
                  "Responsible Entities", "SLCM Criticality", "SLCM Frequency", "SLCM Method", "SLCM Reporting", "SLCM Tracking", "SLCM Comments",
                  "Common Control Provider", "N/A Justification", "Severity", "Relevance of Threat", "Likelihood", "Impact", "Residual Risk Level",
                  "Impact Description", "Mitigations", "Compliance Status"]


def sample_rows(system_rows: list[dict] | None = None) -> list[list]:
    """SYNTHETIC sample: conflicts vs the seeded demo values, an unknown control, invalid values, a duplicate row."""
    by = {r["control_id"]: r for r in (system_rows or [])}

    def r(cid, **kw):
        base = by.get(cid, {})
        g = lambda k: kw.get(k, base.get(k))  # noqa: E731
        d = g("estimated_completion_date")
        return [fmt_acronym(cid), base.get("title") or "", kw.get("implementation_status", base.get("implementation_status") or "Planned"),
                g("control_designation"), d.strftime("%m/%d/%Y") if isinstance(d, date) else d, g("responsible_entities"), g("slcm_criticality"),
                g("slcm_frequency"), g("slcm_method"), g("slcm_reporting"), g("slcm_tracking"), g("slcm_comments"), g("common_control_provider"),
                g("na_justification"), g("severity"), g("relevance_of_threat"), g("likelihood"), g("impact"), g("residual_risk_level"),
                g("impact_description"), g("mitigations"), "SYNTHETIC – ignored column"]

    full = dict(control_designation="System-Specific", estimated_completion_date="06/30/2027", responsible_entities="ISSO; System Owner",
                slcm_criticality="Moderate", slcm_frequency="Monthly", slcm_method="Automated", slcm_reporting="Monthly ConMon report to the ISSM",
                slcm_tracking="eMASS POA&M", slcm_comments="SYNTHETIC sample value from the eMASS Control Information export")
    return [
        r("ac-2", slcm_frequency="Weekly", severity="High", slcm_comments="SYNTHETIC: eMASS shows weekly account reviews"),
        r("ac-17", slcm_method="Semi-Automated", likelihood="Moderate"),
        r("at-2", responsible_entities="Training Office; ISSO"),
        r("cm-6", slcm_frequency="Biweekly"),
        r("ra-5", residual_risk_level="Extreme", estimated_completion_date="31/31/2027"),
        r("si-2", **{"slcm_comments": "x" * 4100}),
        r("cp-4", **{**full, "slcm_reporting": "Quarterly ConMon report to the AO", "slcm_tracking": "ConMon tracker reviewed by the ISSO"}),
        r("au-2", implementation_status="Not Applicable"),
        r("ia-2.1", severity="Low"),
        r("cm-2", implementation_status="Implemented", **full),
        r("zz-99"),
        r("ac-2", slcm_frequency="Daily"),
    ]


def sample_xlsx(system_rows: list[dict] | None = None) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    ws = wb.active
    ws.title = "Control Information"
    ws.append(["SYNTHETIC SAMPLE – eMASS Control Information export layout (compatibility pending verification against a real eMASS export). "
               "Not real eMASS data."])
    ws["A1"].font = Font(bold=True, color="9C0006")
    ws.append([])
    ws.append(SAMPLE_HEADERS)
    for row in sample_rows(system_rows):
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def out(imp: EmassImport, full: bool = False) -> dict:
    from app.services import emass_reports as er
    d = {"id": str(imp.id), "system_id": str(imp.system_id) if imp.system_id else None, "file_name": imp.file_name, "report_type": imp.report_type,
         "report_type_label": er.TYPES.get(imp.report_type), "detected": imp.detected, "sha256": imp.sha256, "size_bytes": imp.size_bytes,
         "status": imp.status, "uploaded_by": imp.uploaded_by_email, "created_at": imp.created_at.isoformat(), "counts": imp.counts or {},
         "applied_by": imp.applied_by_email, "finished_at": imp.finished_at.isoformat() if imp.finished_at else None, "label": er.label(imp.report_type)}
    if full:
        d.update(detection=imp.detection, report=imp.report, items=imp.items)
    return d
