"""HW/SW inventory on the REAL eMASS .xlsm template + STIG→asset correlation, consistency checks, AI rank/explain (mocked provider)."""
import io
import uuid
import zipfile
from pathlib import Path
from types import SimpleNamespace

import openpyxl
import pytest
from sqlalchemy import func, select

from app.db import SessionLocal
from app.models import (Asset, AssetCorrelation, AuditEvent, Evidence, Finding, InventoryDiscrepancy, InventoryImport, Job)
from app.services.ai.prompts import CorrelationRankOutput, DiscrepancyExplainOutput
from app.services.inventory import ai as inv_ai
from app.services.inventory import consistency as cs
from app.services.tr import xlsx
from tests.conftest import auth
from tests.test_sandbox import _signup, signup_on  # noqa: F401

TPL = (Path(__file__).parent / "fixtures/emass/HWSWList_Template-eMASS3.xlsm").read_bytes()
FAKE = "00000000-dead-beef-0000-000000000000"
HW_ROWS = [("Server", "web01", "10.0.0.1", "SN-WEB01", "RHEL 8", "Dell", "R740"),
           ("Server", "app01", "10.0.0.50", "SN-APP01", "RHEL 8", "", ""),
           ("Server", "app02", "10.0.0.50", "SN-APP02", "RHEL 8", "", ""),
           ("Server", "db01", "10.0.0.3", "SN-DB01", "RHEL 9", "", ""),
           ("Workstation", "printer01", "10.0.0.9", "SN-PR01", "FW 2.1", "HP", ""),
           ("Server", "web_01", "10.0.0.2", "SN-WEB01B", "RHEL 8", "", "")]
SW_ROWS = [("COTS Application", "Apache", "HTTP Server", "2.4.57"), ("Operating System", "Red Hat", "Enterprise Linux", "8.9")]


def workbook(hw=HW_ROWS, sw=SW_ROWS) -> bytes:
    cells = {}
    for i, (t, name, ip, serial, os_, man, model) in enumerate(hw):
        r = 8 + i
        cells.update({f"B{r}": t, f"C{r}": name, f"E{r}": ip, f"G{r}": man, f"H{r}": model, f"I{r}": serial, f"J{r}": os_})
    out = xlsx.write(TPL, "Hardware", cells)
    sc = {}
    for i, (t, v, n, ver) in enumerate(sw):
        r = 8 + i
        sc.update({f"B{r}": t, f"C{r}": v, f"D{r}": n, f"E{r}": ver})
    return xlsx.write(out, "Software", sc)


def ckl(host, ip, fqdn="", mac="", vuln="V-7001", status="Open"):
    return f"""<?xml version="1.0" encoding="UTF-8"?><CHECKLIST><ASSET><ROLE>Member Server</ROLE><ASSET_TYPE>Computing</ASSET_TYPE><HOST_NAME>{host}</HOST_NAME>
<HOST_IP>{ip}</HOST_IP><HOST_MAC>{mac}</HOST_MAC><HOST_FQDN>{fqdn}</HOST_FQDN></ASSET><STIGS><iSTIG><STIG_INFO><SI_DATA><SID_NAME>title</SID_NAME>
<SID_DATA>Inventory Test STIG (SYNTHETIC)</SID_DATA></SI_DATA><SI_DATA><SID_NAME>version</SID_NAME><SID_DATA>1</SID_DATA></SI_DATA></STIG_INFO><VULN>
<STIG_DATA><VULN_ATTRIBUTE>Vuln_Num</VULN_ATTRIBUTE><ATTRIBUTE_DATA>{vuln}</ATTRIBUTE_DATA></STIG_DATA><STIG_DATA><VULN_ATTRIBUTE>Rule_ID</VULN_ATTRIBUTE>
<ATTRIBUTE_DATA>SV-{vuln}r1_rule</ATTRIBUTE_DATA></STIG_DATA><STIG_DATA><VULN_ATTRIBUTE>Rule_Title</VULN_ATTRIBUTE><ATTRIBUTE_DATA>Synthetic rule {vuln}</ATTRIBUTE_DATA></STIG_DATA>
<STIG_DATA><VULN_ATTRIBUTE>Severity</VULN_ATTRIBUTE><ATTRIBUTE_DATA>medium</ATTRIBUTE_DATA></STIG_DATA><STIG_DATA><VULN_ATTRIBUTE>CCI_REF</VULN_ATTRIBUTE>
<ATTRIBUTE_DATA>CCI-000366</ATTRIBUTE_DATA></STIG_DATA><STATUS>{status}</STATUS></VULN></iSTIG></STIGS></CHECKLIST>""".encode()


def _jobs(db, jtype):
    from worker.handlers import HANDLERS

    for j in db.scalars(select(Job).where(Job.job_type == jtype, Job.status == "queued")):
        HANDLERS[jtype](db, j.payload)
        j.status = "succeeded"
        db.commit()


@pytest.fixture(scope="module")
def inv(client, signup_on):  # noqa: F811
    from tests.test_processing_pipeline import _process

    t = _signup(client, f"inv-{uuid.uuid4().hex[:6]}@example.com", "Inventory Org")
    h = auth(t["access_token"])
    sid = client.post("/api/v1/systems", json={"name": "Inv System", "acronym": f"IV{uuid.uuid4().hex[:3]}", "baseline_key": "moderate",
                                               "categorization_mode": "fips199", "information_types": [{"id": "D.16.3"}]}, headers=h).json()["id"]
    r = client.post(f"/api/v1/systems/{sid}/inventory/imports", files={"file": ("HWSW.xlsm", workbook())}, headers=h)
    assert r.status_code == 201, r.text
    with SessionLocal() as db:
        _jobs(db, "inventory.import")
    runs = {}
    for name, data in (("web01", ckl("web01", "10.0.0.1", "web01.example.mil", "00-11-22-33-44-55")), ("natgw", ckl("nat-gw", "10.0.0.50", vuln="V-7002")),
                       ("db01", ckl("db01", "10.0.0.99", vuln="V-7003")), ("ghost", ckl("ghost01", "10.0.0.200", vuln="V-7001"))):
        run = client.post(f"/api/v1/systems/{sid}/uploads", files={"file": (f"{name}.ckl", data)}, headers=h).json()["id"]
        _process(run)
        runs[name] = run
    with SessionLocal() as db:
        cs.run(db, uuid.UUID(sid))
        db.commit()
    role = lambda r: auth(client.post("/api/v1/sandbox/act-as", json={"role": r}, headers=h).json()["access_token"])  # noqa: E731
    return {"h": h, "sid": sid, "role": role}


def _asset(sid, name):
    with SessionLocal() as db:
        return db.scalar(select(Asset).where(Asset.system_id == uuid.UUID(sid), Asset.name == name, Asset.is_unmatched.is_(False),
                                         Asset.correlated_asset_id.is_(None)))


def _discs(client, inv, **p):
    return client.get(f"/api/v1/systems/{inv['sid']}/inventory/discrepancies", params=p, headers=inv["h"]).json()


def test_import_real_template_creates_inventory(client, inv):
    o = client.get(f"/api/v1/systems/{inv['sid']}/inventory", headers=inv["h"]).json()
    imp = o["imports"][0]
    assert imp["status"] == "imported" and imp["has_vba"] and imp["counts"]["hardware_rows"] == 6 and imp["counts"]["software_rows"] == 2
    names = {a["name"] for a in o["hardware"]}
    assert {"web01", "app01", "app02", "db01", "printer01", "web_01"} <= names
    w = next(a for a in o["hardware"] if a["name"] == "web01")
    assert w["source_label"] == "eMASS HW/SW import" and w["inventory"]["manufacturer"] == "Dell" and w["serial_number"] == "SN-WEB01"
    assert {s["name"] for s in o["software"]} == {"HTTP Server", "Enterprise Linux"}


def test_deterministic_match_attaches_findings_to_inventory_asset(inv):
    web = _asset(inv["sid"], "web01")
    with SessionLocal() as db:
        f = db.scalar(select(Finding).where(Finding.system_id == uuid.UUID(inv["sid"]), Finding.vuln_id == "V-7001", Finding.asset_id == web.id))
        assert f is not None
        c = db.scalar(select(AssetCorrelation).where(AssetCorrelation.chosen_asset_id == web.id))
        assert c.status == "matched" and c.method.startswith("deterministic:") and "hostname" in c.method


def test_shared_ip_and_conflicts_go_to_review_not_auto_linked(client, inv):
    app01, app02, db01 = (_asset(inv["sid"], n) for n in ("app01", "app02", "db01"))
    corrs = client.get(f"/api/v1/systems/{inv['sid']}/inventory/correlations", headers=inv["h"]).json()
    by = {c["observed_asset"]["name"]: c for c in corrs}
    assert "Shared IP 10.0.0.50" in by["nat-gw"]["reason"] and {x["asset_id"] for x in by["nat-gw"]["candidates"]} == {str(app01.id), str(app02.id)}
    assert "conflicts on ip" in by["db01"]["reason"]
    with SessionLocal() as db:
        for v in ("V-7002", "V-7003"):
            f = db.scalar(select(Finding).where(Finding.system_id == uuid.UUID(inv["sid"]), Finding.vuln_id == v))
            assert f.asset_id not in (app01.id, app02.id, db01.id), "ambiguous scan never attached to an inventory asset automatically"
    mr = client.get("/api/v1/mapping-review", params={"system_id": inv["sid"]}, headers=inv["h"]).json()
    assert mr["asset_correlations_pending"] == 2


def test_consistency_checks_detect_each_kind(client, inv):
    kinds = {}
    for d in _discs(client, inv):
        kinds.setdefault(d["kind"], []).append(d)
    assert any("ghost01" in d["title"] for d in kinds["missing_from_inventory"])
    assert {d["title"].split("“")[1].split("”")[0] for d in kinds["no_scan"]} >= {"printer01", "app01", "app02"}
    assert any("10.0.0.50" in d["title"] for d in kinds["shared_ip"])
    assert any("web01" in d["title"] and "web_01" in d["title"] for d in kinds["likely_duplicate"])
    assert any(d["proposed_change"]["field"] == "fqdn" and d["proposed_change"]["to"] == "web01.example.mil" for d in kinds["conflict"])
    assert any(d["proposed_change"]["field"] == "mac_address" for d in kinds["conflict"])
    assert len(kinds["ambiguous_correlation"]) == 2 and any("Manufacturer" in d["title"] for d in kinds["template_rule"])
    d = kinds["missing_from_inventory"][0]
    assert any(r["type"] == "evidence" and r["source_label"] == "Scan/checklist file import" and r["timestamp"] for r in d["records"])
    with SessionLocal() as db:
        web = db.get(Asset, _asset(inv["sid"], "web01").id)
        for f in db.scalars(select(Finding).where(Finding.asset_id == web.id)):
            f.last_seen_at = f.last_seen_at.replace(year=f.last_seen_at.year - 1)
        db.commit()
        cs.run(db, uuid.UUID(inv["sid"]))
        db.commit()
    assert any("web01" in d["title"] for d in _discs(client, inv) if d["kind"] == "stale_scan")


def test_no_overwrite_without_confirmation(client, inv):
    web = _asset(inv["sid"], "web01")
    rows = [("Server", "web01", "10.0.0.77", "SN-WEB01-NEW", "RHEL 8", "Dell", "R740")]
    client.post(f"/api/v1/systems/{inv['sid']}/inventory/imports", files={"file": ("HWSW-v2.xlsm", workbook(rows, []))}, headers=inv["h"])
    with SessionLocal() as db:
        _jobs(db, "inventory.import")
    after = _asset(inv["sid"], "web01")
    assert after.serial_number == "SN-WEB01" and after.ip_address == "10.0.0.1", "import never overwrites"
    d = next(x for x in _discs(client, inv) if x["kind"] == "import_conflict" and x["proposed_change"].get("field") == "serial_number")
    url = f"/api/v1/inventory/discrepancies/{d['id']}"
    assert client.post(url, json={"action": "approve"}, headers=inv["h"]).status_code == 409, "must be accepted first"
    assert client.post(url, json={"action": "accept"}, headers=inv["role"]("leadership_viewer")).status_code == 403
    assert client.post(url, json={"action": "accept"}, headers=inv["h"]).json()["status"] == "pending_approval"
    assert _asset(inv["sid"], "web01").serial_number == "SN-WEB01", "accepting is not applying"
    r = client.post(url, json={"action": "approve"}, headers=inv["h"])
    assert r.status_code == 200 and r.json()["status"] == "applied" and _asset(inv["sid"], "web01").serial_number == "SN-WEB01-NEW"
    with SessionLocal() as db:
        ev = db.scalar(select(AuditEvent).where(AuditEvent.action == "inventory.change_applied", AuditEvent.entity_id == d["id"]))
        assert ev.before == {"serial_number": "SN-WEB01"} and ev.after["serial_number"] == "SN-WEB01-NEW"
    ip = next(x for x in _discs(client, inv) if x["kind"] == "import_conflict" and x["proposed_change"]["type"] == "add_ip")
    iu = f"/api/v1/inventory/discrepancies/{ip['id']}"
    assert client.post(iu, json={"action": "ignore", "reason": "short"}, headers=inv["h"]).status_code == 422
    assert client.post(iu, json={"action": "ignore", "reason": "Old DHCP lease, not this server"}, headers=inv["h"]).json()["status"] == "ignored"
    assert _asset(inv["sid"], "web01").ip_address == "10.0.0.1" and web.id == _asset(inv["sid"], "web01").id


def test_create_asset_from_scan_needs_approval(client, inv):
    d = next(x for x in _discs(client, inv) if x["kind"] == "missing_from_inventory" and "ghost01" in x["title"])
    url = f"/api/v1/inventory/discrepancies/{d['id']}"
    client.post(url, json={"action": "accept"}, headers=inv["h"])
    assert _asset(inv["sid"], "ghost01") is None
    assert client.post(url, json={"action": "approve"}, headers=inv["h"]).json()["status"] == "applied"
    g = _asset(inv["sid"], "ghost01")
    assert g is not None and g.inventory_source == "inventory:created_from_scan"


def _mock(build):
    class M:
        configured = True

        async def run(self, pid, ver, variables, actor_user_id=None):
            raw = build(variables)
            model = CorrelationRankOutput if pid == "inventory.correlation_rank" else DiscrepancyExplainOutput
            return SimpleNamespace(output=model.model_validate(raw), provider="mock", model="mock-model-1")
    return M()


def test_ai_rank_validates_refs_and_creates_nothing(client, inv, monkeypatch):
    app01 = _asset(inv["sid"], "app01")
    corr = next(c for c in client.get(f"/api/v1/systems/{inv['sid']}/inventory/correlations", headers=inv["h"]).json() if c["observed_asset"]["name"] == "nat-gw")
    monkeypatch.setattr(inv_ai, "get_ai_service", lambda: _mock(lambda v: {"summary": "Likely app01", "ranking": [
        {"asset_id": str(app01.id), "confidence": 1.7, "explanation": "same IP", "identifiers": ["ip", "mac"]},
        {"asset_id": FAKE, "confidence": 0.9, "explanation": "invented", "identifiers": ["hostname"]}]}))
    with SessionLocal() as db:
        nf, na = db.scalar(select(func.count()).select_from(Finding)), db.scalar(select(func.count()).select_from(Asset))
        inv_ai.run(db, "correlation", uuid.UUID(corr["id"]))
        db.commit()
        c = db.get(AssetCorrelation, uuid.UUID(corr["id"]))
        assert c.ai["mode"] == "ai" and c.ai["label"] == "AI suggestion" and c.status == "needs_review", "AI never links"
        assert [r["asset_id"] for r in c.ai["ranking"]] == [str(app01.id)] and c.ai["ranking"][0]["confidence"] == 1.0
        assert c.ai["ranking"][0]["identifiers"] == ["ip"] and {d["id"] for d in c.ai["dropped"]} == {FAKE, str(app01.id)}
        assert db.scalar(select(func.count()).select_from(Finding)) == nf and db.scalar(select(func.count()).select_from(Asset)) == na
    d = next(x for x in _discs(client, inv) if x["kind"] == "likely_duplicate")
    rid = d["records"][0]["id"]
    monkeypatch.setattr(inv_ai, "get_ai_service", lambda: _mock(lambda v: {"explanation": "Same box", "refs": [rid, FAKE],
                                                                           "ranking": [{"record_id": FAKE, "score": 0.9, "reason": "x"}]}))
    with SessionLocal() as db:
        inv_ai.run(db, "discrepancy", uuid.UUID(d["id"]))
        db.commit()
        x = db.get(InventoryDiscrepancy, uuid.UUID(d["id"]))
        assert x.ai["refs"] == [rid] and x.ai["ranking"] == [] and len(x.ai["dropped"]) == 2 and x.status == "open"
    monkeypatch.setattr(inv_ai, "get_ai_service", lambda: SimpleNamespace(configured=False))
    with SessionLocal() as db:
        inv_ai.run(db, "correlation", uuid.UUID(corr["id"]))
        db.commit()
        assert db.get(AssetCorrelation, uuid.UUID(corr["id"])).ai["label"] == "Deterministic analysis (no AI)"


def test_confirm_correlation_links_without_changing_findings(client, inv):
    app01 = _asset(inv["sid"], "app01")
    corr = next(c for c in client.get(f"/api/v1/systems/{inv['sid']}/inventory/correlations", headers=inv["h"]).json() if c["observed_asset"]["name"] == "nat-gw")
    with SessionLocal() as db:
        f = db.scalar(select(Finding).where(Finding.system_id == uuid.UUID(inv["sid"]), Finding.vuln_id == "V-7002"))
        before = (f.id, f.asset_id, f.dedup_key, f.current_state, f.latest_evidence_id)
    url = f"/api/v1/inventory/correlations/{corr['id']}"
    assert client.post(url, json={"action": "confirm", "asset_id": FAKE, "justification": "Confirmed with network team"}, headers=inv["h"]).status_code == 422
    r = client.post(url, json={"action": "confirm", "asset_id": str(app01.id), "justification": "Confirmed with network team"}, headers=inv["h"])
    assert r.status_code == 200 and r.json()["status"] == "confirmed"
    with SessionLocal() as db:
        f = db.get(Finding, before[0])
        assert (f.id, f.asset_id, f.dedup_key, f.current_state, f.latest_evidence_id) == before, "original finding and source IDs unchanged"
        assert db.get(Asset, before[1]).correlated_asset_id == app01.id
    d = client.get(f"/api/v1/systems/{inv['sid']}/inventory/assets/{app01.id}", headers=inv["h"]).json()
    assert any(x["vuln_id"] == "V-7002" for x in d["findings"]) and "Inventory Test STIG (SYNTHETIC)" in d["applies_to"]


def test_export_preserves_vba_and_writes_only_permitted_cells(client, inv):
    r = client.post(f"/api/v1/systems/{inv['sid']}/inventory/export", headers=inv["h"])
    assert r.status_code == 200, r.text
    v = r.json()["verification"]
    assert v["ok"] and not v["outside_permitted"] and v["vba_parts"] and all(x["unchanged"] for x in v["vba_parts"].values())
    assert all(s["ok"] and not s["issues"] for s in v["sheets"].values())
    out = client.get(f"/api/v1/inventory/exports/{r.json()['id']}/download", headers=inv["h"]).content
    za, zb = zipfile.ZipFile(io.BytesIO(TPL)), zipfile.ZipFile(io.BytesIO(out))
    assert za.namelist() == zb.namelist()
    for n in za.namelist():
        if "vba" in n.lower() or n in ("[Content_Types].xml", "xl/workbook.xml", "xl/styles.xml") or "tables/" in n or "sheet3" in n or "sheet4" in n:
            assert za.read(n) == zb.read(n), n
    wb = openpyxl.load_workbook(io.BytesIO(out), keep_vba=True)
    assert wb.sheetnames == openpyxl.load_workbook(io.BytesIO(TPL)).sheetnames
    names = {wb["Hardware"][f"C{r}"].value for r in range(8, 20)}
    assert {"web01", "ghost01", "printer01"} <= names
    assert wb["Software"]["D8"].value in ("HTTP Server", "Enterprise Linux")
    rep = client.get(f"/api/v1/inventory/exports/{r.json()['id']}/validation-report", headers=inv["h"]).json()
    assert rep["human_required"] and rep["changed_cells"]


def test_mocked_sources_are_labelled():
    assert cs.src_label("mock:aws_govcloud") == "MOCKED source"
    ev = Evidence(id=uuid.uuid4(), file_name="mock-aws.json", source="mock_adapter", source_id=None)
    with SessionLocal() as db:
        assert cs.rec_evidence(db, ev)["source_label"] == "MOCKED source"
    assert cs.src_label("ingest:ckl") == "Scan/checklist file import"


def test_inventory_tenant_isolation(client, inv):
    other = auth(_signup(client, f"inv2-{uuid.uuid4().hex[:6]}@example.com", "Other Inv Org")["access_token"])
    assert client.get(f"/api/v1/systems/{inv['sid']}/inventory", headers=other).status_code == 404
    d = _discs(client, inv, status="all")[0]
    assert client.post(f"/api/v1/inventory/discrepancies/{d['id']}", json={"action": "ignore", "reason": "cross-tenant attempt"}, headers=other).status_code == 404
    c = client.get(f"/api/v1/systems/{inv['sid']}/inventory/correlations", params={"status": "all"}, headers=inv["h"]).json()[0]
    assert client.post(f"/api/v1/inventory/correlations/{c['id']}", json={"action": "reject", "justification": "cross-tenant attempt"}, headers=other).status_code == 404
    with SessionLocal() as db:
        imp = db.scalar(select(InventoryImport).where(InventoryImport.system_id == uuid.UUID(inv["sid"])))
        assert imp is not None
