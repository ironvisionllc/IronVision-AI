"""Continuous Monitoring Strategy (Task S-5) → effective eMASS SLCM values: resolution order, follow vs override, ISSM approval + SoD,
promote preview (reviewed, never automatic), bulk "follow strategy", missing filter, xlsx/OSCAL/payload use resolved values, Monitor due list."""
import io
import uuid

import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from app.db import SessionLocal
from app.models import AuditEvent, EmassControlFields
from tests.conftest import auth
from tests.test_sandbox import _signup, signup_on  # noqa: F401

BASE = {"control_designation": "System-Specific", "estimated_completion_date": "2027-03-31", "responsible_entities": "ISSO; System Owner"}
SLCM_ALL = {"slcm_criticality": "Moderate", "slcm_frequency": "Annually", "slcm_method": "Manual", "slcm_reporting": "Annual report to the ISSM",
            "slcm_tracking": "ConMon tracker", "slcm_comments": "SYNTHETIC strategy default"}


@pytest.fixture(scope="module")
def ctx(client, signup_on):  # noqa: F811
    t = _signup(client, f"is-{uuid.uuid4().hex[:6]}@example.com", "ISCM Org")
    h = auth(t["access_token"])
    sid = client.post("/api/v1/systems", json={"name": "ISCM System", "acronym": f"IS{uuid.uuid4().hex[:3]}", "baseline_key": "moderate",
                                                "categorization_mode": "fips199", "information_types": [{"id": "D.16.3"}]}, headers=h).json()["id"]
    role = lambda r: auth(client.post("/api/v1/sandbox/act-as", json={"role": r}, headers=h).json()["access_token"])  # noqa: E731
    return {"isso": role("isso"), "issm": role("issm"), "role": role, "sid": sid, "u": f"/api/v1/systems/{sid}",
            "base": f"/api/v1/systems/{sid}/workbench/controls"}


def _approve(client, ctx, body):
    assert client.put(f"{ctx['u']}/iscm", json=body, headers=ctx["isso"]).status_code == 200
    assert client.post(f"{ctx['u']}/iscm/submit", headers=ctx["isso"]).status_code == 200
    r = client.post(f"{ctx['u']}/iscm/decision", json={"approve": True, "comment": "ok"}, headers=ctx["issm"])
    assert r.status_code == 200, r.text
    return r.json()


def test_draft_does_not_feed_controls_and_only_issm_approves(client, ctx):
    client.put(f"{ctx['u']}/iscm", json={"narrative": "SYNTHETIC", "defaults": SLCM_ALL}, headers=ctx["isso"])
    e = client.get(f"{ctx['base']}/ac-2/emass", headers=ctx["isso"]).json()
    assert e["values"]["slcm_frequency"] is None, "draft strategy never feeds controls"
    assert client.put(f"{ctx['u']}/iscm", json={"defaults": {"slcm_frequency": "Hourly"}}, headers=ctx["isso"]).status_code == 422
    client.put(f"{ctx['u']}/iscm", json={"narrative": "SYNTHETIC", "defaults": SLCM_ALL}, headers=ctx["isso"])
    client.post(f"{ctx['u']}/iscm/submit", headers=ctx["isso"])
    r = client.post(f"{ctx['u']}/iscm/decision", json={"approve": True}, headers=ctx["isso"])
    assert r.status_code == 403 and "ISSM" in r.json()["detail"]
    assert client.put(f"{ctx['u']}/iscm", json={}, headers=ctx["role"]("leadership_viewer")).status_code == 403


def test_resolution_order_override_profile_default(client, ctx):
    v = _approve(client, ctx, {"narrative": "SYNTHETIC", "defaults": SLCM_ALL,
                                "profiles": [{"name": "AU", "match_type": "family", "match_value": "au", "values": {"slcm_frequency": "Monthly", "slcm_method": "Automated"}}]})
    assert v["strategy"]["status"] == "approved" and v["feeds"]["controls"] > 0
    au = client.get(f"{ctx['base']}/au-2/emass", headers=ctx["isso"]).json()
    assert au["values"]["slcm_frequency"] == "Monthly" and au["slcm_source"]["slcm_frequency"] == "Profile: AU"
    assert au["values"]["slcm_criticality"] == "Moderate" and au["slcm_source"]["slcm_criticality"] == "System default"
    # saving the pre-filled strategy values keeps following the strategy; a different value is an override
    r = client.put(f"{ctx['base']}/ac-2/emass/plan", json={"implementation_status": "Implemented", **BASE, **SLCM_ALL, "slcm_frequency": "Weekly"}, headers=ctx["isso"])
    assert r.status_code == 200, r.text
    e = client.get(f"{ctx['base']}/ac-2/emass", headers=ctx["isso"]).json()
    assert e["slcm_source"]["slcm_frequency"] == "Override" and e["values"]["slcm_frequency"] == "Weekly"
    assert e["slcm_source"]["slcm_method"] == "System default"
    with SessionLocal() as db:
        f = db.scalar(select(EmassControlFields).where(EmassControlFields.system_id == uuid.UUID(ctx["sid"]), EmassControlFields.control_id == "ac-2"))
        assert f.slcm_method is None and f.slcm_frequency == "Weekly"
    r = client.post(f"{ctx['base']}/ac-2/emass/slcm-reset", json={"fields": ["slcm_frequency"]}, headers=ctx["isso"])
    assert r.json()["reset"] == ["slcm_frequency"]
    e = client.get(f"{ctx['base']}/ac-2/emass", headers=ctx["isso"]).json()
    assert e["values"]["slcm_frequency"] == "Annually" and e["slcm_source"]["slcm_frequency"] == "System default"
    with SessionLocal() as db:
        assert db.scalar(select(AuditEvent.id).where(AuditEvent.action == "emass.slcm_reset_to_strategy", AuditEvent.details["control_id"].astext == "ac-2"))


def test_exports_and_filter_use_resolved_values(client, ctx):
    wb = load_workbook(io.BytesIO(client.get(f"{ctx['u']}/emass/implementation-plan.xlsx", headers=ctx["isso"]).content))
    ws = wb["Implementation Plan"]
    hdr = [c.value for c in ws[2]]
    row = next(r for r in ws.iter_rows(min_row=3, values_only=True) if r[0] == "AU-2")
    assert row[hdr.index("SLCM Frequency")] == "Monthly"
    p = client.get(f"{ctx['u']}/emass/controls-payload", headers=ctx["isso"])
    assert p.status_code == 200 and '"slcmFrequency": "Monthly"' in p.text.replace('":"', '": "')
    from app.models import System
    from app.services import oscal_export
    with SessionLocal() as db:
        ssp = oscal_export.build_ssp(db, db.get(System, uuid.UUID(ctx["sid"])))
    reqs = ssp["system-security-plan"]["control-implementation"]["implemented-requirements"]
    au = next(r for r in reqs if r["control-id"] == "au-2")
    assert {"name": "emass-slcm-frequency", "ns": "https://ironvisionai.com/ns/oscal/emass", "value": "Monthly"} in au["props"]
    r = client.get(f"{ctx['base']}?emass_missing=yes", headers=ctx["isso"])
    assert r.status_code == 200


def test_bulk_follow_strategy_and_merge_on_approval(client, ctx):
    r = client.post(f"{ctx['u']}/workbench/emass/bulk", json={"control_ids": ["ac-3", "ac-4"], "fields": {"slcm_frequency": "Quarterly"}, "confirm": True}, headers=ctx["isso"])
    assert r.status_code == 200, r.text
    prev = client.get(f"{ctx['u']}/iscm/promote-preview", headers=ctx["isso"]).json()["items"]
    q = next(i for i in prev if i["field"] == "slcm_frequency" and i["value"] == "Quarterly")
    assert q["controls"] == 2 and q["families"] == ["AC"]
    assert client.get(f"{ctx['base']}/ac-3/emass", headers=ctx["isso"]).json()["slcm_source"]["slcm_frequency"] == "Override", "promote is never automatic"
    r = client.post(f"{ctx['u']}/iscm/promote", json={"field": "slcm_frequency", "value": "Quarterly", "target": "family_profile", "families": ["AC"]}, headers=ctx["isso"])
    assert r.json()["strategy"]["status"] == "draft"
    client.post(f"{ctx['u']}/iscm/submit", headers=ctx["isso"])
    client.post(f"{ctx['u']}/iscm/decision", json={"approve": True}, headers=ctx["issm"])
    e = client.get(f"{ctx['base']}/ac-3/emass", headers=ctx["isso"]).json()
    assert e["values"]["slcm_frequency"] == "Quarterly" and e["slcm_source"]["slcm_frequency"] == "Profile: AC"
    r = client.post(f"{ctx['u']}/workbench/emass/bulk", json={"control_ids": ["ac-4"], "fields": {"slcm_method": "__strategy__"}, "confirm": True}, headers=ctx["isso"])
    assert r.status_code == 200 and r.json()["fields"] == ["slcm_method"]


def test_monitor_due_list(client, ctx):
    d = client.get(f"{ctx['u']}/iscm/due", headers=ctx["isso"]).json()
    au = next(i for i in d["items"] if i["control_id"] == "au-2")
    assert au["slcm_frequency"] == "Monthly" and au["status"] in ("never_monitored", "overdue", "due_soon", "on_track")
    assert sum(d["counts"].values()) == len(d["items"])
    from app.services import iscm
    with SessionLocal() as db:
        out = iscm.run_due_job(db)
    assert out["systems"]
