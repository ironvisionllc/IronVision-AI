"""eMASS Implementation Plan + Risk Assessment fields: conditional validation, limits, Other text, RBAC, audit, suggestions, bulk, exports, OSCAL."""
import io
import uuid

import pytest
from sqlalchemy import select

from app.db import SessionLocal
from app.models import AuditEvent, ControlAssignment, EmassControlFields
from tests.conftest import auth
from tests.test_sandbox import _signup, signup_on  # noqa: F401

FULL = {"control_designation": "System-Specific", "estimated_completion_date": "2027-03-31", "responsible_entities": "ISSO; System Owner",
        "slcm_criticality": "Moderate", "slcm_frequency": "Monthly", "slcm_method": "Automated", "slcm_reporting": "Monthly ConMon report to the ISSM",
        "slcm_tracking": "eMASS POA&M", "slcm_comments": "SYNTHETIC test values"}


@pytest.fixture(scope="module")
def ctx(client, signup_on):  # noqa: F811
    t = _signup(client, f"em-{uuid.uuid4().hex[:6]}@example.com", "EM Org")
    h = auth(t["access_token"])
    sid = client.post("/api/v1/systems", json={"name": "EM System", "acronym": f"EM{uuid.uuid4().hex[:3]}", "baseline_key": "moderate",
                                                "categorization_mode": "fips199", "information_types": [{"id": "D.16.3"}]}, headers=h).json()["id"]
    role = lambda r: auth(client.post("/api/v1/sandbox/act-as", json={"role": r}, headers=h).json()["access_token"])  # noqa: E731
    return {"h": role("isso"), "sid": sid, "role": role, "base": f"/api/v1/systems/{sid}/workbench/controls"}


def _plan(client, c, cid, body, h=None):
    return client.put(f"{c['base']}/{cid}/emass/plan", json=body, headers=h or c["h"])


def test_conditional_required_fields_422_and_status_sync(client, ctx):
    r = _plan(client, ctx, "ac-2", {"implementation_status": "Implemented", **{**FULL, "slcm_method": None}})
    assert r.status_code == 422 and "SLCM Method" in r.json()["detail"]
    assert _plan(client, ctx, "ac-2", {"implementation_status": "Not Applicable"}).status_code == 422
    r = _plan(client, ctx, "ac-2", {"implementation_status": "Manually Inherited", **{k: v for k, v in FULL.items() if k.startswith("slcm")}})
    assert r.status_code == 422 and "Common Control Provider" in r.json()["detail"]
    r = _plan(client, ctx, "ac-2", {"implementation_status": "Implemented", **FULL})
    assert r.status_code == 200, r.text
    d = client.get(f"{ctx['base']}/ac-2", headers=ctx["h"]).json()
    assert d["assignment"]["implementation_status"] == "implemented"  # single status: eMASS Implemented <-> internal implemented
    assert _plan(client, ctx, "ac-2", {"implementation_status": "Not Applicable", "na_justification": "Capability not in the boundary (SYNTHETIC)"}).status_code == 200
    assert client.get(f"{ctx['base']}/ac-2/emass", headers=ctx["h"]).json()["internal_status"] == "not_applicable"
    r = _plan(client, ctx, "ac-2", {"implementation_status": "Manually Inherited", "common_control_provider": "DoD", **{k: v for k, v in FULL.items() if k.startswith("slcm")}})
    assert r.status_code == 200 and r.json()["emass"]["implementation_status"] == "Manually Inherited"
    assert r.json()["emass"]["internal_status"] == "inherited"


def test_enums_char_limits_and_other_text(client, ctx):
    assert _plan(client, ctx, "ac-3", {"implementation_status": "Planned", **{**FULL, "slcm_frequency": "Biweekly"}}).status_code == 422
    assert _plan(client, ctx, "ac-3", {"implementation_status": "Planned", **{**FULL, "responsible_entities": "x" * 2001}}).status_code == 422
    assert _plan(client, ctx, "ac-3", {"implementation_status": "Planned", **{**FULL, "slcm_comments": "y" * 4001}}).status_code == 422
    assert _plan(client, ctx, "ac-3", {"implementation_status": "Planned", **{**FULL, "slcm_comments": "y" * 4000}}).status_code == 200
    other = "Weekly stand-up with the enclave NOC (custom text)"
    assert _plan(client, ctx, "ac-3", {"implementation_status": "Planned", **{**FULL, "slcm_tracking": other}}).status_code == 200
    assert client.get(f"{ctx['base']}/ac-3/emass", headers=ctx["h"]).json()["values"]["slcm_tracking"] == other
    r = client.put(f"{ctx['base']}/ac-3/emass/risk", json={"severity": "Extreme"}, headers=ctx["h"])
    assert r.status_code == 422
    assert client.put(f"{ctx['base']}/ac-3/emass/risk", json={"mitigations": "m" * 2001}, headers=ctx["h"]).status_code == 422


def test_inherited_requires_inheritance_and_locks_risk(client, ctx):
    r = _plan(client, ctx, "ac-4", {"implementation_status": "Inherited"})
    assert r.status_code == 422 and "Manually Inherited" in r.json()["detail"]
    with SessionLocal() as db:
        a = db.scalar(select(ControlAssignment).where(ControlAssignment.system_id == uuid.UUID(ctx["sid"]), ControlAssignment.control_id == "ac-4"))
        a.responsibility = "inherited"
        db.commit()
    assert _plan(client, ctx, "ac-4", {"implementation_status": "Inherited"}).status_code == 200
    r = client.put(f"{ctx['base']}/ac-4/emass/risk", json={"severity": "Low"}, headers=ctx["h"])
    assert r.status_code == 422 and "Inherited" in r.json()["detail"]
    v = client.get(f"{ctx['base']}/ac-4/emass", headers=ctx["h"]).json()
    assert v["inherited_locked"] and v["inheritance"]["responsibility"] == "inherited"


def test_rbac_read_only_roles_and_audit_before_after(client, ctx):
    sca = ctx["role"]("sca")
    assert _plan(client, ctx, "ac-5", {"implementation_status": "Planned", **FULL}, h=sca).status_code == 403
    assert client.put(f"{ctx['base']}/ac-5/emass/risk", json={"severity": "Low"}, headers=sca).status_code == 403
    assert client.get(f"{ctx['base']}/ac-5/emass", headers=sca).json()["can_edit"] is False
    so = ctx["role"]("system_owner")
    assert _plan(client, ctx, "ac-5", {"implementation_status": "Planned", **FULL}, h=so).status_code == 200
    assert client.put(f"{ctx['base']}/ac-5/emass/risk", json={"severity": "High", "likelihood": "Moderate"}, headers=ctx["h"]).status_code == 200
    with SessionLocal() as db:
        aid = str(db.scalar(select(ControlAssignment.id).where(ControlAssignment.system_id == uuid.UUID(ctx["sid"]), ControlAssignment.control_id == "ac-5")))
        ev = db.scalar(select(AuditEvent).where(AuditEvent.entity_id == aid, AuditEvent.action == "emass.risk_updated").order_by(AuditEvent.occurred_at.desc()))
        assert ev.before == {"likelihood": None, "severity": None} and ev.after == {"likelihood": "Moderate", "severity": "High"}
        assert db.scalar(select(AuditEvent.id).where(AuditEvent.entity_id == aid, AuditEvent.action == "emass.plan_updated"))


def test_suggestions_never_autosave_and_rule_fallback(client, ctx, monkeypatch):
    from app.services.ai import service as svc

    class Off:
        configured = False
    monkeypatch.setattr(svc, "get_ai_service", lambda: Off())
    before = client.get(f"{ctx['base']}/ac-6/emass", headers=ctx["h"]).json()["values"]
    r = client.post(f"{ctx['base']}/ac-6/emass/suggest", json={"implementation_status": "Planned"}, headers=ctx["h"])
    assert r.status_code == 200
    d = r.json()
    assert d["source"] == "rules" and d["interaction_id"] is None and "deterministic" in d["note"]
    fields = {s["field"] for s in d["suggestions"]}
    assert {"control_designation", "slcm_frequency", "slcm_method", "severity"} <= fields
    assert not fields & {"compliance_status", "test_result", "assessment_status"}
    assert client.get(f"{ctx['base']}/ac-6/emass", headers=ctx["h"]).json()["values"] == before  # nothing saved
    vals = {s["field"]: s["value"] for s in d["suggestions"] if s["field"] in FULL}
    r = _plan(client, ctx, "ac-6", {"implementation_status": "Planned", **vals,
                                    "assist": {"source": "rules", "fields": list(vals)}})
    assert r.status_code == 200, r.text
    assert set(r.json()["emass"]["assisted"]) >= set(vals)
    bad = _plan(client, ctx, "ac-6", {"implementation_status": "Planned", **vals, "assist": {"source": "ai", "fields": ["slcm_method"]}})
    assert bad.status_code == 422  # AI-assisted values need a logged interaction id


def test_ai_suggestions_validated_and_truncation_batched(client, ctx, monkeypatch):
    from app.services.ai import service as svc
    from app.services.ai.service import AIOutputTruncatedError
    calls = []

    class Fake:
        configured = True

        async def run(self, prompt_id, version, variables, actor_user_id=None, max_tokens=None):
            calls.append(variables["fields"])
            if len(calls) == 1:
                raise AIOutputTruncatedError("cut")
            from app.services.ai.prompts import EmassFieldsSuggestOutput
            out = EmassFieldsSuggestOutput.model_validate({"suggestions": [
                {"field": "slcm_method", "value": "Automated", "reason": "ACAS"}, {"field": "slcm_frequency", "value": "Biweekly", "reason": "bad enum"},
                {"field": "severity", "value": "High", "reason": "CAT I"}, {"field": "compliance_status", "value": "Compliant", "reason": "not allowed"}]})
            return svc.AIResult(output=out, provider="fake", model="m", latency_ms=1, interaction_id=uuid.uuid4())
    monkeypatch.setattr(svc, "get_ai_service", lambda: Fake())
    d = client.post(f"{ctx['base']}/ac-7/emass/suggest", json={"implementation_status": "Planned"}, headers=ctx["h"]).json()
    assert d["source"] == "ai" and len(calls) == 3  # truncated once, then plan + risk batches
    assert {s["field"] for s in d["suggestions"]} == {"slcm_method", "severity"}
    assert {s["field"] for s in d["rejected"]} == {"slcm_frequency"}  # compliance_status is never even considered
    with SessionLocal() as db:
        assert db.scalar(select(EmassControlFields.slcm_method).join(ControlAssignment, ControlAssignment.id == EmassControlFields.assignment_id)
                         .where(ControlAssignment.system_id == uuid.UUID(ctx["sid"]), ControlAssignment.control_id == "ac-7")) is None


def test_bulk_confirm_filters_and_exports(client, ctx):
    url = f"/api/v1/systems/{ctx['sid']}/workbench/emass/bulk"
    body = {"control_ids": ["ac-8", "ac-11", "ac-4"], "fields": {"slcm_frequency": "Quarterly", "responsible_entities": "Custom team (Other)"}}
    assert client.post(url, json=body, headers=ctx["h"]).status_code == 422
    r = client.post(url, json={**body, "confirm": True}, headers=ctx["h"]).json()
    assert set(r["updated"]) == {"ac-8", "ac-11"} and r["skipped"][0]["control_id"] == "ac-4"
    assert client.post(url, json={**body, "fields": {"slcm_method": "Sometimes"}, "confirm": True}, headers=ctx["h"]).status_code == 422
    lst = client.get(f"/api/v1/systems/{ctx['sid']}/workbench/controls", params={"emass_status": "Implemented"}, headers=ctx["h"]).json()
    assert lst == [] or all(x["implementation_status"] == "implemented" for x in lst)
    miss = {x["control_id"] for x in client.get(f"/api/v1/systems/{ctx['sid']}/workbench/controls", params={"emass_missing": "yes"}, headers=ctx["h"]).json()}
    assert "ac-8" in miss and "ac-4" not in miss
    from openpyxl import load_workbook
    x = client.get(f"/api/v1/systems/{ctx['sid']}/emass/implementation-plan.xlsx", headers=ctx["h"])
    assert x.status_code == 200
    wb = load_workbook(io.BytesIO(x.content))
    ws = wb["Implementation Plan"]
    assert "compatibility pending human verification" in ws["A1"].value
    hdr = [c.value for c in ws[2]]
    for col in ("Implementation Status", "SLCM Frequency", "SLCM Comments", "Residual Risk Level", "Common Control Provider", "N/A Justification"):
        assert col in hdr
    row = next(r for r in ws.iter_rows(min_row=3, values_only=True) if r[0] == "AC-3")
    assert row[hdr.index("SLCM Tracking")] == "Weekly stand-up with the enclave NOC (custom text)"
    assert {"Field mapping", "Validation"} <= set(wb.sheetnames)
    rep = client.get(f"/api/v1/systems/{ctx['sid']}/emass/implementation-plan/validation", headers=ctx["h"]).json()
    assert rep["rows_with_issues"] > 0 and any(i["control"] == "AC-8" for i in rep["items"])
    p = client.get(f"/api/v1/systems/{ctx['sid']}/emass/controls-payload", params={"control_id": "ac-5"}, headers=ctx["h"]).json()
    assert p["mock"] and "MOCKED" in p["label"]
    b = p["body"][0]
    assert b["acronym"] == "AC-5" and b["implementationStatus"] == "Planned" and b["slcmFrequency"] == "Monthly" and isinstance(b["estimatedCompletionDate"], int)
    assert b["severity"] == "High"


def test_oscal_ssp_carries_emass_fields_and_validates(client, ctx):
    from app.models import System
    from app.services import oscal_export as ox
    with SessionLocal() as db:
        doc = ox.build_ssp(db, db.get(System, uuid.UUID(ctx["sid"])))
        assert ox.validate("ssp", doc)["valid"]
    req = next(r for r in doc["system-security-plan"]["control-implementation"]["implemented-requirements"] if r["control-id"] == "ac-5")
    props = {p["name"]: p for p in req["props"]}
    assert props["emass-implementation-status"]["value"] == "Planned" and props["emass-slcm-frequency"]["ns"] == "https://ironvisionai.com/ns/oscal/emass"
    assert "Responsible Entities" in req["remarks"]


def test_openapi_lists_new_routes(client):
    paths = client.get("/api/openapi.json").json()["paths"]
    for p in ("/api/v1/systems/{system_id}/workbench/controls/{control_id}/emass/plan", "/api/v1/systems/{system_id}/workbench/emass/bulk",
              "/api/v1/systems/{system_id}/emass/implementation-plan.xlsx", "/api/v1/systems/{system_id}/emass/controls-payload"):
        assert p in paths


def test_migration_0029_up_down_up():
    from alembic import command
    from alembic.config import Config
    from sqlalchemy import create_engine, inspect, text
    from sqlalchemy.engine import make_url

    from tests.conftest import _BACKEND, _test_db
    url = make_url(_test_db)
    name = f"mig_{uuid.uuid4().hex[:8]}"
    admin = create_engine(url, isolation_level="AUTOCOMMIT")
    with admin.connect() as c:
        c.execute(text(f'CREATE DATABASE "{name}"'))
    target = url.set(database=name).render_as_string(hide_password=False)
    try:
        cfg = Config(str(_BACKEND / "alembic.ini"))
        cfg.set_main_option("script_location", str(_BACKEND / "alembic"))
        cfg.attributes["database_url"] = target
        command.upgrade(cfg, "0029")
        eng = create_engine(target)
        assert "emass_control_fields" in inspect(eng).get_table_names()
        command.downgrade(cfg, "0028")
        assert "emass_control_fields" not in inspect(eng).get_table_names()
        command.upgrade(cfg, "head")
        assert "emass_control_fields" in inspect(eng).get_table_names()
        eng.dispose()
    finally:
        with admin.connect() as c:
            c.execute(text(f'DROP DATABASE IF EXISTS "{name}" WITH (FORCE)'))
        admin.dispose()


def test_narrative_rule_shared_by_panel_filter_and_report(client, ctx):
    # ac-3 has complete Planned fields but no implementation statement -> missing narrative everywhere
    v = client.get(f"{ctx['base']}/ac-3/emass", headers=ctx["h"]).json()
    assert v["missing_required"] == ["implementation_narrative"]
    for flag in ("yes", "true", "1"):
        ids = {x["control_id"] for x in client.get(f"/api/v1/systems/{ctx['sid']}/workbench/controls", params={"emass_missing": flag}, headers=ctx["h"]).json()}
        assert "ac-3" in ids and "ac-4" not in ids  # ac-4 is Inherited: narrative comes from the provider
    rep = client.get(f"/api/v1/systems/{ctx['sid']}/emass/implementation-plan/validation", headers=ctx["h"]).json()
    assert next(i for i in rep["items"] if i["control"] == "AC-3")["missing"] == ["Implementation Narrative"]
    assert client.put(f"{ctx['base']}/ac-3/implementation", json={"narrative": "SYNTHETIC statement for AC-3.", "implementation_status": "planned"}, headers=ctx["h"]).status_code in (200, 201)
    assert client.get(f"{ctx['base']}/ac-3/emass", headers=ctx["h"]).json()["missing_required"] == []
    for flag in ("no", "false", "0"):
        ids = {x["control_id"] for x in client.get(f"/api/v1/systems/{ctx['sid']}/workbench/controls", params={"emass_missing": flag}, headers=ctx["h"]).json()}
        assert "ac-3" in ids
    rep = client.get(f"/api/v1/systems/{ctx['sid']}/emass/implementation-plan/validation", headers=ctx["h"]).json()
    assert all(i["control"] != "AC-3" for i in rep["items"])


def test_filter_rejects_bad_values(client, ctx):
    url = f"/api/v1/systems/{ctx['sid']}/workbench/controls"
    assert client.get(url, params={"emass_missing": "maybe"}, headers=ctx["h"]).status_code == 422
    assert client.get(url, params={"emass_status": "Done"}, headers=ctx["h"]).status_code == 422


def test_ai_suggest_accept_save_is_persisted_and_audited(client, ctx, monkeypatch):
    from app.models import AIInteraction
    from app.services.ai import service as svc
    from app.services.ai.prompts import EmassFieldsSuggestOutput

    class Fake:
        configured = True

        async def run(self, prompt_id, version, variables, actor_user_id=None, max_tokens=None):
            with SessionLocal() as db:
                i = AIInteraction(provider="fake", model="m", prompt_id=prompt_id, prompt_version=version, input_sha256="0" * 64,
                                  validation_status="valid", status="succeeded", actor_user_id=actor_user_id)
                db.add(i)
                db.commit()
                iid = i.id
            out = EmassFieldsSuggestOutput.model_validate({"suggestions": [
                {"field": "slcm_method", "value": "Semi-Automated", "reason": "ACAS + manual review"},
                {"field": "slcm_frequency", "value": "Weekly", "reason": "weekly scans"},
                {"field": "severity", "value": "Moderate", "reason": "CAT II"}]})
            return svc.AIResult(output=out, provider="fake", model="m", latency_ms=1, interaction_id=iid)
    monkeypatch.setattr(svc, "get_ai_service", lambda: Fake())
    d = client.post(f"{ctx['base']}/ac-12/emass/suggest", json={"implementation_status": "Planned"}, headers=ctx["h"]).json()
    assert d["source"] == "ai" and d["interaction_id"]
    acc = {s["field"]: s["value"] for s in d["suggestions"]}
    assist = {"source": "ai", "interaction_id": d["interaction_id"]}
    plan = {**FULL, "slcm_method": acc["slcm_method"], "slcm_frequency": acc["slcm_frequency"]}
    r = _plan(client, ctx, "ac-12", {"implementation_status": "Planned", **plan, "assist": {**assist, "fields": ["slcm_method", "slcm_frequency"]}})
    assert r.status_code == 200, r.text
    r = client.put(f"{ctx['base']}/ac-12/emass/risk", json={"severity": acc["severity"], "assist": {**assist, "fields": ["severity"]}}, headers=ctx["h"])
    assert r.status_code == 200
    v = client.get(f"{ctx['base']}/ac-12/emass", headers=ctx["h"]).json()
    assert v["values"]["slcm_method"] == "Semi-Automated" and v["values"]["slcm_frequency"] == "Weekly" and v["values"]["severity"] == "Moderate"
    assert {k: v["assisted"][k]["source"] for k in ("slcm_method", "slcm_frequency", "severity")} == {"slcm_method": "ai", "slcm_frequency": "ai", "severity": "ai"}
    assert v["assisted"]["slcm_method"]["interaction_id"] == d["interaction_id"]
    with SessionLocal() as db:
        aid = str(db.scalar(select(ControlAssignment.id).where(ControlAssignment.system_id == uuid.UUID(ctx["sid"]), ControlAssignment.control_id == "ac-12")))
        ev = db.scalar(select(AuditEvent).where(AuditEvent.entity_id == aid, AuditEvent.action == "emass.plan_updated"))
        assert ev.details["ai_assisted"] is True and ev.details["suggestion_source"] == "ai" and ev.details["interaction_id"] == d["interaction_id"]
        assert set(ev.details["suggested_fields_accepted"]) == {"slcm_method", "slcm_frequency"} and ev.after["slcm_method"] == "Semi-Automated"
        rv = db.scalar(select(AuditEvent).where(AuditEvent.entity_id == aid, AuditEvent.action == "emass.risk_updated"))
        assert rv.details["ai_assisted"] is True and rv.after == {"severity": "Moderate"}
        assert db.scalar(select(AuditEvent.id).where(AuditEvent.entity_id == aid, AuditEvent.action == "ai.emass_fields_suggest"))
