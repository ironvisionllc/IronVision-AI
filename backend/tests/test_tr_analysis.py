"""TR evidence analysis: server-side reference/quote validation, no AI control over status/tester/date, deterministic fallback."""
import json
import uuid
from types import SimpleNamespace

import pytest
from sqlalchemy import select

from app.db import SessionLocal
from app.models import Assessment, AuditEvent, ControlAssignment, Finding, Job, TrAnalysis, TrRow, TrWorkbook, User
from app.services.ai.prompts import EvidenceAnalysisOutput
from app.services.tr import analysis
from tests.test_sandbox import _signup, signup_on  # noqa: F401
from tests.test_tr import _ap, _import, _rows, populated, ws  # noqa: F401

FAKE = "00000000-dead-beef-0000-000000000000"


class MockAI:
    """Mocked provider: mimics AIService.run (strict schema validation of the raw JSON)."""
    configured = True

    def __init__(self, build):
        self.build, self.calls = build, 0

    async def run(self, prompt_id, version, variables, actor_user_id=None):
        self.calls += 1
        raw = self.build(json.loads(variables["records"]))
        return SimpleNamespace(output=EvidenceAnalysisOutput.model_validate(raw), provider="mock", model="mock-model-1")


@pytest.fixture(scope="module")
def twb(client, ws):  # noqa: F811
    c = ws["stig_ctl"]
    wid = _import(client, ws, populated([(_ap(c).split("(")[0], f"{_ap(c)}.1", ws["stig_cci"].replace("CCI-", "")), ("PE-3", "PE-3.1", "000919")]))
    return wid


def _row(twb, n):
    with SessionLocal() as db:
        return db.scalar(select(TrRow).where(TrRow.workbook_id == uuid.UUID(twb), TrRow.row_number == n)).id


def _run(monkeypatch, twb, n, ai=None):
    monkeypatch.setattr(analysis, "get_ai_service", lambda: ai or SimpleNamespace(configured=False))
    with SessionLocal() as db:
        wb, r = db.get(TrWorkbook, uuid.UUID(twb)), db.get(TrRow, _row(twb, n))
        a = analysis.analyze_row(db, wb, r)
        db.commit()
        return a.id


def _get(client, ws, twb, n):  # noqa: F811
    return client.get(f"/api/v1/tr/workbooks/{twb}/rows/{_row(twb, n)}/analysis", headers=ws["h"]).json()["latest"]


def test_deterministic_fallback_labelled(client, ws, twb, monkeypatch):  # noqa: F811
    _run(monkeypatch, twb, 7)
    a = _get(client, ws, twb, 7)
    assert a["mode"] == "deterministic" and a["label"] == "Deterministic analysis (no AI)" and a["prompt_version"] == analysis.DET_VERSION
    assert a["input_ids"] and all(s["id"] in a["input_ids"] for s in a["output"]["sources"])
    assert any(s["verdict"] in ("supports", "contradicts") for s in a["output"]["sources"])
    assert "No recorded assessor determination" in " ".join(a["output"]["missing"])
    _run(monkeypatch, twb, 8)
    b = _get(client, ws, twb, 8)
    assert b["output"]["sources"] == [] and any("No technical test evidence" in m for m in b["output"]["missing"])


def test_fabricated_reference_and_quote_rejected(client, ws, twb, monkeypatch):  # noqa: F811
    def build(recs):
        r0 = recs[0]
        return {"sources": [{"id": r0["id"], "verdict": "contradicts", "reason": "ok", "quote": r0["text"][:20]},
                            {"id": FAKE, "verdict": "supports", "reason": "invented record"},
                            {"id": recs[-1]["id"], "verdict": "supports", "reason": "bad quote", "quote": "firewall rule 99 verified by vendor"}],
                "stale": [FAKE], "conflicts": [{"ids": [FAKE], "reason": "invented"}], "explanation": f"See record {FAKE} and {r0['id']}.",
                "explanation_refs": [r0["id"], FAKE]}
    ai = MockAI(build)
    _run(monkeypatch, twb, 7, ai)
    a = _get(client, ws, twb, 7)
    o, ids = a["output"], set(a["input_ids"])
    assert ai.calls == 1 and a["mode"] == "ai" and a["label"] == "AI suggestion" and a["provider"] == "mock" and a["model"] == "mock-model-1"
    assert FAKE not in json.dumps(o) and all(s["id"] in ids for s in o["sources"]) and set(o["explanation_refs"]) <= ids
    assert "[removed reference]" in o["explanation"]
    assert o["sources"][0]["quote"], "an exact substring quote is kept"
    assert o["sources"][-1]["quote"] == "", "a fabricated quote is removed"
    fields = {(d["field"], d["id"]) for d in a["dropped"]}
    assert {("sources", FAKE), ("stale", FAKE), ("conflicts", FAKE), ("explanation", FAKE), ("explanation_refs", FAKE)} <= fields
    assert any(d["field"] == "sources.quote" for d in a["dropped"])
    with SessionLocal() as db:
        st = db.get(TrAnalysis, uuid.UUID(a["id"]))
        assert st.prompt_version == "tr.evidence_analysis@1.1.0" and st.input_bundle and st.dropped


def test_ai_cannot_set_status_tester_or_date(client, ws, twb, monkeypatch):  # noqa: F811
    with SessionLocal() as db:
        before = dict(db.get(TrRow, _row(twb, 7)).proposed)
    ai = MockAI(lambda recs: {"sources": [], "explanation": "x", "compliance_status": "Compliant", "tested_by": "AI", "date_tested": "2026-01-01"})
    _run(monkeypatch, twb, 7, ai)
    a = _get(client, ws, twb, 7)
    assert a["mode"] == "deterministic" and "AI output not used" in a["error"], "schema rejects status/tester/date → deterministic stands"
    with SessionLocal() as db:
        r = db.get(TrRow, _row(twb, 7))
        assert r.proposed == before and (r.final or {}).get("compliance_status") is None
    for extra in ({"compliance_status": "Compliant"}, {"tested_by": "AI"}, {"date_tested": "2026-01-01"}):
        resp = client.post(f"/api/v1/tr/workbooks/{twb}/rows/{_row(twb, 7)}/use-analysis", json={"analysis_id": a["id"], "text": "t", **extra}, headers=ws["h"])
        assert resp.status_code == 422, resp.text
    resp = client.post(f"/api/v1/tr/workbooks/{twb}/decide", json={"action": "edit", "row_ids": [str(_row(twb, 7))], "justification": "AI says so",
                                                               "values": {"compliance_status": "Compliant"}}, headers=ws["role"]("isso"))
    assert resp.status_code == 403


def test_conflict_open_stig_vs_compliant_assessment_and_stale(client, ws, twb, monkeypatch):  # noqa: F811
    from worker.handlers import handle_tr_generate

    sid = uuid.UUID(ws["sid"])
    with SessionLocal() as db:
        f = db.scalars(select(Finding).where(Finding.system_id == sid, Finding.source_type == "stig_checklist")).all()
        f = next(x for x in f if ws["stig_cci"] in (x.ccis or []))
        f.current_state = "open"
        a = db.scalar(select(ControlAssignment).where(ControlAssignment.system_id == sid, ControlAssignment.control_id == ws["stig_ctl"]))
        u = db.scalar(select(User).where(User.org_id == db.get(TrWorkbook, uuid.UUID(twb)).org_id))
        db.add(Assessment(system_id=sid, assignment_id=a.id, control_id=ws["stig_ctl"], assessor_id=u.id, objectives=[], methods=["test"],
                          result="satisfied", findings_text="Tested OK."))
        db.commit()
        handle_tr_generate(db, {"workbook_id": twb})
        db.commit()
    _run(monkeypatch, twb, 7)
    o = _get(client, ws, twb, 7)["output"]
    assert o["conflicts"] and "Compliant" in o["conflicts"][0]["reason"] and str(f.id) in o["conflicts"][0]["ids"]
    assert any(s["id"] == str(f.id) and s["verdict"] == "contradicts" for s in o["sources"])
    # AI that omits the conflict cannot hide it
    _run(monkeypatch, twb, 7, MockAI(lambda recs: {"sources": [], "explanation": "Looks fine."}))
    assert _get(client, ws, twb, 7)["output"]["conflicts"], "deterministic conflict is merged into AI output"
    stale = analysis.deterministic([{"id": "e1", "type": "evidence", "text": "Evidence a.txt", "freshness": "stale"},
                                    {"id": "e2", "type": "evidence", "text": "Evidence b.txt", "freshness": "current"}])
    assert stale["stale"] == ["e1"] and next(s for s in stale["sources"] if s["id"] == "e1")["reason"].startswith("Evidence is stale")


def _tr_cell(client, ws, twb, e):  # noqa: F811
    import io

    import openpyxl
    d = client.get(f"/api/v1/tr/workbooks/{twb}", headers=ws["h"]).json()["definition"]
    wbk = openpyxl.load_workbook(io.BytesIO(client.get(f"/api/v1/tr/exports/{e['id']}/download", headers=ws["h"]).content))
    return wbk[d["sheet"]][f"{d['roles']['test_results']}7"].value


def test_export_unchanged_unless_text_applied_then_attributed(client, ws, twb, monkeypatch):  # noqa: F811
    from worker.handlers import handle_tr_analyze

    rid = str(_row(twb, 7))
    sca = ws["role"]("sca")
    assert client.post(f"/api/v1/tr/workbooks/{twb}/decide", json={"action": "approve", "row_ids": [rid], "justification": "ok",
                                                                                                     "tested_by": "Lee, Sam", "date_tested": "2026-09-01"}, headers=sca).status_code == 200
    e1 = client.post(f"/api/v1/tr/workbooks/{twb}/export", headers=ws["h"]).json()
    j = client.post(f"/api/v1/tr/workbooks/{twb}/analyze", json={"row_ids": [rid]}, headers=ws["h"])
    assert j.status_code == 202, j.text
    monkeypatch.setattr(analysis, "get_ai_service", lambda: SimpleNamespace(configured=False))
    with SessionLocal() as db:
        assert handle_tr_analyze(db, db.get(Job, uuid.UUID(j.json()["job_id"])).payload) == {"deterministic": 1}
        db.commit()
    e2 = client.post(f"/api/v1/tr/workbooks/{twb}/export", headers=ws["h"]).json()
    assert e1["sha256"] == e2["sha256"] and _tr_cell(client, ws, twb, e1) == _tr_cell(client, ws, twb, e2), "analysis alone changes no exported cell"
    a = _get(client, ws, twb, 7)
    text = a["output"]["explanation"] + " Reviewed by assessor."
    r = client.post(f"/api/v1/tr/workbooks/{twb}/rows/{rid}/use-analysis", json={"analysis_id": a["id"], "text": text}, headers=ws["role"]("isso"))
    assert r.status_code == 200 and r.json()["state"] == "edited" and r.json()["final"]["ai_assisted"]["edited_by_user"] is True
    assert r.json()["final"]["compliance_status"] == "Compliant", "status untouched by applying text"
    e3 = client.post(f"/api/v1/tr/workbooks/{twb}/export", headers=ws["h"]).json()
    assert _tr_cell(client, ws, twb, e3) != text, "not exported until the SCA re-approves"
    assert client.post(f"/api/v1/tr/workbooks/{twb}/decide", json={"action": "approve", "row_ids": [rid], "justification": "ok",
                                                                                                     "tested_by": "Lee, Sam", "date_tested": "2026-09-01"}, headers=sca).status_code == 200
    e4 = client.post(f"/api/v1/tr/workbooks/{twb}/export", headers=ws["h"]).json()
    assert _tr_cell(client, ws, twb, e4) == text and e4["verification"]["ok"] and not e4["verification"]["outside_permitted"]
    with SessionLocal() as db:
        row = db.get(TrRow, uuid.UUID(rid))
        assert row.final["ai_assisted"]["analysis_id"] == a["id"] and row.final["test_results"] == text
        ev = db.scalar(select(AuditEvent).where(AuditEvent.action == "tr.ai_text_applied", AuditEvent.entity_id == rid))
        assert ev is not None and ev.after["final"]["ai_assisted"]["mode"] == "deterministic"


def test_analysis_tenant_isolation_and_permissions(client, ws, twb):  # noqa: F811
    from tests.conftest import auth

    other = auth(_signup(client, f"tra-{uuid.uuid4().hex[:6]}@example.com", "Other Org")["access_token"])
    rid = str(_row(twb, 7))
    assert client.get(f"/api/v1/tr/workbooks/{twb}/rows/{rid}/analysis", headers=other).status_code == 404
    assert client.post(f"/api/v1/tr/workbooks/{twb}/analyze", json={"row_ids": [rid]}, headers=other).status_code == 404
    assert client.post(f"/api/v1/tr/workbooks/{twb}/analyze", json={"row_ids": [rid]}, headers=ws["role"]("leadership_viewer")).status_code == 403
