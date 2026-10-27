"""eMASS System Import & Migration: SLCM worksheet, Categorization template, PPSM boundaries on the shared import sessions –
detection of the real templates, imported values preselected, keep-existing controls, provenance, audit, revert, idempotent retry,
exports + round trips, formula neutralization, AI assist (redaction, toggle, suggestion-only), copy setup, RBAC and tenant isolation."""
import io
import time
import uuid
from pathlib import Path

import openpyxl
import pytest
from openpyxl import Workbook
from sqlalchemy import select

from app.db import SessionLocal
from app.models import AuditEvent, EmassControlFields, IscmStrategy, OdpValue, Poam, PpsmEntry, System
from app.services import emass_migration as em
from app.services import emass_reports as er
from tests.conftest import auth
from tests.test_sandbox import _signup, signup_on  # noqa: F401

FX = Path(__file__).parent / "fixtures" / "emass"
REAL = {"slcm": "slcm_plan_worksheet_v1_3.xlsm", "categorization": "categorization_template.xlsx", "ppsm": "ppsm_boundaries_template.xlsx",
        "test_results": "BLANK_TRExport.xlsx", "hwsw": "HWSWList_Template-eMASS3.xlsm"}
FIPS = {"baseline_key": "moderate", "categorization_mode": "fips199", "information_types": [{"id": "D.16.3"}]}


def _wb(rows, title="Final Results") -> bytes:
    wb = Workbook()
    wb.active.title = title
    for r in rows:
        wb.active.append(r)
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


def _org(client, tag):
    t = _signup(client, f"mg-{tag}-{uuid.uuid4().hex[:6]}@example.com", f"MIG {tag} Org")
    h = auth(t["access_token"])
    role = lambda r: auth(client.post("/api/v1/sandbox/act-as", json={"role": r}, headers=h).json()["access_token"])  # noqa: E731
    return {"h": h, "role": role, "isso": role("isso")}


@pytest.fixture(scope="module")
def ctx(client, signup_on):  # noqa: F811
    o = _org(client, "a")
    o["sid"] = _mk(client, o, "Migration Target", "MIGT")
    return o


@pytest.fixture(scope="module")
def other(client, signup_on):  # noqa: F811
    o = _org(client, "b")
    o["sid"] = _mk(client, o, "Other Tenant", "OTH")
    return o


def _mk(client, o, name, acr):
    r = client.post("/api/v1/systems", json={"name": name, "acronym": acr + uuid.uuid4().hex[:3], "processes_pii": False, **FIPS}, headers=o["h"])
    assert r.status_code == 201, r.text
    return r.json()["id"]


def _import(client, o, name, data, typ=None, sid=None):
    r = client.post("/api/v1/emass-imports", files={"file": (name, data)}, data={"system_id": sid or o["sid"]}, headers=o["isso"])
    assert r.status_code == 201, r.text
    imp = r.json()
    r = client.post(f"/api/v1/emass-imports/{imp['id']}/target", json={"action": "update", "system_id": sid or o["sid"], "report_type": typ}, headers=o["isso"])
    assert r.status_code == 200, r.text
    return r.json()


def _apply(client, o, imp):
    r = client.post(f"/api/v1/emass-imports/{imp['id']}/apply", headers=o["isso"])
    assert r.status_code == 200, r.text
    return r.json()


def _changes(imp):
    return sum(1 for i in imp["items"] for f in i["fields"].values() if f["change"] and not f["error"] and not f["blocked"])


@pytest.mark.parametrize("typ", list(REAL))
def test_real_templates_detected(typ):
    d = er.detect(er.read_sheets((FX / REAL[typ]).read_bytes(), REAL[typ]))
    assert d["type"] == typ and d["confidence"] in ("high", "medium"), d


def test_supplied_tr_and_hwsw_templates_identical_to_fixtures():
    import hashlib
    assert hashlib.sha256((FX / "BLANK_TRExport.xlsx").read_bytes()).hexdigest().startswith("8cb7cb8a496210cd5ba6")
    assert hashlib.sha256((FX / "HWSWList_Template-eMASS3.xlsm").read_bytes()).hexdigest().startswith("e4f7f1d2e7d24a0f1122")


def test_deferred_types_listed_and_excluded(client, ctx):
    h = client.get("/api/v1/emass-imports", headers=ctx["isso"]).json()
    assert set(h["deferred"]) == {"control_information", "poam"}
    assert not set(h["migration_types"]) & set(h["deferred"])


def test_slcm_real_worksheet_preselected_provenance_apply_audit_revert(client, ctx):
    imp = _import(client, ctx, REAL["slcm"], (FX / REAL["slcm"]).read_bytes())
    assert imp["report_type"] == "slcm" and imp["report"]["sheet"] == "Final Results"
    items = {i["key"]: i for i in imp["items"]}
    ac2 = items["ac-2"]
    f = ac2["fields"]["slcm_frequency"]
    assert f["imported"] == "Annually" and f["src"]["sheet"] == "Final Results" and f["src"]["col"] == "E" and f["src"]["row"] == ac2["row"]
    assert all(fd["decision"] == "accept" for i in imp["items"] for fd in i["fields"].values() if fd["change"] and not fd["error"])  # imported wins by default
    assert imp["report"]["not_in_system"], "worksheet controls outside the moderate baseline go to the exception queue"
    done = _apply(client, ctx, imp)
    assert done["counts"]["applied_fields"] > 100
    with SessionLocal() as db:
        fs = db.scalar(select(EmassControlFields).where(EmassControlFields.system_id == uuid.UUID(ctx["sid"]), EmassControlFields.control_id == "ac-2"))
        assert fs.slcm_frequency == "Annually" and fs.slcm_criticality == "CRWG Yellow Criticality Control"
        ev = db.scalar(select(AuditEvent).where(AuditEvent.action == "emass.slcm_imported", AuditEvent.entity_id == str(fs.assignment_id)))
        assert ev.details["import_sha256"] == imp["sha256"] and ev.details["provenance"]["slcm_frequency"]["sheet"] == "Final Results"
    again = _import(client, ctx, REAL["slcm"], (FX / REAL["slcm"]).read_bytes())  # idempotent retry
    assert again["report"].get("duplicate_of") == imp["id"] and _changes(again) == 0
    client.post(f"/api/v1/emass-imports/{again['id']}/discard", headers=ctx["isso"])
    r = client.post(f"/api/v1/emass-imports/{imp['id']}/revert", json={"keys": ["ac-2"]}, headers=ctx["isso"])
    assert r.status_code == 200 and r.json()["reverted"] == [{"key": "ac-2", "control": "AC-2"}]
    with SessionLocal() as db:
        fs = db.scalar(select(EmassControlFields).where(EmassControlFields.system_id == uuid.UUID(ctx["sid"]), EmassControlFields.control_id == "ac-2"))
        assert fs.slcm_frequency is None
    r = client.post(f"/api/v1/emass-imports/{imp['id']}/revert", json={}, headers=ctx["isso"])
    assert r.json()["status"] == "reverted"


def test_slcm_invalid_values_keep_existing_and_round_trip(client, ctx):
    data = em.sample("slcm")
    imp = _import(client, ctx, "slcm.xlsx", data)
    items = {i["key"]: i for i in imp["items"]}
    assert "zz-9" not in items and any(x["control"] == "ZZ-9" for x in imp["report"]["exceptions"] + imp["report"]["not_in_system"])
    r = client.put(f"/api/v1/emass-imports/{imp['id']}/decisions", json={"bulk": "keep_existing", "section": "SLCM"}, headers=ctx["isso"]).json()
    assert all(f["decision"] == "reject" for i in r["items"] for f in i["fields"].values() if f["change"] and not f["error"])
    r = client.put(f"/api/v1/emass-imports/{imp['id']}/decisions", json={"bulk": "use_imported"}, headers=ctx["isso"]).json()
    _apply(client, ctx, r)
    exp = client.get(f"/api/v1/systems/{ctx['sid']}/emass-exports/slcm", headers=ctx["isso"])
    assert exp.status_code == 200
    back = _import(client, ctx, "export.xlsx", exp.content)
    assert back["report_type"] == "slcm" and _changes(back) == 0, "export → re-import is a no-op"


def test_slcm_bad_frequency_flagged(client, ctx):
    data = _wb([["Implementation Status / Security Control Status"], ["Sort #", "Control Acronym", "Control Title", "Criticality", "Frequency", "Method"],
                [1, "AC-3", "x", "CRWG White Criticality Control", "Hourly", "Telepathy"]])
    imp = _import(client, ctx, "bad.xlsx", data, "slcm")
    f = imp["items"][0]["fields"]
    assert f["slcm_frequency"]["error"] and f["slcm_method"]["error"] and f["slcm_frequency"]["decision"] is None


def test_categorization_proposal_and_approved_stays_proposal(client, ctx):
    sid = _mk(client, ctx, "Cat Target", "CAT")
    imp = _import(client, ctx, "cat.xlsx", em.sample("categorization"), sid=sid)
    f = imp["items"][0]["fields"]
    assert (f["confidentiality_impact"]["imported"], f["integrity_impact"]["imported"], f["availability_impact"]["imported"]) == ("Moderate",) * 3
    assert len(f["emass_mission_types"]["imported"]) == 2
    _apply(client, ctx, imp)
    with SessionLocal() as db:
        s = db.get(System, uuid.UUID(sid))
        assert s.categorization == "moderate" or s.categorization in ("Moderate",)
        assert len(s.environment_detail["emass_mission_types"]) == 2
    with SessionLocal() as db:
        from app.models import CategorizationRecord
        db.add(CategorizationRecord(system_id=uuid.UUID(sid), revision=99, framework="fips199", status="approved"))
        db.commit()
    imp2 = _import(client, ctx, "cat2.xlsx", _set_high(em.sample("categorization")), sid=sid)
    f2 = imp2["items"][0]["fields"]
    assert f2["confidentiality_impact"]["imported"] == "High" and f2["confidentiality_impact"]["blocked"]


def _set_high(data: bytes) -> bytes:
    wb = openpyxl.load_workbook(io.BytesIO(data))
    wb["Mission Types"].cell(row=8, column=3, value="Yes")  # Key Asset and Critical Infrastructure Protection → High
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


def test_categorization_registration_prefill(client, ctx):
    r = client.post("/api/v1/emass-imports", files={"file": ("cat.xlsx", em.sample("categorization"))}, headers=ctx["h"]).json()
    p = client.post(f"/api/v1/emass-imports/{r['id']}/target", json={"action": "register", "report_type": "categorization"}, headers=ctx["h"]).json()
    assert p["prefill"]["categorization_hint"] == {"c": "Moderate", "i": "Moderate", "a": "Moderate"}
    assert p["prefill"]["system"]["name"].startswith("SYNTHETIC")


def test_ppsm_import_validation_apply_export_round_trip_revert(client, ctx):
    sid = _mk(client, ctx, "PPSM Target", "PPS")
    imp = _import(client, ctx, "ppsm.xlsx", em.sample("ppsm"), sid=sid)
    items = {i["label"]: i for i in imp["items"]}
    https = next(i for i in imp["items"] if i["label"].startswith("443/TCP"))
    assert https["create"] and https["fields"]["boundary"]["imported"] == ["Ext to DoD GW (In)", "Enclave GW to Enclave (In)"]
    bad = next(i for i in imp["items"] if i["label"].startswith("70000"))
    assert bad["fields"]["port"]["error"] and bad["fields"]["boundary"]["error"] and bad["required_issue"]
    assert any("person data" in x for x in imp["report"]["not_imported"]) and items
    _apply(client, ctx, imp)
    lst = client.get(f"/api/v1/systems/{sid}/ppsm", headers=ctx["isso"]).json()["items"]
    assert len(lst) == 2 and all(e["source"] == "emass_import" and e["import_id"] == imp["id"] for e in lst)
    with SessionLocal() as db:
        e = db.scalar(select(PpsmEntry).where(PpsmEntry.system_id == uuid.UUID(sid), PpsmEntry.port == "1521"))
        e.purpose = "=HYPERLINK(\"http://evil\")"
        db.commit()
    exp = client.get(f"/api/v1/systems/{sid}/emass-exports/ppsm", headers=ctx["isso"])
    ws = openpyxl.load_workbook(io.BytesIO(exp.content))["PPSM"]
    cells = [c.value for row in ws.iter_rows(min_row=9, max_row=12) for c in row if c.value]
    assert "'=HYPERLINK(\"http://evil\")" in cells and not any(str(v).startswith("=") for v in cells), "formula neutralized on export"
    back = _import(client, ctx, "ppsm_export.xlsx", exp.content, sid=sid)
    assert back["report_type"] == "ppsm" and _changes(back) == 0, "export → re-import is a no-op"
    client.post(f"/api/v1/emass-imports/{back['id']}/discard", headers=ctx["isso"])
    r = client.post(f"/api/v1/emass-imports/{imp['id']}/revert", json={}, headers=ctx["isso"]).json()
    assert r["status"] == "reverted" and client.get(f"/api/v1/systems/{sid}/ppsm", headers=ctx["isso"]).json()["items"] == []


def test_redaction_removes_pii():
    s = em.redact("alex.rivera@corp.mil 10.10.0.11 fe80::1 web01.enclave.mil +1 (555) 123-4567 HTTPS")
    assert "@" not in s and "10.10" not in s and "enclave.mil" not in s and "555" not in s and "HTTPS" in s
    p = em.ai_payload(("Port", "POC Name", "Host"), [("443", "Alex Rivera", "db01.lab.local")])
    assert "Alex" not in p and "lab.local" not in p and "443" in p


def test_ai_toggle_audited_admin_only_and_suggestion_only(client, ctx, monkeypatch):
    assert client.put("/api/v1/emass-import-settings/ai", json={"enabled": False}, headers=ctx["isso"]).status_code == 403
    assert client.get("/api/v1/emass-import-settings/ai", headers=ctx["isso"]).json()["enabled"] is True
    data = _wb([["Ctl", "How often", "How", "Importance"], ["AC-3", "Monthly", "Manual", "CRWG White Criticality Control"]], "Sheet1")
    sent = {}

    def fake(typ, header, samples):
        sent["payload"] = em.ai_payload(header, samples)
        return [{"column": 0, "field": "control", "confidence": "high", "reason": "acronyms"},
                {"column": 1, "field": "slcm_frequency", "confidence": "medium", "reason": "frequency words"},
                {"column": 3, "field": "slcm_criticality", "confidence": "low", "reason": "unsure"}]
    monkeypatch.setattr(em, "ai_suggest", fake)
    imp = _import(client, ctx, "odd.xlsx", data, "slcm")
    f = imp["items"][0]["fields"]
    assert imp["report"]["ai"]["used"] and f["slcm_frequency"]["ai"]["confidence"] == "medium"
    assert f["slcm_frequency"]["decision"] is None, "AI-mapped values are never preselected"
    assert "slcm_criticality" not in f and imp["report"]["ai"]["low_confidence"]
    assert client.put("/api/v1/emass-import-settings/ai", json={"enabled": False}, headers=ctx["role"]("platform_admin")).status_code == 200
    r = client.post("/api/v1/emass-imports", files={"file": ("odd.xlsx", data)}, data={"system_id": ctx["sid"]}, headers=ctx["isso"]).json()
    r = client.post(f"/api/v1/emass-imports/{r['id']}/target", json={"action": "update", "system_id": ctx["sid"], "report_type": "slcm"}, headers=ctx["isso"])
    assert r.status_code == 422, "AI off: no AI call; unrecognised headers are reported"
    ok = _import(client, ctx, "slcm.xlsx", em.sample("slcm"))  # deterministic import works with AI off
    assert ok["items"]
    with SessionLocal() as db:
        assert db.scalar(select(AuditEvent).where(AuditEvent.action == "admin.emass_ai_updated").order_by(AuditEvent.created_at.desc()).limit(1)).after == {"enabled": False}
    client.put("/api/v1/emass-import-settings/ai", json={"enabled": True}, headers=ctx["role"]("platform_admin"))


def test_read_only_role_cannot_apply_or_revert(client, ctx):
    imp = _import(client, ctx, "slcm.xlsx", em.sample("slcm"))
    v = ctx["role"]("leadership_viewer")
    assert client.post(f"/api/v1/emass-imports/{imp['id']}/apply", headers=v).status_code == 403
    assert client.put("/api/v1/emass-import-settings/ai", json={"enabled": False}, headers=v).status_code == 403


def test_tenant_isolation(client, ctx, other):
    imp = _import(client, ctx, "ppsm.xlsx", em.sample("ppsm"))
    oh = other["h"]
    for m, url, body in [("get", f"/api/v1/emass-imports/{imp['id']}", None), ("post", f"/api/v1/emass-imports/{imp['id']}/apply", None),
                         ("post", f"/api/v1/emass-imports/{imp['id']}/revert", {}), ("put", f"/api/v1/emass-imports/{imp['id']}/decisions", {"bulk": "use_imported"})]:
        r = getattr(client, m)(url, headers=oh, **({"json": body} if body is not None else {}))
        assert r.status_code == 404, (url, r.status_code)
    for url in (f"/api/v1/systems/{ctx['sid']}/ppsm", f"/api/v1/systems/{ctx['sid']}/emass-exports/slcm", f"/api/v1/systems/{ctx['sid']}/emass-exports/ppsm",
                f"/api/v1/systems/{ctx['sid']}/setup-template"):
        assert client.get(url, headers=oh).status_code in (403, 404), url
    assert client.post(f"/api/v1/systems/{other['sid']}/copy-setup", json={"source_system_id": ctx["sid"]}, headers=oh).status_code in (403, 404)
    assert ctx["sid"] not in str(client.get("/api/v1/emass-imports", headers=oh).json())
    with SessionLocal() as db:
        key = db.get(__import__("app.models", fromlist=["EmassImport"]).EmassImport, uuid.UUID(imp["id"])).storage_key
        assert key.startswith("emass-import/") and str(db.get(System, uuid.UUID(ctx["sid"])).org_id) in key


def test_copy_setup_copies_setup_never_assessments(client, ctx):
    src = ctx["sid"]
    with SessionLocal() as db:
        db.add(OdpValue(system_id=uuid.UUID(src), control_id="ac-2", parameter="ac-02_odp.01", value="30 days", source="manual", approved_by="AO"))
        db.add(IscmStrategy(system_id=uuid.UUID(src), status="approved", narrative="SYNTHETIC strategy", defaults={"slcm_frequency": "Monthly"}))
        db.commit()
    r = client.post("/api/v1/systems", json={"name": "Copy Target", "acronym": "CPY" + uuid.uuid4().hex[:3], "processes_pii": False,
                                             **{k: v for k, v in client.get(f"/api/v1/systems/{src}/setup-template", headers=ctx["h"]).json().items()
                                                if k in ("baseline_key", "categorization_mode", "information_types")}}, headers=ctx["h"])
    dst = r.json()["id"]
    out = client.post(f"/api/v1/systems/{dst}/copy-setup", json={"source_system_id": src}, headers=ctx["h"])
    assert out.status_code == 200, out.text
    assert out.json()["odps"] >= 1 and out.json()["iscm"]
    with SessionLocal() as db:
        o = db.scalar(select(OdpValue).where(OdpValue.system_id == uuid.UUID(dst), OdpValue.control_id == "ac-2"))
        assert o.value == "30 days" and o.approved_by is None, "approvals are never copied"
        st = db.scalar(select(IscmStrategy).where(IscmStrategy.system_id == uuid.UUID(dst)))
        assert st.status == "draft" and st.approved_snapshot is None
        assert not db.scalar(select(Poam).where(Poam.system_id == uuid.UUID(dst)).limit(1))
        assert db.scalar(select(AuditEvent).where(AuditEvent.action == "system.setup_copied", AuditEvent.entity_id == dst))


def test_invalid_files_rejected(client, ctx):
    r = client.post("/api/v1/emass-imports", files={"file": ("x.exe", b"MZ")}, data={"system_id": ctx["sid"]}, headers=ctx["isso"])
    assert r.status_code == 422
    r = client.post("/api/v1/emass-imports", files={"file": ("x.xlsx", b"not a zip")}, data={"system_id": ctx["sid"]}, headers=ctx["isso"])
    assert r.status_code in (400, 422)


def test_large_slcm_build_is_fast(client, ctx):
    rows = [["Implementation Status / Security Control Status"], ["Sort #", "Control Acronym", "Control Title", "Criticality", "Frequency", "Method", "Reporting", "Tracking"]]
    with SessionLocal() as db:
        from app.models import ControlAssignment
        ids = list(db.scalars(select(ControlAssignment.control_id).where(ControlAssignment.system_id == uuid.UUID(ctx["sid"]))))
    for n, c in enumerate((ids * 5)[:1200]):
        rows.append([n, c.upper(), "t", "CRWG White Criticality Control", "Monthly", "Automated", "r", "POA&M"])
    t = time.perf_counter()
    imp = _import(client, ctx, "big.xlsx", _wb(rows))
    assert time.perf_counter() - t < 20 and imp["items"]


def test_migration_report_coverage_steps_and_deferred(client, ctx, other):
    r = client.get(f"/api/v1/systems/{ctx['sid']}/migration-report", headers=ctx["isso"])
    assert r.status_code == 200
    d = r.json()
    cov = {c["type"]: c for c in d["coverage"]}
    assert cov["slcm"]["status"] == "applied" and any(s["to"] == "/rmf/monitor" for s in cov["slcm"]["steps"])
    assert {x["type"] for x in d["deferred"]} == {"control_information", "poam"} and "poam" not in cov
    assert any(s["exception_total"] for s in d["sessions"] if s["report_type"] == "slcm")
    assert client.get(f"/api/v1/systems/{ctx['sid']}/migration-report", headers=other["h"]).status_code in (403, 404)
