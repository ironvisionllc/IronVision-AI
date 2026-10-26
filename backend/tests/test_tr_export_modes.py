"""GUARD TESTS – eMASS Test Results export modes on the REAL blank fixture (tests/fixtures/emass/BLANK_TRExport.xlsx).

Mode 1 (populated template): only M–P may change, only on matched rows that have an SCA-approved, tester-confirmed result.
Mode 2 (blank template): only A, F, G, H (generated rows) and M–P (approved generated rows) may change; headers untouched.
The diff is independent of the exporter's own verify(): every cell value/type/style/number format, merges, data validations,
conditional formatting, column widths, row heights, defined names, sheet list/order and every other zip part are compared.
Set TR_VERIFY_OUT=<dir> to also write the inputs, outputs and diff reports for human review."""
import io
import json
import os
import uuid
from datetime import date, timedelta
from pathlib import Path

import openpyxl
import pytest
from sqlalchemy import select

from app.db import SessionLocal
from app.models import Cci, CciControlMap, ControlAssignment, TrRow
from app.services.tr import service as trs
from app.services.tr import xlsx
from tests import xlsx_diff
from tests.test_sandbox import signup_on  # noqa: F401
from tests.test_tr import BLANK, _ap, _import, _rows, ws  # noqa: F401  (ws = synthetic sandbox system fixture)

pytestmark = pytest.mark.guard
OUT = os.environ.get("TR_VERIFY_OUT")
RESULT = set("MNOP")


def _all_rows(wid) -> dict:
    """Every TrRow of the workbook (the rows API pages at 500)."""
    with SessionLocal() as db:
        return {r.row_number: {"id": str(r.id), "row": r.row_number, "added": r.added}
                for r in db.scalars(select(TrRow).where(TrRow.workbook_id == uuid.UUID(wid)))}


def _save(name: str, data) -> None:
    if OUT:
        Path(OUT).mkdir(parents=True, exist_ok=True)
        p = Path(OUT) / name
        p.write_bytes(data) if isinstance(data, bytes) else p.write_text(json.dumps(data, indent=2, default=str))


def _export(client, ws, wid):  # noqa: F811
    e = client.post(f"/api/v1/tr/workbooks/{wid}/export", headers=ws["h"])
    assert e.status_code == 200, e.text
    e = e.json()
    return e, client.get(f"/api/v1/tr/exports/{e['id']}/download", headers=ws["h"]).content


def _report(name, mode, e, d, sheet="Template"):
    v = e["verification"]
    rep = {"mode": v["mode"], "mode_label": v["mode_label"], "mode_explanation": v["mode_explanation"], "modified_columns_from_report": v["modified_columns"],
           "modified_columns_from_diff": xlsx_diff.by_column(d, sheet), "row_counts": v["row_counts"], "unmatched_controls": v["unmatched_controls"],
           "controls_not_in_system": v["controls_not_in_system"], "rows_not_populated": v["rows_not_populated"][:60],
           "rows_not_populated_total": len(v["rows_not_populated"]), "structure_changes": d["structure"], "zip_parts_changed": d["zip_parts_changed"],
           "other_sheets_changed": [s for s in d["cells"] if s != sheet], "sheet_order_changed": d["sheet_order_changed"],
           "defined_names_changed": d["defined_names_changed"], "changed_cells_sample": d["cells"].get(sheet, [])[:40]}
    _save(f"{name}_diff_report.json", rep)
    return rep


def _approve(client, ws, wid, ids, testers, status="Compliant"):  # noqa: F811
    per = {i: {"tested_by": t, "date_tested": dt} for i, (t, dt) in zip(ids, testers)}
    r = client.post(f"/api/v1/tr/workbooks/{wid}/decide", json={"action": "approve", "row_ids": ids, "justification": "guard test approval",
                                                                 "status": status, "per_row": per,
                                                                 "values": {"test_results": "SYNTHETIC guard test result: examined the configuration."}},
                    headers=ws["role"]("sca"))
    assert r.status_code == 200, r.text
    return r.json()


@pytest.fixture(scope="module")
def mode2(client, ws):  # noqa: F811
    from worker.handlers import handle_tr_detect, handle_tr_generate

    wid = client.post(f"/api/v1/systems/{ws['sid']}/tr/workbooks", files={"file": ("BLANK_TRExport.xlsx", BLANK)}, headers=ws["h"]).json()["id"]
    with SessionLocal() as db:
        handle_tr_detect(db, {"workbook_id": wid})
        db.commit()
        handle_tr_generate(db, {"workbook_id": wid, "populate": True})
        db.commit()
    rows = _all_rows(wid)
    ids = [rows[n]["id"] for n in sorted(rows)[:6]]
    bad = client.post(f"/api/v1/tr/workbooks/{wid}/decide", json={"action": "approve", "row_ids": ids[:1], "status": "Compliant", "tested_by": "Lee, Sam",
                                                                   "date_tested": (date.today() + timedelta(days=1)).isoformat()}, headers=ws["role"]("sca"))
    assert bad.status_code == 422 and "future" in bad.json()["detail"]
    _approve(client, ws, wid, ids[:5], [(f"Tester {i}", f"2026-08-0{i + 1}") for i in range(5)])
    with SessionLocal() as db:  # row 5: an old approval without tester confirmation must be excluded
        r = db.get(TrRow, uuid.UUID(ids[4]))
        r.final = {k: v for k, v in r.final.items() if k != "tester_confirmed"}
        db.commit()
    e, out = _export(client, ws, wid)
    _save("mode2_input_BLANK_TRExport.xlsx", BLANK)
    _save("mode2_output_blank_template_export.xlsx", out)
    return {"wid": wid, "rows": rows, "approved": [rows[n]["row"] for n in sorted(rows)[:4]], "unconfirmed": rows[sorted(rows)[4]]["row"], "e": e, "out": out}


def test_guard_mode2_blank_only_A_F_G_H_and_M_to_P(client, ws, mode2):  # noqa: F811
    d = xlsx_diff.diff(BLANK, mode2["out"])
    rep = _report("mode2", "blank", mode2["e"], d)
    gen = set(mode2["rows"])
    allowed = {**{c: gen for c in "AFGH"}, **{c: set(mode2["approved"]) for c in RESULT}}
    assert xlsx_diff.violations(d, "Template", allowed) == []
    cols = rep["modified_columns_from_diff"]
    assert set(cols) == set("AFGHMNOP"), cols
    assert all(cols[c]["cells"] == len(gen) for c in "AFG") and all(cols[c]["cells"] == 4 for c in RESULT)
    assert min(gen) == 7, "header rows 1–6 are never touched"
    assert mode2["e"]["verification"]["mode"] == "blank" and "had no control rows" in rep["mode_explanation"]
    assert {x["column"] for x in rep["modified_columns_from_report"]} == set(cols), "the report lists exactly the columns the diff shows"
    rc = rep["row_counts"]
    assert rc["template_rows"] == len(gen) and rc["rows_written"] == 4 and rc["rows_not_populated"] == len(gen) - 4
    why = {x["row"]: x["reason"] for x in mode2["e"]["verification"]["rows_not_populated"]}
    assert "re-confirm" in why[mode2["unconfirmed"]] and all(why[r] == "No SCA-approved result" for r in sorted(gen)[6:20])


def test_guard_mode2_identifiers_match_reference_data(client, ws, mode2):  # noqa: F811
    t = openpyxl.load_workbook(io.BytesIO(mode2["out"]))["Template"]
    with SessionLocal() as db:
        defs = dict(db.execute(select(Cci.cci_id, Cci.definition)).all())
        cmap: dict = {}
        for cc, ct in db.execute(select(CciControlMap.cci_id, CciControlMap.control_id)):
            cmap.setdefault(cc, set()).add(ct)
        assigned = sorted(db.scalars(select(ControlAssignment.control_id).where(ControlAssignment.system_id == uuid.UUID(ws["sid"]))))
    expected = sum(len({c for c, s in cmap.items() if a in s}) for a in assigned)
    assert len(mode2["rows"]) == expected, "one generated row per CCI of every assigned control"
    for n in sorted(mode2["rows"]):
        ctl_cell, ap, cci, definition = t[f"A{n}"].value, t[f"F{n}"].value, t[f"G{n}"].value, t[f"H{n}"].value
        cid, proc = trs.parse_ap(ap)
        assert cid in assigned and proc and ctl_cell == _ap(cid), (n, ctl_cell, ap)
        full = trs.norm_cci(cci)
        assert cid in cmap[full], f"row {n}: {full} does not map to {cid}"
        assert definition == defs[full], f"row {n}: CCI Definition must come from the DISA CCI reference data"
    first = min(mode2["rows"])
    assert t[f"M{first}"].value == "Compliant" and t[f"O{first}"].value == "Tester 0" and t[f"N{first}"].value.date() == date(2026, 8, 1)
    assert t[f"M{mode2['unconfirmed']}"].value is None, "unconfirmed approval excluded"


def _mode1_template(db_pairs, outsider, ws):  # noqa: F811
    """Populated eMASS-style template built from the real blank fixture: identifiers A–L, a formula in E and after P, formatting in A–L
    and after P, a column width after P; Example / Assignment Values / Instructions untouched."""
    from openpyxl.styles import Font, PatternFill

    wb = openpyxl.load_workbook(io.BytesIO(BLANK))
    t = wb["Template"]
    rows = [(_ap(c), f"{_ap(c)}.{i}", cci.replace("CCI-", "")) for i, (c, cci) in enumerate(db_pairs, start=1)]
    rows += [(_ap(outsider[0]), f"{_ap(outsider[0])}.1", outsider[1].replace("CCI-", "")), ("ZZ-1", "ZZ-1.1", "999999")]
    for i, (ctl, ap, cci) in enumerate(rows):
        r = 7 + i
        t[f"A{r}"], t[f"F{r}"], t[f"G{r}"] = ctl, ap, cci
        for col in "BCDHIJKL":
            t[f"{col}{r}"] = f"eMASS {col}{r}"
        t[f"E{r}"] = f'=A{r}&" narrative"'
        t[f"U{r}"] = f"=LEN(P{r})"
    t["B7"].font = Font(bold=True, color="FF0000")
    t["V8"] = "note after P"
    t["V8"].fill = PatternFill("solid", fgColor="FFFF00")
    t.column_dimensions["U"].width = 31
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), len(rows)


@pytest.fixture(scope="module")
def mode1(client, ws):  # noqa: F811
    with SessionLocal() as db:
        assigned = set(db.scalars(select(ControlAssignment.control_id).where(ControlAssignment.system_id == uuid.UUID(ws["sid"]))))
        cmap: dict = {}
        for cc, ct in db.execute(select(CciControlMap.cci_id, CciControlMap.control_id)):
            cmap.setdefault(cc, set()).add(ct)
        uniq = sorted((next(iter(s)), c) for c, s in cmap.items() if len(s) == 1)
        catalog = trs._ref(db)[1]
    pairs = []
    for ctl, cci in uniq:
        if ctl in assigned and ctl not in {p[0] for p in pairs}:
            pairs.append((ctl, cci))
        if len(pairs) == 5:
            break
    outsider = next((ctl, cci) for ctl, cci in uniq if ctl not in assigned and ctl in catalog)
    data, n = _mode1_template(pairs, outsider, ws)
    wid = _import(client, ws, data, name="TR-populated-guard.xlsx")
    rows = _rows(client, ws, wid)
    _approve(client, ws, wid, [rows[7]["id"], rows[8]["id"], rows[9]["id"]], [("Lee, Sam", "2026-09-01"), ("Lee, Sam", "2026-09-02"), ("Kim, Ana", "2026-09-03")],
             status="Non-Compliant")
    with SessionLocal() as db:
        r = db.get(TrRow, uuid.UUID(rows[9]["id"]))
        r.final = {k: v for k, v in r.final.items() if k != "tester_confirmed"}
        db.commit()
    e, out = _export(client, ws, wid)
    _save("mode1_input_populated_template.xlsx", data)
    _save("mode1_output_populated_template_export.xlsx", out)
    return {"wid": wid, "data": data, "out": out, "e": e, "n": n, "outsider": outsider[0]}


def test_guard_mode1_populated_only_M_to_P(client, ws, mode1):  # noqa: F811
    d = xlsx_diff.diff(mode1["data"], mode1["out"])
    rep = _report("mode1", "prepopulated", mode1["e"], d)
    assert xlsx_diff.violations(d, "Template", {c: {7, 8} for c in RESULT}) == []
    cols = rep["modified_columns_from_diff"]
    assert set(cols) == RESULT and all(c["cells"] == 2 for c in cols.values()), cols
    assert {x["column"] for x in rep["modified_columns_from_report"]} == RESULT
    assert mode1["e"]["verification"]["mode"] == "prepopulated" and "No rows were inserted" in rep["mode_explanation"]
    a, b = openpyxl.load_workbook(io.BytesIO(mode1["data"]))["Template"], openpyxl.load_workbook(io.BytesIO(mode1["out"]))["Template"]
    assert b["E7"].value == '=A7&" narrative"' and b["U8"].value == "=LEN(P8)" and b["B7"].font.b and b["V8"].fill.fgColor.rgb == a["V8"].fill.fgColor.rgb
    assert b.max_row == a.max_row and b.column_dimensions["U"].width == 31
    assert b["M7"].value == "Non-Compliant" and b["O8"].value == "Lee, Sam" and b["N7"].value.date() == date(2026, 9, 1)


def test_guard_mode1_matching_and_unpopulated_report(client, ws, mode1):  # noqa: F811
    v = mode1["e"]["verification"]
    rc = v["row_counts"]
    assert rc == {"template_rows": mode1["n"], "rows_written": 2, "rows_not_populated": mode1["n"] - 2, "unmatched_rows": 1, "needs_mapping_review": 0,
                  "rows_for_controls_not_in_system": 1}, rc
    why = {x["row"]: x["reason"] for x in v["rows_not_populated"]}
    assert "re-confirm" in why[9] and why[10] == why[11] == "No SCA-approved result"
    assert "not in this system's control set" in why[12] and why[13].startswith("Unmatched")
    assert v["unmatched_controls"] == ["ZZ-1"] and v["controls_not_in_system"] == [mode1["outsider"].upper()]
    assert v["system_controls_without_template_row"], "system controls absent from the template are reported"


def test_guard_mode2_output_roundtrip_as_populated_template(client, ws, mode2):  # noqa: F811
    wid = _import(client, ws, mode2["out"], name="TR-roundtrip.xlsx")
    rows = _all_rows(wid)
    assert set(rows) == set(mode2["rows"]) and not any(r["added"] for r in rows.values())
    target = sorted(rows)[10]
    _approve(client, ws, wid, [rows[target]["id"]], [("Roundtrip, Tess", "2026-09-05")])
    e, out = _export(client, ws, wid)
    assert e["verification"]["mode"] == "prepopulated"
    d = xlsx_diff.diff(mode2["out"], out)
    _report("mode2_roundtrip", "prepopulated", e, d)
    _save("mode2_roundtrip_output.xlsx", out)
    already = set(mode2["approved"])  # previously written M–P cells are rewritten with identical values only if approved again
    assert xlsx_diff.violations(d, "Template", {c: {target} | already for c in RESULT}) == []
    assert {c["cell"] for c in d["cells"]["Template"]} >= {f"{c}{target}" for c in "MNOP"}


def test_guard_diff_catches_writes_outside_permitted_columns():
    """Self-test of the guard: any write to an identifier column, a column after P, another sheet or a header is reported."""
    for cells, sheet in (({"L7": "x"}, "Template"), ({"Q7": "x"}, "Template"), ({"M6": "x"}, "Template"), ({"A7": "x"}, "Example")):
        out = xlsx.write(BLANK, sheet, cells)
        assert xlsx_diff.violations(xlsx_diff.diff(BLANK, out), "Template", {c: set(range(7, 3000)) for c in RESULT}), cells
