"""Export Test Plan: importer layout (sheet/headers/no formulas), selected controls only, round trip (same system → unchanged, 0 errors;
fresh system → same procedures), RBAC 403, cross-tenant 404, audit with row count + SHA-256."""
import hashlib
import io
import uuid

import openpyxl
import pytest
from sqlalchemy import func, select

from app.db import SessionLocal
from app.models import AuditEvent, ControlAssignment, TrRow, TrWorkbook
from app.services.tr import test_plan_export as tpe
from tests.conftest import auth
from tests.test_sandbox import _signup, signup_on  # noqa: F401

SYS = {"baseline_key": "low", "categorization_mode": "fips199", "information_types": [{"id": "D.16.3"}]}


@pytest.fixture(scope="module")
def ctx(client, signup_on):  # noqa: F811
    t = _signup(client, f"tpx-{uuid.uuid4().hex[:6]}@example.com", "TPX Org")
    h = auth(t["access_token"])
    role = lambda r: auth(client.post("/api/v1/sandbox/act-as", json={"role": r}, headers=h).json()["access_token"])  # noqa: E731
    isso = role("isso")
    a = client.post("/api/v1/systems", json={"name": "Export Source", "acronym": "TPXA", **SYS}, headers=isso).json()["id"]
    b = client.post("/api/v1/systems", json={"name": "Fresh Copy", "acronym": "TPXB", **SYS}, headers=isso).json()["id"]
    with SessionLocal() as db:  # tailor one control out of A (not selected) – it must not be exported
        gone = db.scalar(select(ControlAssignment).where(ControlAssignment.system_id == uuid.UUID(a)).order_by(ControlAssignment.control_id.desc()))
        removed = gone.control_id
        db.delete(gone)
        db.commit()
    return {"a": a, "b": b, "isso": isso, "role": role, "removed": removed}


def _export(client, ctx, sid, h=None):
    return client.get(f"/api/v1/systems/{sid}/test-plan/export", headers=h or ctx["isso"])


def _rows(data: bytes) -> list[tuple]:
    ws = openpyxl.load_workbook(io.BytesIO(data))[tpe.SHEET]
    return [tuple(c.value for c in r[:5]) for r in ws.iter_rows(min_row=tpe.HEADER_ROW + 1)]


def _post(client, ctx, sid, step, data, name="plan.xlsx"):
    return client.post(f"/api/v1/systems/{sid}/tr/test-plan/{step}", files={"file": (name, data)}, headers=ctx["isso"])


def test_workbook_layout_and_selected_controls_only(client, ctx):
    r = _export(client, ctx, ctx["a"])
    assert r.status_code == 200, r.text
    assert r.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert "TestPlan_TPXA_" in r.headers["content-disposition"]
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Template"]
    ws = wb["Template"]
    assert "pending verification against a real eMASS export" in ws["A1"].value
    assert [c.value for c in ws[tpe.HEADER_ROW]] == tpe.COLUMNS
    assert not [c.coordinate for row in ws.iter_rows() for c in row if c.data_type == "f"], "no formulas"
    with SessionLocal() as db:
        selected = set(db.scalars(select(ControlAssignment.control_id).where(ControlAssignment.system_id == uuid.UUID(ctx["a"]))))
    exported = {tpe.control_acronym(c) for c in selected}
    rows = list(ws.iter_rows(min_row=3, values_only=True))
    assert int(r.headers["x-row-count"]) == len(rows) > 0
    assert {x[0] for x in rows} <= exported and tpe.control_acronym(ctx["removed"]) not in {x[0] for x in rows}, "only selected controls"
    assert all(x[1] and x[2] and x[2].startswith("CCI-") for x in rows)
    assert {x[6] for x in rows} <= {"Yes", "No", "Hybrid (shared)"}
    assert all(x[7] is None for x in rows), "no results before the SCA approves any"
    with SessionLocal() as db:
        ev = db.scalar(select(AuditEvent).where(AuditEvent.action == "test_plan.exported", AuditEvent.entity_id == ctx["a"]).order_by(AuditEvent.occurred_at.desc()))
        assert ev.details["rows"] == len(rows) and ev.details["sha256"] == hashlib.sha256(r.content).hexdigest()
    s = client.get(f"/api/v1/systems/{ctx['a']}/test-plan/export/summary", headers=ctx["isso"]).json()
    assert s["rows"] == len(rows) and s["columns"] == tpe.COLUMNS


def test_round_trip_same_system_then_fresh_copy(client, ctx):
    first = _export(client, ctx, ctx["a"]).content
    p = _post(client, ctx, ctx["a"], "import", first)
    assert p.status_code == 201, p.text
    c = p.json()["preview"]["counts"]
    assert c["rejected"] == 0 and c["accepted"] == len(_rows(first))
    again = _export(client, ctx, ctx["a"]).content
    assert _rows(again) == _rows(first), "export is stable after re-import"
    with SessionLocal() as db:
        before = db.scalar(select(func.count()).select_from(TrRow).join(TrWorkbook).where(TrWorkbook.system_id == uuid.UUID(ctx["a"])))
    pv = _post(client, ctx, ctx["a"], "preview", again).json()
    assert pv["counts"]["rejected"] == 0 and pv["counts"]["accepted"] == 0 and pv["counts"]["unchanged"] == len(_rows(again)), "0 errors"
    r = _post(client, ctx, ctx["a"], "import", again)
    assert r.status_code == 409 and "Nothing to import" in r.json()["detail"]["message"]
    with SessionLocal() as db:
        after = db.scalar(select(func.count()).select_from(TrRow).join(TrWorkbook).where(TrWorkbook.system_id == uuid.UUID(ctx["a"])))
    assert after == before, "no unintended changes"
    # fresh copy: system B has the same control set except the one tailored out of A → same procedures, 0 errors
    fb = _post(client, ctx, ctx["b"], "import", first)
    assert fb.status_code == 201 and fb.json()["preview"]["counts"]["rejected"] == 0
    with SessionLocal() as db:
        got = {(r.control_acronym, r.ap_acronym, r.cci_id) for r in db.scalars(select(TrRow).where(TrRow.workbook_id == uuid.UUID(fb.json()["id"])))}
    assert got == {(x[0], x[1], x[2]) for x in _rows(first)}


def test_rbac_and_tenant_isolation(client, ctx, signup_on):  # noqa: F811
    for r in ("issm", "sca", "ao", "system_owner"):
        assert _export(client, ctx, ctx["a"], ctx["role"](r)).status_code == 200, r
    for r in ("leadership_viewer", "devsecops"):
        assert _export(client, ctx, ctx["a"], ctx["role"](r)).status_code == 403, r
    other = _signup(client, f"tpx-b-{uuid.uuid4().hex[:6]}@example.com", "TPX Other")
    hb = auth(client.post("/api/v1/sandbox/act-as", json={"role": "isso"}, headers=auth(other["access_token"])).json()["access_token"])
    assert _export(client, ctx, ctx["a"], hb).status_code == 404
    assert client.get(f"/api/v1/systems/{ctx['a']}/test-plan/export/summary", headers=hb).status_code == 404
