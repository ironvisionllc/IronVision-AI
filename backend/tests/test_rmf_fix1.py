"""Fix task after b6564d2: server-side IL rules, unknown role keys, excerpt sections/pages, CSV injection end to end,
no-supporting-text path, Fully Addressed only when all requirements MET, logical system archive."""
import csv
import io
import types
import uuid

import docx
import pytest
from docx.enum.text import WD_BREAK
from sqlalchemy import select

from app.db import SessionLocal
from app.models import AuditEvent, OdpValue, Policy, PolicyRequirementCoverage, System
from app.services import mapping_sbs, rmf_accuracy as ra
from app.services import policy_coverage as pc
from app.services import policy_mapping as pm
from tests.conftest import auth
from tests.test_policy_auditor_v2 import _analyze, _upload
from tests.test_sandbox import _signup, signup_on  # noqa: F401

SYS = {"baseline_key": "moderate", "categorization_mode": "fips199", "information_types": [{"id": "D.16.3"}]}


@pytest.fixture(scope="module")
def org(client, signup_on):  # noqa: F811
    t = _signup(client, f"fx-{uuid.uuid4().hex[:6]}@example.com", "Fix Org")
    h = auth(t["access_token"])
    sid = client.post("/api/v1/systems", json={"name": "Fix System", "acronym": f"FX{uuid.uuid4().hex[:3]}", **SYS}, headers=h).json()["id"]
    role = lambda r: auth(client.post("/api/v1/sandbox/act-as", json={"role": r}, headers=h).json()["access_token"])  # noqa: E731
    return {"h": h, "sid": sid, "role": role}


def _post(client, org, **kw):
    return client.post("/api/v1/systems", json={"name": f"S {uuid.uuid4().hex[:4]}", "acronym": f"S{uuid.uuid4().hex[:3]}", **SYS, **kw}, headers=org["h"])


def test_impact_level_rules_enforced_on_create_and_update(client, org):
    r = _post(client, org, operating_environment="on_premises", environment_detail={"impact_level": "IL6"})
    assert r.status_code == 422 and "on_prem_gov" in str(r.json()["detail"]) and "cloud_iaas" in str(r.json()["detail"]), "enum lists allowed values"
    r = _post(client, org, operating_environment="on_prem_gov", environment_detail={"impact_level": "IL6"})
    assert r.status_code == 422 and "cloud" in r.json()["detail"]
    r = _post(client, org, operating_environment="on_prem_gov", environment_detail={"impact_level": "IL2"})
    assert r.status_code == 422 and "no cloud component" in r.json()["detail"]
    r = _post(client, org, operating_environment="cloud_iaas", environment_detail={"impact_level": "IL6"})
    assert r.status_code == 422 and "IL6" in r.json()["detail"] and "CNSSI 1253" in r.json()["detail"]
    r = _post(client, org, operating_environment="cloud_iaas", environment_detail={"impact_level": "IL6"}, categorization_mode="cnssi1253",
              nss_determination="yes", classified_information="no")
    assert r.status_code == 422
    ok = _post(client, org, operating_environment="cloud_iaas", environment_detail={"impact_level": "IL6"}, categorization_mode="cnssi1253",
               nss_determination="yes", classified_information="yes")
    assert ok.status_code == 201 and ok.json()["applicability_flags"]["classified_information"] == "yes"
    il5 = _post(client, org, operating_environment="cloud_iaas", environment_detail={"impact_level": "IL5"}).json()
    url = f"/api/v1/systems/{il5['id']}"
    assert client.patch(url, json={"environment_detail": {"impact_level": "IL6"}}, headers=org["h"]).status_code == 422
    r = client.patch(url, json={"operating_environment": "on_prem_gov"}, headers=org["h"])
    assert r.status_code == 422 and "no cloud component" in r.json()["detail"]
    assert client.patch(url, json={"operating_environment": "on_prem_gov", "environment_detail": {}}, headers=org["h"]).status_code == 200
    env = client.get("/api/openapi.json").json()["components"]["schemas"]["SystemCreate"]["properties"]["operating_environment"]
    assert "cloud_saas" in str(env) and "hybrid" in str(env), "documented enum in OpenAPI"
    with SessionLocal() as db:  # legacy inconsistent record: kept, flagged, and unrelated edits still work
        s = db.get(System, uuid.UUID(il5["id"]))
        s.operating_environment, s.environment_detail = "on_prem_gov", {"impact_level": "IL6"}
        db.commit()
    assert client.patch(url, json={"description": "unrelated edit"}, headers=org["h"]).status_code == 200
    with SessionLocal() as db:
        s = db.get(System, uuid.UUID(il5["id"]))
        assert s.environment_detail["impact_level"] == "IL6" and any(f["field"] == "impact_level" for f in ra.review_flags(s))


def test_unknown_and_alias_role_keys_rejected(client, org):
    allowed = "Allowed at registration: system_owner, isso"
    for keys, want, not_in in ((("ccp", "common_control_provider"), "assign CCP under Inheritance providers", "AODR"),
                               (("aodr", "ao_designated_representative"), "assign AODR under Authorize", "CCP"),
                               (("user_rep", "user_representative", "userrep"), "User Representative isn't assigned in this app's registration", "CCP")):
        for k in keys:
            d = _post(client, org, key_roles={k: {"name": "Pat Doe"}}).json()["detail"]
            assert want in d and allowed in d and not_in not in d.split("Allowed")[0], (k, d)
    assert "AODR" not in _post(client, org, key_roles={"userrep": {"name": "P"}}).json()["detail"].split("Allowed")[0]
    r = _post(client, org, key_roles={"isso": {"name": "Ann"}, "chief_wizard": {"name": "X"}})
    d = r.json()["detail"]
    assert r.status_code == 422 and "Unknown role key(s): chief_wizard" in d and "Allowed at registration: system_owner, isso" in d and "ccp" not in d.split("Allowed")[1]
    r = client.put(f"/api/v1/systems/{org['sid']}/prepare", json={"roles": {"chief_wizard": {"name": "X"}}}, headers=org["h"])
    assert r.status_code == 422 and "Unknown role key(s): chief_wizard" in r.json()["detail"]


def _paged_docx() -> bytes:
    d = docx.Document()
    for t in ("Access Control Policy", "1. Purpose", "This policy establishes access control for the organization."):
        d.add_paragraph(t)
    d.add_paragraph("Introductory text ends here.").runs[0].add_break(WD_BREAK.PAGE)
    for t in ("2. Requirements", "AC-2 Account Management: the ISSO shall review accounts every 30 days and must disable inactive accounts."):
        d.add_paragraph(t)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def test_excerpts_cite_section_and_page_and_old_policies_are_reprocessed(client, org):
    r = client.post("/api/v1/policies", files={"file": ("paged.docx", _paged_docx())}, data={"title": "Paged Policy", "system_ids": org["sid"]}, headers=org["h"])
    assert r.status_code == 201, r.text
    pid = r.json()["id"]
    _analyze(pid)
    ex = client.get(f"/api/v1/policies/{pid}/mappings/ac-2/review", headers=org["h"]).json()["mapping"]["excerpts"][0]
    assert ex["section"] == "AC-2 Account Management" and ex["page"] == 2, "same section finder as coverage snippets"
    t = "Policy\n1. Purpose\nThis policy establishes access control.\n\f3.2 Account reviews\nAccounts shall be reviewed monthly."
    assert mapping_sbs.locate(t, "This policy establishes access control.")["section"] == "1. Purpose"
    loc = mapping_sbs.locate(t, "Accounts shall be reviewed monthly.")
    assert loc["section"] == "3.2 Account reviews" and loc["page"] == 2
    assert mapping_sbs.locate(t.replace("\f", ""), "Accounts shall be reviewed monthly.")["page"] is None, "no page breaks → page not available"
    with SessionLocal() as db:  # a policy extracted by the old extractor (no page breaks, no text_version) is re-extracted once
        p = db.get(Policy, uuid.UUID(pid))
        p.text_content = p.text_content.replace("\f", "")
        p.analysis = {k: v for k, v in (p.analysis or {}).items() if k != "text_version"}
        db.commit()
    assert client.get(f"/api/v1/policies/{pid}/mappings/ac-2/review", headers=org["h"]).json()["mapping"]["excerpts"][0]["page"] is None
    with SessionLocal() as db:
        assert pc.ensure_current(db, [db.get(Policy, uuid.UUID(pid))]) == 1
        db.commit()
        p = db.get(Policy, uuid.UUID(pid))
        from app.services.rmf import TEXT_VERSION
        assert "\f" in p.text_content and p.analysis["text_version"] == TEXT_VERSION
        assert p.analysis_status == "queued", "re-analysis is a queued background job, never inline"
        assert pc.ensure_current(db, [p]) == 0, "idempotent"
        db.commit()
    from worker.handlers import handle_policy_analyze
    with SessionLocal() as db:  # what the worker does with the queued policy.analyze job
        handle_policy_analyze(db, {"policy_id": pid})
        db.commit()
    ex = client.get(f"/api/v1/policies/{pid}/mappings/ac-2/review", headers=org["h"]).json()["mapping"]["excerpts"][0]
    assert ex["page"] == 2 and ex["section"] == "AC-2 Account Management"


def test_excerpt_before_first_heading_is_preamble():
    t = "Risk Assessment Policy and Procedures (RA-1)\nThe organization shall assess risk annually.\n1. Purpose\nRisk assessments shall be documented."
    title = mapping_sbs.locate(t, "Risk Assessment Policy and Procedures (RA-1)")
    assert title["start"] == 0 and title["end"] == 44 and title["section"] == "Preamble / title"
    assert mapping_sbs.locate(t, "The organization shall assess risk annually.")["section"] == "Preamble / title"
    assert mapping_sbs.locate(t, "Risk assessments shall be documented.")["section"] == "1. Purpose"
    far = "Title line (x)\n" + "\n".join(f"text line {i}." for i in range(80)) + "\n2. Scope\n" + "\n".join(f"more {i}." for i in range(60)) + "\nTarget sentence."
    assert mapping_sbs.locate(far, "Target sentence.")["section"] == "2. Scope", "heading beyond the 40-line window is still found"


def _csv(text: str) -> list[dict]:
    return list(csv.DictReader(io.StringIO(text.lstrip("\ufeff"))))


def test_formula_value_exported_with_apostrophe_end_to_end(client, org):
    sid = org["sid"]
    with SessionLocal() as db:
        db.add(OdpValue(system_id=uuid.UUID(sid), control_id="ac-7", parameter="ac-07_odp.01", value="=1+1", source="SSP", approved_by="ISSM"))
        p = Policy(org_id=db.get(System, uuid.UUID(sid)).org_id, policy_key=f"f-{uuid.uuid4().hex[:6]}", title="Formula Policy", version=1, system_ids=[sid],
                   analysis={"values": [{"control_id": "ac-7", "parameter": "ac-07_odp.02", "value": "=1+1", "quote": "=1+1"}]})
        db.add(p)
        db.commit()
        pid = p.id
    rows = {r["Parameter ID"]: r for r in _csv(client.get(f"/api/v1/systems/{sid}/parameters.csv", params={"family": "ac"}, headers=org["h"]).text)}
    assert rows["ac-07_odp.01"]["System ODP value (approved)"] == "'=1+1" and rows["ac-07_odp.01"]["Effective value"] == "'=1+1"
    assert rows["ac-07_odp.02"]["Organization-defined value (policy)"] == "'=1+1"
    prow = [r for r in _csv(client.get(f"/api/v1/policies/{pid}/parameters.csv", headers=org["h"]).text) if r["Parameter ID"] == "ac-07_odp.02"]
    assert prow and prow[0]["Organization-defined value (policy)"] == "'=1+1"


@pytest.fixture(scope="module")
def pol(client, org):
    lines = ["1 Purpose", "This policy establishes the access control policy.", "2 Requirements",
             "AC-2 Account Management: the ISSO shall review accounts every 30 days and must disable inactive accounts."]
    return _upload(client, org["h"], org["sid"], lines, "Fix Access Policy")


def test_no_supporting_text_path(client, org, pol):
    with SessionLocal() as db:
        pm.decide(db, db.get(Policy, uuid.UUID(pol)), "sc-8", "mapped", types.SimpleNamespace(id=None, email="t@example.com"), "manual")
        db.commit()
    r = client.get(f"/api/v1/policies/{pol}/mappings/sc-8/review", headers=org["h"]).json()
    m = r["mapping"]
    assert m["no_supporting_text"] and m["excerpts"] == [] and m["evidence_confidence"] == "Requires manual review"
    assert m["confidence_explanation"] == "No supporting text found – manual review required" == mapping_sbs.NO_TEXT
    assert r["coverage"]["outcome"] == "Requires Manual Review"
    detail = next(x for x in client.get(f"/api/v1/policies/{pol}", headers=org["h"]).json()["control_mappings"] if x["control_id"] == "sc-8")
    assert detail["no_supporting_text"] and detail["evidence_confidence"] == "Requires manual review"


def test_fully_addressed_only_when_every_requirement_met(client, org, pol):
    m = types.SimpleNamespace(coverage_override=None)
    ev = {"no_supporting_text": False, "evidence_confidence": "High"}
    cov = lambda *st: [types.SimpleNamespace(status=s) for s in st]  # noqa: E731
    assert mapping_sbs.outcome(m, cov("met", "met"), ev)["outcome"] == "Fully Addressed"
    assert mapping_sbs.outcome(m, cov("met", "missing"), ev)["outcome"] == "Partially Addressed"
    assert mapping_sbs.outcome(m, cov("missing"), ev)["outcome"] == "Not Addressed"
    url = f"/api/v1/policies/{pol}/mappings/ac-2/review"
    assert client.get(url, headers=org["h"]).json()["mapping"]["evidence_confidence"] == "High"
    with SessionLocal() as db:  # High-confidence mapping, one requirement MISSING → not Fully Addressed
        rows = list(db.scalars(select(PolicyRequirementCoverage).where(PolicyRequirementCoverage.policy_id == uuid.UUID(pol),
                                                                       PolicyRequirementCoverage.control_id == "ac-2")))
        assert len(rows) >= 2
        for c in rows:
            c.status = "met"
        db.commit()
    assert client.get(url, headers=org["h"]).json()["coverage"]["outcome"] == "Fully Addressed"
    with SessionLocal() as db:
        db.get(PolicyRequirementCoverage, rows[0].id).status = "missing"
        db.commit()
    r = client.get(url, headers=org["h"]).json()
    assert r["mapping"]["evidence_confidence"] == "High" and r["coverage"]["outcome"] != "Fully Addressed"
    assert any(g["type"] == "Missing requirement" for g in r["gaps"])


def test_archive_and_restore_system(client, org):
    s = _post(client, org).json()
    url = f"/api/v1/systems/{s['id']}"
    assert client.post(f"{url}/archive", json={"reason": "Junk test system"}, headers=org["role"]("leadership_viewer")).status_code == 403
    assert client.post(f"{url}/archive", json={"reason": "x"}, headers=org["h"]).status_code == 422
    a = client.post(f"{url}/archive", json={"reason": "Junk test system"}, headers=org["h"]).json()
    assert a["archived_at"] and a["archive_reason"] == "Junk test system"
    assert s["id"] not in {x["id"] for x in client.get("/api/v1/systems", headers=org["h"]).json()}
    assert s["id"] in {x["id"] for x in client.get("/api/v1/systems", params={"include_archived": True}, headers=org["h"]).json()}
    assert client.get(url, headers=org["h"]).status_code == 200, "logical: the record is still readable"
    assert client.post(f"{url}/archive", json={"reason": "again please"}, headers=org["h"]).status_code == 409
    r = client.post(f"{url}/restore", json={"reason": "Archived by mistake"}, headers=org["h"]).json()
    assert r["archived_at"] is None and s["id"] in {x["id"] for x in client.get("/api/v1/systems", headers=org["h"]).json()}
    with SessionLocal() as db:
        acts = {e.action for e in db.scalars(select(AuditEvent).where(AuditEvent.entity_id == s["id"]))}
        assert {"system.archived", "system.restored"} <= acts
