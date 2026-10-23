"""eMASS Test Result workflow on the REAL blank template (tests/fixtures/emass/BLANK_TRExport.xlsx)."""
import io
import re
import time
import uuid
import zipfile
from datetime import date, datetime, timedelta, timezone
from pathlib import Path

import openpyxl
import pytest
from sqlalchemy import select

from app.db import SessionLocal
from app.models import Assessment, AuditEvent, ControlAssignment, Finding, TrRow, TrWorkbook, User
from app.services.tr import service as trs
from app.services.tr import xlsx
from tests.conftest import auth
from tests.test_sandbox import _signup, signup_on  # noqa: F401

BLANK = (Path(__file__).parent / "fixtures/emass/BLANK_TRExport.xlsx").read_bytes()


def populated(rows: list[tuple], base: bytes = BLANK) -> bytes:
    """An eMASS-style export: the blank template with AP rows (Control/AP/CCI) filled – written with the same XML engine."""
    cells = {}
    for i, (ctl, ap, cci) in enumerate(rows):
        r = 7 + i
        cells.update({f"A{r}": ctl, f"F{r}": ap, f"G{r}": cci})
    return xlsx.write(base, "Template", cells)


def test_detect_real_blank_template():
    d = xlsx.detect(BLANK)
    assert d["sheet"] == "Template" and d["header_row"] == 6
    assert {k: d["roles"][k] for k in xlsx.RESULT_ROLES} == {"compliance_status": "M", "date_tested": "N", "tested_by": "O", "test_results": "P"}
    assert d["roles"]["ap_acronym"] == "F" and d["roles"]["cci"] == "G"
    latest = {c["col"]: c["role"] for c in d["columns"] if (c["role"] or "").startswith("latest_")}
    assert latest == {"Q": "latest_compliance_status", "R": "latest_date_tested", "S": "latest_tested_by", "T": "latest_test_results"}
    assert d["vocabulary"]["compliance_status"] == ["Compliant", "Non-Compliant", "Not Applicable"] and "Example" in d["vocabulary"]["source"]
    assert {"Q", "R", "S", "T", "F", "G"} <= set(d["readonly_cols"]) and not {"M", "N", "O", "P"} & set(d["readonly_cols"])
    assert d["rows"] == [] and d["limits"]["test_results"] == 4000 and d["limits"]["tested_by"] == 150
    assert [s["title"] for s in d["sheets"]] == ["Template", "Example", "Assignment Values", "Instructions"]


def test_detect_shuffled_and_renamed_variant():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "My TR"
    ws.append(["Our org TR sheet"])
    ws.append(["Results", "CCI #", "Assessor", "AP", "Test Date", "Status", "Notes"])
    ws.append(["", "000015", "", "AC-2(1).1", "", "", ""])
    buf = io.BytesIO()
    wb.save(buf)
    d = xlsx.detect(buf.getvalue())
    assert d["header_row"] == 2
    assert d["roles"] == {"ap_acronym": "D", "cci": "B", "compliance_status": "F", "date_tested": "E", "tested_by": "C", "test_results": "A"}
    assert d["rows"][0]["ap_acronym"] == "AC-2(1).1" and d["vocabulary"]["compliance_status"] == xlsx.DEFAULT_VOCAB


@pytest.fixture(scope="module")
def ref():
    with SessionLocal() as db:
        return trs._ref(db)


@pytest.mark.parametrize("ap,cci,ctl,status,control", [
    ("AC-1.1", "000002", None, "matched", "ac-1"),
    ("AC-2(1).1", "000015", None, "matched", "ac-2.1"),           # enhancement AP acronym
    ("AC-2.1", "000015", None, "ambiguous", "ac-2"),              # CCI belongs to the enhancement
    ("AC-1.1", "999999", None, "unmatched", None),                # CCI not in the DISA list
    ("ZZ-9.1", "000002", None, "unmatched", None),                # control not in the catalog
    ("garbage", "000002", None, "ambiguous", "ac-1"),             # AP unreadable, single-control CCI → confirm
    (None, "002153", None, "ambiguous", None),                    # missing AP, CCI maps to several controls
    (None, "000002", None, "ambiguous", "ac-1"),                  # missing AP, single candidate still needs confirmation
    ("AC-1.1", None, None, "ambiguous", "ac-1"),                  # missing CCI
    ("AC-3.1", "000002", None, "ambiguous", "ac-3"),              # CCI maps to a different control
])
def test_ap_cci_matching(ref, ap, cci, ctl, status, control):
    m = trs.match(*ref, ap, cci, ctl)
    assert m["match_status"] == status and m["control_id"] == control, m
    assert m["match_reason"]


@pytest.fixture(scope="module")
def ws(client, signup_on):  # noqa: F811
    t = _signup(client, f"tr-{uuid.uuid4().hex[:6]}@example.com", "TR Org")
    h = auth(t["access_token"])
    r = client.post("/api/v1/systems", json={"name": "TR System", "acronym": f"TR{uuid.uuid4().hex[:3]}", "baseline_key": "moderate",
                                             "categorization_mode": "fips199", "information_types": [{"id": "D.16.3"}]}, headers=h)
    sid = r.json()["id"]
    from tests.test_processing_pipeline import FILES, _process
    run = client.post(f"/api/v1/systems/{sid}/uploads", files={"file": ("w.ckl", FILES["ckl"].read_bytes())}, headers=h).json()["id"]
    _process(run)
    with SessionLocal() as db:
        f = db.scalars(select(Finding).where(Finding.system_id == uuid.UUID(sid), Finding.source_type == "stig_checklist")).all()
        from app.models import CciControlMap
        assigned = set(db.scalars(select(ControlAssignment.control_id).where(ControlAssignment.system_id == uuid.UUID(sid))))
        cmap = {}
        for cc, ct in db.execute(select(CciControlMap.cci_id, CciControlMap.control_id)):
            cmap.setdefault(cc, set()).add(ct)
        stig_cci, ctl = next((c, next(iter(cmap[c]))) for x in sorted(f, key=lambda x: x.rule_id or "") for c in sorted(x.ccis or [])
                             if len(cmap.get(c, ())) == 1 and next(iter(cmap[c])) in assigned)
    role = lambda r: auth(client.post("/api/v1/sandbox/act-as", json={"role": r}, headers=h).json()["access_token"])  # noqa: E731
    return {"h": h, "sid": sid, "stig_cci": stig_cci, "stig_ctl": ctl, "role": role}


def _ap(ctl):
    return ctl.upper().replace(".", "(") + (")" if "." in ctl else "")


def _import(client, ws, data, name="TR.xlsx"):
    from worker.handlers import handle_tr_detect, handle_tr_generate

    r = client.post(f"/api/v1/systems/{ws['sid']}/tr/workbooks", files={"file": (name, data)}, headers=ws["h"])
    assert r.status_code == 201, r.text
    wid = r.json()["id"]
    with SessionLocal() as db:
        handle_tr_detect(db, {"workbook_id": wid})
        db.commit()
        handle_tr_generate(db, {"workbook_id": wid})
        db.commit()
    return wid


def _rows(client, ws, wid, **p):
    return {r["row"]: r for r in client.get(f"/api/v1/tr/workbooks/{wid}/rows", params={"size": 500, **p}, headers=ws["h"]).json()["items"]}


@pytest.fixture(scope="module")
def wb(client, ws):
    c = ws["stig_ctl"]
    data = populated([(_ap(c).split("(")[0], f"{_ap(c)}.1", ws["stig_cci"].replace("CCI-", "")),   # row 7: STIG data
                      ("PE-3", "PE-3.1", "000919"),                                                # row 8: no data at all
                      ("ZZ-1", "ZZ-1.1", "999999"),                                                # row 9: unmatched
                      ("AC-3", "AC-3.1", "000002")])                                               # row 10: ambiguous
    return _import(client, ws, data)


def test_generation_uses_real_data_and_never_fabricates(client, ws, wb):
    rows = _rows(client, ws, wb)
    r7, r8, r9, r10 = rows[7], rows[8], rows[9], rows[10]
    assert r7["proposed"]["compliance_status"] is None and r7["proposed"]["suggestion"] in ("Compliant", "Non-Compliant")
    assert any(s["type"] == "finding" and s["label"].startswith("STIG ") for s in r7["sources"]) and "STIG" in r7["proposed"]["test_results"]
    assert "assessor_input_required" in r7["flags"] and r7["proposed"]["tested_by"] is None and r7["proposed"]["date_tested"] is None
    assert r8["state"] == "missing" and not r8["sources"] and r8["proposed"]["compliance_status"] is None and r8["proposed"]["tested_by"] is None
    assert r9["match_status"] == "unmatched" and "needs_mapping_review" in r9["flags"] and r9["state"] == "missing"
    assert r10["match_status"] == "ambiguous" and not r10["sources"], "uncertain matches are not used until reviewed"


@pytest.mark.parametrize("role", ["isso", "issm", "ao", "system_owner", "devsecops", "leadership_viewer"])
def test_only_sca_can_approve_or_set_status(client, ws, wb, role):
    h = ws["role"](role)
    rid = _rows(client, ws, wb)[7]["id"]
    for body in ({"action": "approve", "row_ids": [rid], "use_suggestion": True},
                 {"action": "approve", "row_ids": [rid, rid], "justification": "bulk approve", "use_suggestion": True},
                 {"action": "edit", "row_ids": [rid], "justification": "set status", "values": {"compliance_status": "Compliant"}}):
        r = client.post(f"/api/v1/tr/workbooks/{wb}/decide", json=body, headers=h)
        assert r.status_code == 403 and "Only an assessor (SCA role)" in r.json()["detail"], (role, r.text)
    assert _rows(client, ws, wb)[7]["state"] != "approved"


def test_sca_approval_partial_export_and_validation_report(client, ws, wb):
    sca = ws["role"]("sca")
    rows = _rows(client, ws, wb)
    body = {"action": "approve", "row_ids": [rows[7]["id"], rows[8]["id"]], "justification": "Reviewed STIG results", "use_suggestion": True}
    for missing, msg in (({}, "Date tested"), ({"tested_by": "Lee, Sam"}, "Date tested"), ({"date_tested": "2026-09-01"}, "Tested by"),
                         ({"tested_by": "   ", "date_tested": "2026-09-01"}, "Tested by"),
                         ({"tested_by": "Lee, Sam", "date_tested": (date.today() + timedelta(days=1)).isoformat()}, "future")):
        r = client.post(f"/api/v1/tr/workbooks/{wb}/decide", json={**body, **missing}, headers=sca)
        assert r.status_code == 422 and msg in r.json()["detail"], (missing, r.text)
    assert _rows(client, ws, wb)[7]["state"] != "approved"
    before = client.get(f"/api/v1/tr/workbooks/{wb}", headers=ws["h"]).json()["counts"]["state"].get("approved", 0)
    r = client.post(f"/api/v1/tr/workbooks/{wb}/decide", json={**body, "tested_by": "Lee, Sam", "date_tested": "2026-09-01"}, headers=sca)
    assert r.status_code == 200 and r.json()["approved"] == [7] and r.json()["skipped"][0]["row"] == 8, r.text
    assert r.json()["counts"]["state"]["approved"] == before + 1, "counts reflect this action (computed after the commit)"
    with SessionLocal() as db:
        row = db.get(TrRow, uuid.UUID(rows[7]["id"]))
        assert row.final["date_tested"] == "2026-09-01" and row.final["tested_by"] == "Lee, Sam" and row.final["tester_confirmed"]
        assert row.final["approved_by"] != row.final["tested_by"] and row.final["approved_at"] and row.final["tester_source"] == "entered by the SCA at approval"
        ev = db.scalar(select(AuditEvent).where(AuditEvent.action == "tr.row_approved", AuditEvent.entity_id == rows[7]["id"]))
        assert ev.after["tester"] == "Lee, Sam" and ev.after["test_date"] == "2026-09-01" and ev.after["approver"] and ev.created_at
    e = client.post(f"/api/v1/tr/workbooks/{wb}/export", headers=ws["h"])
    assert e.status_code == 200, e.text
    v = e.json()["verification"]
    assert v["ok"] and not v["outside_permitted"] and v["issues"] == []
    assert set(v["changed_cells"]) == {"M7", "N7", "O7", "P7"}, v["changed_cells"]
    assert v["permitted_range"] == ["Template!M7:M2714", "Template!N7:N2714", "Template!O7:O2714", "Template!P7:P2714"]
    assert v["mode"] == "prepopulated" and all("Result column" in x["why"] for x in v["permitted_ranges"]) and "never written" in v["proof"]
    out = client.get(f"/api/v1/tr/exports/{e.json()['id']}/download", headers=ws["h"]).content
    wbo = openpyxl.load_workbook(io.BytesIO(out))
    t = wbo["Template"]
    assert t["M7"].value == row.final["compliance_status"] and t["M7"].value in ("Compliant", "Non-Compliant")
    assert t["N7"].value.date() == date(2026, 9, 1) and t["O7"].value == "Lee, Sam", "entered tester/date – not the approver, not today"
    assert all(t[f"{c}8"].value is None for c in "MNOP"), "unapproved rows stay blank"
    assert wbo.sheetnames == ["Template", "Example", "Assignment Values", "Instructions"]
    za, zb = zipfile.ZipFile(io.BytesIO(BLANK)), zipfile.ZipFile(io.BytesIO(out))
    assert [i.filename for i in za.infolist()] == [i.filename for i in zb.infolist()] and zb.testzip() is None
    rep = client.get(f"/api/v1/tr/exports/{e.json()['id']}/validation-report", headers=ws["h"]).json()
    assert rep["ok"] and rep["changed_cells"] == v["changed_cells"] and rep["human_required"]


def test_repeated_export_is_deterministic_with_history_and_audit(client, ws, wb):
    a = client.post(f"/api/v1/tr/workbooks/{wb}/export", headers=ws["h"]).json()
    b = client.post(f"/api/v1/tr/workbooks/{wb}/export", headers=ws["h"]).json()
    assert a["sha256"] == b["sha256"] and b["version"] == a["version"] + 1
    d = client.get(f"/api/v1/tr/workbooks/{wb}", headers=ws["h"]).json()
    assert [x["version"] for x in d["exports"]][:2] == [b["version"], a["version"]] and all(x["validation_ok"] for x in d["exports"])
    with SessionLocal() as db:
        assert db.scalar(select(AuditEvent).where(AuditEvent.action == "tr.exported", AuditEvent.entity_id == b["id"])) is not None
        assert db.scalar(select(AuditEvent).where(AuditEvent.action == "tr.row_approved")) is not None


def test_assessment_record_is_the_source_of_status_tester_and_date(client, ws):
    """The exporter and today's date are never substituted for the tester and the test date."""
    c = ws["stig_ctl"]
    with SessionLocal() as db:
        a = db.scalar(select(ControlAssignment).where(ControlAssignment.system_id == uuid.UUID(ws["sid"]), ControlAssignment.control_id == c))
        if a is None:
            pytest.skip("STIG control not in baseline")
        from app.models import System
        sca = User(org_id=db.get(System, uuid.UUID(ws["sid"])).org_id, email=f"pat.assessor-{uuid.uuid4().hex[:4]}@example.com", display_name="Pat Assessor")
        db.add(sca)
        db.flush()
        tested = datetime(2026, 9, 14, 15, 0, tzinfo=timezone.utc)
        db.add(Assessment(system_id=a.system_id, assignment_id=a.id, control_id=c, assessor_id=sca.id, objectives=[], methods=["examine"],
                          findings_text="Examined the STIG checklist.", result="satisfied", created_at=tested))
        db.commit()
    wid = _import(client, ws, populated([(_ap(c), f"{_ap(c)}.1", ws["stig_cci"].replace("CCI-", ""))]), name="TR-assess.xlsx")
    r = _rows(client, ws, wid)[7]
    assert r["proposed"]["compliance_status"] == "Compliant" and r["proposed"]["tested_by"] == "Assessor, Pat" and r["proposed"]["date_tested"] == "2026-09-14"
    assert any(s["type"] == "assessment" for s in r["sources"]) and "Assessor notes" in r["proposed"]["test_results"]
    if any(s["label"].startswith("STIG") and " Open " in s["label"] for s in r["sources"]):
        assert "conflict" in r["flags"] and r["state"] == "conflict"
    sca_h = ws["role"]("sca")
    assert client.post(f"/api/v1/tr/workbooks/{wid}/decide", json={"action": "approve", "row_ids": [r["id"]]}, headers=sca_h).status_code == 422, \
        "an assessment record pre-fills, but the SCA must still confirm the tester and date"
    ok = client.post(f"/api/v1/tr/workbooks/{wid}/decide", json={"action": "approve", "row_ids": [r["id"]], "tested_by": r["proposed"]["tested_by"],
                                                                 "date_tested": r["proposed"]["date_tested"]}, headers=sca_h).json()
    assert ok["approved"] == [7]
    with SessionLocal() as db:
        assert db.get(TrRow, uuid.UUID(r["id"])).final["tester_source"].startswith("confirmed from assessment:")
    e = client.post(f"/api/v1/tr/workbooks/{wid}/export", headers=ws["h"]).json()
    t = openpyxl.load_workbook(io.BytesIO(client.get(f"/api/v1/tr/exports/{e['id']}/download", headers=ws["h"]).content))["Template"]
    assert t["O7"].value == "Assessor, Pat" and t["N7"].value.date() == date(2026, 9, 14) and t["M7"].value == "Compliant"


def test_blank_template_populated_from_system(client, ws):
    from worker.handlers import handle_tr_detect, handle_tr_generate

    wid = client.post(f"/api/v1/systems/{ws['sid']}/tr/workbooks", files={"file": ("BLANK_TRExport.xlsx", BLANK)}, headers=ws["h"]).json()["id"]
    with SessionLocal() as db:
        handle_tr_detect(db, {"workbook_id": wid})
        db.commit()
        assert db.get(TrWorkbook, uuid.UUID(wid)).counts["rows"] == 0
        handle_tr_generate(db, {"workbook_id": wid, "populate": True})
        db.commit()
    rows = _rows(client, ws, wid)
    assert rows and all(r["added"] and "derived_ap" in r["flags"] for r in rows.values())
    e = client.post(f"/api/v1/tr/workbooks/{wid}/export", headers=ws["h"]).json()
    v = e["verification"]
    assert v["ok"] and {c[0] for c in v["changed_cells"]} <= {"A", "F", "G", "H"}, "no result cells without approval"
    assert v["mode"] == "blank" and client.get(f"/api/v1/tr/workbooks/{wid}", headers=ws["h"]).json()["export_mode"] == "blank"
    assert any("Blank-template mode only" in x["why"] and x["column"] == "cci_definition" for x in v["permitted_ranges"])
    ids = [r["id"] for r in list(rows.values())[:3]]
    per = {i: {"tested_by": f"Tester {n}", "date_tested": f"2026-08-0{n + 1}"} for n, i in enumerate(ids)}
    sca = ws["role"]("sca")
    d = client.post(f"/api/v1/tr/workbooks/{wid}/decide", json={"action": "approve", "row_ids": ids, "justification": "bulk per-row testers",
                                                                "status": "Compliant", "per_row": per}, headers=sca)
    assert d.status_code == 200 and len(d.json()["approved"]) == 3, d.text
    e = client.post(f"/api/v1/tr/workbooks/{wid}/export", headers=ws["h"]).json()
    t = openpyxl.load_workbook(io.BytesIO(client.get(f"/api/v1/tr/exports/{e['id']}/download", headers=ws["h"]).content))["Template"]
    first = min(r["row"] for r in rows.values())
    assert t[f"O{first}"].value == "Tester 0" and t[f"N{first}"].value.date() == date(2026, 8, 1) and t[f"M{first}"].value == "Compliant"
    assert all(t[f"{c}{first}"].value is None for c in "BCDEIJKL"), "B–E and I–L stay blank in blank-template mode"
    assert t[f"A{first}"].value and t[f"G{first}"].value


def test_prepopulated_mode_leaves_a_to_l_untouched(client, ws):
    """eMASS-style pre-populated export: every Control-Information column A–L filled; only M–P may change."""
    c = ws["stig_ctl"]
    cells = {}
    for i in range(3):
        r = 7 + i
        cells.update({f"A{r}": _ap(c).split("(")[0], f"F{r}": f"{_ap(c)}.{i + 1}", f"G{r}": ws["stig_cci"].replace("CCI-", "")})
        cells.update({f"{col}{r}": f"eMASS {col}{r} text" for col in "BCDEHIJKL"})
    data = xlsx.write(BLANK, "Template", cells)
    wid = _import(client, ws, data, name="TR-emass-prepop.xlsx")
    rows = _rows(client, ws, wid)
    assert client.get(f"/api/v1/tr/workbooks/{wid}", headers=ws["h"]).json()["export_mode"] == "prepopulated"
    sca = ws["role"]("sca")
    r = client.post(f"/api/v1/tr/workbooks/{wid}/decide", json={"action": "approve", "row_ids": [rows[7]["id"], rows[8]["id"]], "status": "Non-Compliant",
                                                                "justification": "SCA test", "tested_by": "Lee, Sam", "date_tested": "2026-09-02"}, headers=sca)
    assert r.status_code == 200 and r.json()["approved"] == [7, 8], r.text
    e = client.post(f"/api/v1/tr/workbooks/{wid}/export", headers=ws["h"]).json()
    v = e["verification"]
    assert v["ok"] and v["mode"] == "prepopulated" and not v["outside_permitted"]
    assert {re.match(r"[A-Z]+", x).group(0) for x in v["changed_cells"]} <= set("MNOP") and len(v["permitted_ranges"]) == 4
    a = openpyxl.load_workbook(io.BytesIO(data))["Template"]
    b = openpyxl.load_workbook(io.BytesIO(client.get(f"/api/v1/tr/exports/{e['id']}/download", headers=ws["h"]).content))["Template"]
    for row in range(1, a.max_row + 1):
        for col in "ABCDEFGHIJKL":
            assert a[f"{col}{row}"].value == b[f"{col}{row}"].value and a[f"{col}{row}"].style_id == b[f"{col}{row}"].style_id, f"{col}{row}"
    assert b["M7"].value == "Non-Compliant" and b["O8"].value == "Lee, Sam" and b["M9"].value is None


def test_legacy_approvals_need_tester_confirmation(client, ws, wb):
    rid = _rows(client, ws, wb)[7]["id"]
    with SessionLocal() as db:
        row = db.get(TrRow, uuid.UUID(rid))
        row.state, row.final = "needs_confirmation", {k: v for k, v in (row.final or {}).items() if k != "tester_confirmed"}
        db.commit()
    e = client.post(f"/api/v1/tr/workbooks/{wb}/export", headers=ws["h"]).json()
    assert not any(x.endswith("7") and x[0] in "MNOP" and len(x) == 2 for x in e["verification"]["changed_cells"]), "excluded until confirmed"
    r = client.post(f"/api/v1/tr/workbooks/{wb}/decide", json={"action": "approve", "row_ids": [rid], "tested_by": "Lee, Sam", "date_tested": "2026-09-01"},
                    headers=ws["role"]("sca"))
    assert r.json()["approved"] == [7] and _rows(client, ws, wb)[7]["state"] == "approved"


def test_tenant_isolation(client, ws, wb, signup_on):  # noqa: F811
    h2 = auth(_signup(client, f"tr-other-{uuid.uuid4().hex[:6]}@example.com", "TR Other")["access_token"])
    assert client.get(f"/api/v1/tr/workbooks/{wb}", headers=h2).status_code == 404
    assert client.get(f"/api/v1/tr/workbooks/{wb}/rows", headers=h2).status_code == 404
    assert client.post(f"/api/v1/tr/workbooks/{wb}/export", headers=h2).status_code == 404
    assert client.get(f"/api/v1/systems/{ws['sid']}/tr/workbooks", headers=h2).status_code in (403, 404)


def test_large_template_background_job(client, ws):
    rows = [("AC-2", f"AC-2.{i}", "000015") for i in range(5200)]
    data = populated(rows)
    t0 = time.time()
    wid = _import(client, ws, data, name="TR-large.xlsx")
    assert time.time() - t0 < 240
    d = client.get(f"/api/v1/tr/workbooks/{wid}", headers=ws["h"]).json()
    assert d["counts"]["rows"] == 5200 and d["status"] == "generated"


_ = (timedelta,)
