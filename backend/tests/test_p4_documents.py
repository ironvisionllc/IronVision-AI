"""P4: SSP / SAR / RAR are generated from the system's actual records as DOCX and PDF, with a missing-information check and DRAFT /
DRAFT – INCOMPLETE labelling; the SAR never invents results and the RAR never invents an AO decision; unknown kinds/formats → 422; RBAC."""
import io
import uuid

import pytest
from pypdf import PdfReader

from tests.conftest import auth
from tests.test_sandbox import _signup, signup_on  # noqa: F401


@pytest.fixture(scope="module")
def ctx(client, signup_on):  # noqa: F811
    t = _signup(client, f"p4-{uuid.uuid4().hex[:6]}@example.com", "P4 Org")
    h = auth(t["access_token"])
    sid = client.post("/api/v1/systems", json={"name": "Doc System", "acronym": f"DS{uuid.uuid4().hex[:3]}", "baseline_key": "low",
                                                "categorization_mode": "fips199", "information_types": [{"id": "D.16.3"}]}, headers=h).json()["id"]
    return {"sid": sid, "isso": auth(client.post("/api/v1/sandbox/act-as", json={"role": "isso"}, headers=h).json()["access_token"])}


@pytest.mark.parametrize("kind", ["ssp", "sar", "rar"])
def test_check_reports_missing_and_incomplete(client, ctx, kind):
    r = client.get(f"/api/v1/systems/{ctx['sid']}/documents/{kind}/check", headers=ctx["isso"]).json()
    assert r["status"] == "DRAFT – INCOMPLETE" and r["missing"]
    if kind == "ssp":
        assert any("no implementation statement" in m for m in r["missing"])
    if kind == "sar":
        assert any("no SCA assessment record" in m for m in r["missing"])


def test_docx_and_pdf_contain_real_data_and_labels(client, ctx):
    import docx
    r = client.get(f"/api/v1/systems/{ctx['sid']}/documents/ssp?format=docx", headers=ctx["isso"])
    assert r.status_code == 200 and "DRAFT-INCOMPLETE" in r.headers["content-disposition"]
    d = docx.Document(io.BytesIO(r.content))
    body = "\n".join(p.text for p in d.paragraphs) + "\n".join(c.text for t in d.tables for row in t.rows for c in row.cells)
    assert "Doc System" in body and "AC-1" in body and "No statement recorded" in body and "Missing information" in body
    assert d.sections[0].header.paragraphs[0].text.startswith("DRAFT – INCOMPLETE")
    p = client.get(f"/api/v1/systems/{ctx['sid']}/documents/rar?format=pdf", headers=ctx["isso"])
    assert p.status_code == 200 and p.headers["content-type"] == "application/pdf"
    txt = "".join(pg.extract_text() for pg in PdfReader(io.BytesIO(p.content)).pages)
    assert "DRAFT" in txt and "No authorization decision has been recorded by the AO" in txt


def test_sar_does_not_fabricate_results(client, ctx):
    import docx
    r = client.get(f"/api/v1/systems/{ctx['sid']}/documents/sar?format=docx", headers=ctx["isso"])
    body = "\n".join(p.text for p in docx.Document(io.BytesIO(r.content)).paragraphs)
    assert "No SCA assessment results are recorded" in body and "satisfied:" not in body


def test_bad_kind_format_and_rbac(client, ctx):
    assert client.get(f"/api/v1/systems/{ctx['sid']}/documents/xyz/check", headers=ctx["isso"]).status_code == 422
    assert client.get(f"/api/v1/systems/{ctx['sid']}/documents/ssp?format=exe", headers=ctx["isso"]).status_code == 422
    t = _signup(client, f"p4o-{uuid.uuid4().hex[:6]}@example.com", "Other Org")
    assert client.get(f"/api/v1/systems/{ctx['sid']}/documents/ssp/check", headers=auth(t["access_token"])).status_code in (403, 404)
